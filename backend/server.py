from fastapi import FastAPI, APIRouter, HTTPException, Depends, status, BackgroundTasks, Response, Request, Cookie, UploadFile, File, Form
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Any
import uuid
from datetime import datetime, timezone, timedelta
from passlib.context import CryptContext
from jose import JWTError, jwt
import re
import httpx
import asyncio
import io
import zipfile
import base64
import json

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# JWT Settings
SECRET_KEY = os.environ.get('SECRET_KEY', 'habari-stays-secret-key-change-in-production')
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24 * 7  # 7 days

# Password hashing
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

# Security
security = HTTPBearer(auto_error=False)

# Beem SMS Configuration
BEEM_API_KEY = os.environ.get('BEEM_API_KEY', '')
BEEM_SECRET_KEY = os.environ.get('BEEM_SECRET_KEY', '')
BEEM_SENDER_ID = os.environ.get('BEEM_SENDER_ID', 'HABARISTAYS')

# Payment Configuration
SELCOM_TILL_NUMBER = os.environ.get('SELCOM_TILL_NUMBER', '123456')

# Cloudinary Configuration
CLOUDINARY_CLOUD_NAME = os.environ.get('CLOUDINARY_CLOUD_NAME', 'diupey6vs')
CLOUDINARY_UPLOAD_PRESET = os.environ.get('CLOUDINARY_UPLOAD_PRESET', 'habari_stays_upload')

# Booking expiry time in seconds
BOOKING_EXPIRY_SECONDS = 120  # 2 minutes

# Helper: extract cover photo URL from photos list (handles old string[] and new object[] format)
def get_cover_url(photos, size="cloudinary_mobile"):
    if not photos:
        return None
    # Find primary photo first
    for p in photos:
        if isinstance(p, dict) and p.get("is_primary"):
            return p.get(size) or p.get("cloudinary_original") or p.get("cloudinary_web")
    # Fall back to first photo
    first = photos[0]
    if isinstance(first, dict):
        return first.get(size) or first.get("cloudinary_original") or first.get("cloudinary_web")
    return first if isinstance(first, str) else None

# Hotel-level and Room-level amenity constants
HOTEL_AMENITIES = ["Breakfast", "Parking", "WiFi", "Hot Water", "Bar"]
ROOM_AMENITIES = ["A/C", "Western Toilet", "Squat Toilet", "En-suite Bathroom", "Balcony", "TV", "Safe", "Mini Fridge", "Hot Shower"]

async def create_default_room_type(hotel_id: str):
    """Create a default Standard room type for a newly imported hotel."""
    room_id = str(uuid.uuid4())
    room = {
        "id": room_id,
        "hotel_id": hotel_id,
        "name": "Standard",
        "description": "",
        "price_per_night": 50000,
        "capacity": 2,
        "total_rooms": 1,
        "available_rooms": 1,
        "amenities": [],
        "photos": [],
        "is_default": True,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.room_types.insert_one(room)

app = FastAPI(title="Habari Stays API", version="2.0.0")
api_router = APIRouter(prefix="/api")

# ===================== MODELS =====================

class UserBase(BaseModel):
    email: EmailStr
    full_name: str
    phone: str
    role: str = "traveler"  # traveler, owner, admin, cashier
    assigned_hotel_id: Optional[str] = None

class UserCreate(UserBase):
    password: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserResponse(BaseModel):
    id: str
    email: str
    full_name: str
    phone: str
    role: str
    assigned_hotel_id: Optional[str] = None
    is_verified: bool = True
    picture: Optional[str] = None
    created_at: str

class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    user: UserResponse

class HotelBase(BaseModel):
    name: str
    description: str
    address: str
    city: str
    phone_number: str
    whatsapp_number: Optional[str] = None
    amenities: List[str] = []
    cover_photo: Optional[str] = None
    photos: list = []
    google_maps_url: Optional[str] = None
    google_rating: Optional[float] = None
    google_review_count: Optional[int] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None

class HotelCreate(HotelBase):
    pass

class HotelResponse(HotelBase):
    id: str
    hotel_code: str
    owner_id: Optional[str] = None
    status: str = "pending"  # pending, verified, suspended
    data_source: str = "manual"
    created_at: str
    total_rooms: int = 0
    available_rooms: int = 0
    bookings_this_month: int = 0
    revenue_this_month: float = 0
    average_rating: float = 0
    review_count: int = 0
    min_price: Optional[int] = None

class RoomTypeBase(BaseModel):
    name: str
    description: str = ""
    price_per_night: int
    capacity: int = 2
    total_rooms: int
    available_rooms: int
    amenities: List[str] = []
    photos: list = []
    is_default: bool = False

class RoomTypeCreate(RoomTypeBase):
    hotel_id: str

class RoomTypeResponse(RoomTypeBase):
    id: str
    hotel_id: str
    is_default: bool = False
    created_at: str

class BookingCreate(BaseModel):
    hotel_id: str
    room_type_id: str
    guest_name: str
    guest_phone: str
    guest_email: Optional[EmailStr] = None
    checkin_date: str
    checkout_date: str
    num_guests: int = 1
    booking_type: str = "online"  # online, walkin
    payment_method: str = "mpesa"  # mpesa, cash, card
    notes: Optional[str] = None

class BookingResponse(BaseModel):
    id: str
    booking_ref: str
    hotel_id: str
    hotel_name: str
    room_type_id: str
    room_type_name: str
    guest_name: str
    guest_phone: str
    guest_email: Optional[str] = None
    checkin_date: str
    checkout_date: str
    nights: int
    num_guests: int
    total_amount_tzs: int
    booking_type: str
    payment_method: str
    payment_status: str
    payment_reference: Optional[str] = None
    status: str
    checkin_status: str
    actual_checkin_time: Optional[str] = None
    actual_checkout_time: Optional[str] = None
    created_by: Optional[str] = None
    created_by_name: Optional[str] = None
    notes: Optional[str] = None
    created_at: str
    expires_at: Optional[str] = None

class ReviewCreate(BaseModel):
    booking_id: str
    hotel_id: str
    rating: int
    comment: str

class ReviewResponse(BaseModel):
    id: str
    booking_id: str
    hotel_id: str
    guest_name: str
    guest_phone: str
    rating: int
    comment: str
    created_at: str

class CashierCreate(BaseModel):
    full_name: str
    phone: str
    email: EmailStr
    hotel_id: str

class CashierResponse(BaseModel):
    id: str
    full_name: str
    phone: str
    email: str
    role: str
    assigned_hotel_id: Optional[str] = None
    assigned_hotel_name: str
    is_active: bool
    created_at: str
    last_login: Optional[str] = None

class ShiftSummary(BaseModel):
    cashier_id: str
    cashier_name: str
    date: str
    walkins_count: int
    walkins_total: int
    online_checkins_count: int
    checkouts_count: int
    cash_collected: int

class PaymentConfirmation(BaseModel):
    booking_id: str
    payment_reference: str

# ===================== UTILITIES =====================

def generate_booking_ref():
    year = datetime.now().year
    random_num = str(uuid.uuid4().int)[:4]
    return f"HS-{year}-{random_num}"

def generate_hotel_code():
    random_num = str(uuid.uuid4().int)[:3]
    return f"HTL-{random_num}"

def generate_temp_password():
    return str(uuid.uuid4())[:8]

def verify_password(plain_password, hashed_password):
    return pwd_context.verify(plain_password, hashed_password)

def get_password_hash(password):
    return pwd_context.hash(password)

def create_access_token(data: dict):
    to_encode = data.copy()
    expire = datetime.now(timezone.utc) + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)

async def get_current_user(
    request: Request,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    # Check cookie first
    session_token = request.cookies.get("session_token")
    
    if session_token:
        # Validate session from database
        session = await db.user_sessions.find_one({"session_token": session_token}, {"_id": 0})
        if session:
            expires_at = session.get("expires_at")
            if isinstance(expires_at, str):
                expires_at = datetime.fromisoformat(expires_at)
            if expires_at.tzinfo is None:
                expires_at = expires_at.replace(tzinfo=timezone.utc)
            if expires_at > datetime.now(timezone.utc):
                user = await db.users.find_one({"id": session["user_id"]}, {"_id": 0})
                if user:
                    return user
    
    # Fallback to JWT token
    if credentials:
        token = credentials.credentials
        try:
            payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
            user_id: str = payload.get("sub")
            if user_id:
                user = await db.users.find_one({"id": user_id}, {"_id": 0})
                if user:
                    return user
        except JWTError:
            pass
    
    raise HTTPException(status_code=401, detail="Not authenticated")

async def get_current_user_optional(
    request: Request,
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    try:
        return await get_current_user(request, credentials)
    except HTTPException:
        return None

def validate_tanzania_phone(phone: str) -> bool:
    pattern = r'^\+255[6-8]\d{8}$'
    return bool(re.match(pattern, phone))

def format_phone(phone: str) -> str:
    cleaned = re.sub(r'[\s\-\(\)]', '', phone.strip())
    if cleaned.startswith("+255"):
        return cleaned
    if cleaned.startswith("0"):
        return "+255" + cleaned[1:]
    if cleaned.startswith("255"):
        return "+" + cleaned
    return "+255" + cleaned

def format_tzs(amount: int) -> str:
    return f"TZS {amount:,}"

async def send_sms(phone: str, message: str):
    """Send SMS via Beem Africa API"""
    if not BEEM_API_KEY or not BEEM_SECRET_KEY:
        logging.info(f"SMS (mock): To {phone}: {message}")
        return True
    
    try:
        formatted_phone = format_phone(phone)
        # Remove + for Beem API
        if formatted_phone.startswith("+"):
            formatted_phone = formatted_phone[1:]
        
        url = "https://apisms.beem.africa/v1/send"
        headers = {
            "Content-Type": "application/json"
        }
        payload = {
            "source_addr": BEEM_SENDER_ID,
            "encoding": 0,
            "message": message,
            "recipients": [{"recipient_id": 1, "dest_addr": formatted_phone}]
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(
                url,
                json=payload,
                auth=(BEEM_API_KEY, BEEM_SECRET_KEY),
                timeout=30
            )
            logging.info(f"SMS sent to {formatted_phone}: {response.status_code}")
            return response.status_code == 200
    except Exception as e:
        logging.error(f"SMS failed to {phone}: {str(e)}")
        # Fallback to mock
        logging.info(f"SMS (mock fallback): To {phone}: {message}")
        return False

# ===================== BACKGROUND TASKS =====================

async def expire_pending_bookings():
    """Background task to expire unpaid bookings after 2 minutes"""
    while True:
        try:
            now = datetime.now(timezone.utc).isoformat()
            # Find expired bookings
            expired = await db.bookings.find({
                "status": "pending",
                "payment_status": "unpaid",
                "expires_at": {"$lt": now}
            }, {"_id": 0}).to_list(100)
            
            for booking in expired:
                # Update booking status
                await db.bookings.update_one(
                    {"id": booking["id"]},
                    {"$set": {"status": "expired"}}
                )
                
                # Restore room availability
                await db.room_types.update_one(
                    {"id": booking["room_type_id"]},
                    {"$inc": {"available_rooms": 1}}
                )
                
                # Send expiry SMS
                sms_msg = f"HABARI STAYS: Booking {booking['booking_ref']} imeisha muda. Jaribu tena: habaristays.com"
                await send_sms(booking["guest_phone"], sms_msg)
                
                logging.info(f"Expired booking: {booking['booking_ref']}")
        except Exception as e:
            logging.error(f"Expiry task error: {str(e)}")
        
        await asyncio.sleep(60)  # Check every 60 seconds

# ===================== AUTH ENDPOINTS =====================

@api_router.post("/auth/register", response_model=TokenResponse)
async def register(user_data: UserCreate):
    existing = await db.users.find_one({"email": user_data.email})
    if existing:
        raise HTTPException(status_code=400, detail="Email tayari imetumika / Email already exists")
    
    user_id = str(uuid.uuid4())
    is_verified = user_data.role not in ["owner"]  # Owners need verification
    
    user = {
        "id": user_id,
        "email": user_data.email,
        "full_name": user_data.full_name,
        "phone": format_phone(user_data.phone),
        "role": user_data.role,
        "assigned_hotel_id": user_data.assigned_hotel_id,
        "password_hash": get_password_hash(user_data.password),
        "is_active": True,
        "is_verified": is_verified,
        "picture": None,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "last_login": None
    }
    await db.users.insert_one(user)
    
    if user_data.role == "owner":
        return TokenResponse(
            access_token="",
            user=UserResponse(
                id=user["id"],
                email=user["email"],
                full_name=user["full_name"],
                phone=user["phone"],
                role=user["role"],
                assigned_hotel_id=user.get("assigned_hotel_id"),
                is_verified=False,
                picture=None,
                created_at=user["created_at"]
            )
        )
    
    token = create_access_token({"sub": user_id, "role": user["role"]})
    return TokenResponse(
        access_token=token,
        user=UserResponse(
            id=user["id"],
            email=user["email"],
            full_name=user["full_name"],
            phone=user["phone"],
            role=user["role"],
            assigned_hotel_id=user.get("assigned_hotel_id"),
            is_verified=user["is_verified"],
            picture=None,
            created_at=user["created_at"]
        )
    )

@api_router.post("/auth/login", response_model=TokenResponse)
async def login(credentials: UserLogin, response: Response):
    user = await db.users.find_one({"email": credentials.email}, {"_id": 0})
    if not user or not verify_password(credentials.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email au neno la siri si sahihi / Invalid credentials")
    
    if not user.get("is_active", True):
        raise HTTPException(status_code=403, detail="Akaunti yako imezimwa / Account disabled")
    
    if user.get("role") == "owner" and not user.get("is_verified", False):
        raise HTTPException(status_code=403, detail="Akaunti yako inasubiri uthibitishaji / Account pending verification")
    
    await db.users.update_one(
        {"id": user["id"]},
        {"$set": {"last_login": datetime.now(timezone.utc).isoformat()}}
    )
    
    token = create_access_token({"sub": user["id"], "role": user["role"]})
    
    # Set cookie for session
    response.set_cookie(
        key="session_token",
        value=token,
        httponly=True,
        secure=True,
        samesite="none",
        path="/",
        max_age=60*60*24*7  # 7 days
    )
    
    return TokenResponse(
        access_token=token,
        user=UserResponse(
            id=user["id"],
            email=user["email"],
            full_name=user["full_name"],
            phone=user["phone"],
            role=user["role"],
            assigned_hotel_id=user.get("assigned_hotel_id"),
            is_verified=user.get("is_verified", True),
            picture=user.get("picture"),
            created_at=user["created_at"]
        )
    )

@api_router.get("/auth/me", response_model=UserResponse)
async def get_me(current_user: dict = Depends(get_current_user)):
    return UserResponse(
        id=current_user["id"],
        email=current_user["email"],
        full_name=current_user["full_name"],
        phone=current_user["phone"],
        role=current_user["role"],
        assigned_hotel_id=current_user.get("assigned_hotel_id"),
        is_verified=current_user.get("is_verified", True),
        picture=current_user.get("picture"),
        created_at=current_user["created_at"]
    )

@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie(key="session_token", path="/")
    return {"message": "Logged out"}

# Google OAuth Session Exchange
@api_router.post("/auth/session")
async def exchange_session(request: Request, response: Response):
    """Exchange Emergent session_id for user session"""
    body = await request.json()
    session_id = body.get("session_id")
    
    if not session_id:
        raise HTTPException(status_code=400, detail="session_id required")
    
    # Call Emergent Auth to get user data
    async with httpx.AsyncClient() as client:
        try:
            res = await client.get(
                "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
                headers={"X-Session-ID": session_id},
                timeout=10
            )
            if res.status_code != 200:
                raise HTTPException(status_code=401, detail="Invalid session")
            
            user_data = res.json()
        except Exception as e:
            logging.error(f"OAuth session exchange error: {str(e)}")
            raise HTTPException(status_code=401, detail="Session exchange failed")
    
    # Check if user exists
    existing_user = await db.users.find_one({"email": user_data["email"]}, {"_id": 0})
    
    if existing_user:
        user_id = existing_user["id"]
        # Update user data
        await db.users.update_one(
            {"id": user_id},
            {"$set": {
                "full_name": user_data.get("name", existing_user["full_name"]),
                "picture": user_data.get("picture"),
                "last_login": datetime.now(timezone.utc).isoformat()
            }}
        )
    else:
        # Create new user
        user_id = str(uuid.uuid4())
        new_user = {
            "id": user_id,
            "email": user_data["email"],
            "full_name": user_data.get("name", ""),
            "phone": "",
            "role": "traveler",
            "assigned_hotel_id": None,
            "password_hash": "",
            "is_active": True,
            "is_verified": True,
            "picture": user_data.get("picture"),
            "created_at": datetime.now(timezone.utc).isoformat(),
            "last_login": datetime.now(timezone.utc).isoformat()
        }
        await db.users.insert_one(new_user)
    
    # Create session
    session_token = user_data.get("session_token", str(uuid.uuid4()))
    expires_at = datetime.now(timezone.utc) + timedelta(days=7)
    
    await db.user_sessions.insert_one({
        "user_id": user_id,
        "session_token": session_token,
        "expires_at": expires_at.isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat()
    })
    
    # Set cookie
    response.set_cookie(
        key="session_token",
        value=session_token,
        httponly=True,
        secure=True,
        samesite="none",
        path="/",
        max_age=60*60*24*7
    )
    
    # Get user for response
    user = await db.users.find_one({"id": user_id}, {"_id": 0})
    
    return {
        "user": UserResponse(
            id=user["id"],
            email=user["email"],
            full_name=user["full_name"],
            phone=user["phone"],
            role=user["role"],
            assigned_hotel_id=user.get("assigned_hotel_id"),
            is_verified=user.get("is_verified", True),
            picture=user.get("picture"),
            created_at=user["created_at"]
        )
    }

# ===================== HOTEL ENDPOINTS =====================

@api_router.post("/hotels", response_model=HotelResponse)
async def create_hotel(hotel_data: HotelCreate, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi kuunda hotel")
    
    hotel_id = str(uuid.uuid4())
    hotel = {
        "id": hotel_id,
        "hotel_code": generate_hotel_code(),
        "owner_id": current_user["id"],
        **hotel_data.model_dump(),
        "status": "verified" if current_user["role"] == "admin" else "pending",
        "data_source": "manual",
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.hotels.insert_one(hotel)
    
    return HotelResponse(
        id=hotel["id"],
        hotel_code=hotel["hotel_code"],
        owner_id=hotel["owner_id"],
        name=hotel["name"],
        description=hotel["description"],
        address=hotel["address"],
        city=hotel["city"],
        phone_number=hotel["phone_number"],
        whatsapp_number=hotel.get("whatsapp_number"),
        amenities=hotel.get("amenities", []),
        cover_photo=hotel.get("cover_photo"),
        photos=hotel.get("photos", []),
        google_maps_url=hotel.get("google_maps_url"),
        google_rating=hotel.get("google_rating"),
        google_review_count=hotel.get("google_review_count"),
        latitude=hotel.get("latitude"),
        longitude=hotel.get("longitude"),
        status=hotel["status"],
        data_source=hotel["data_source"],
        created_at=hotel["created_at"]
    )

@api_router.get("/hotels", response_model=List[HotelResponse])
async def get_hotels(
    city: Optional[str] = None,
    min_price: Optional[int] = None,
    max_price: Optional[int] = None,
    amenities: Optional[str] = None,
    sort_by: Optional[str] = "rating",
    verified_only: bool = True
):
    query = {}
    if city:
        query["city"] = {"$regex": city, "$options": "i"}
    if verified_only:
        query["status"] = {"$in": ["verified", "imported", "owner_attached"]}
    
    hotels = await db.hotels.find(query, {"_id": 0}).to_list(100)
    result = []
    
    for hotel in hotels:
        # Get room types for price filtering
        rooms = await db.room_types.find({"hotel_id": hotel["id"]}, {"_id": 0}).to_list(100)
        
        if min_price or max_price:
            prices = [r["price_per_night"] for r in rooms]
            if not prices:
                continue
            min_room_price = min(prices) if prices else 0
            if min_price and min_room_price < min_price:
                continue
            if max_price and min_room_price > max_price:
                continue
        
        if amenities:
            required = amenities.split(",")
            if not all(a in hotel.get("amenities", []) for a in required):
                continue
        
        total_rooms = sum(r.get("total_rooms", 0) for r in rooms)
        available_rooms = sum(r.get("available_rooms", 0) for r in rooms)
        
        # Get reviews
        reviews = await db.reviews.find({"hotel_id": hotel["id"]}, {"_id": 0}).to_list(100)
        avg_rating = sum(r["rating"] for r in reviews) / len(reviews) if reviews else 0
        
        # Get bookings this month
        start_of_month = datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        bookings_count = await db.bookings.count_documents({
            "hotel_id": hotel["id"],
            "created_at": {"$gte": start_of_month.isoformat()}
        })
        
        # Get revenue
        revenue_pipeline = [
            {"$match": {"hotel_id": hotel["id"], "created_at": {"$gte": start_of_month.isoformat()}, "payment_status": "paid"}},
            {"$group": {"_id": None, "total": {"$sum": "$total_amount_tzs"}}}
        ]
        revenue_result = await db.bookings.aggregate(revenue_pipeline).to_list(1)
        revenue = revenue_result[0]["total"] if revenue_result else 0
        
        # Get min price from room types
        room_prices = [r["price_per_night"] for r in rooms if r.get("price_per_night", 0) > 0]
        min_price = min(room_prices) if room_prices else None
        
        result.append(HotelResponse(
            id=hotel["id"],
            hotel_code=hotel.get("hotel_code", ""),
            owner_id=hotel.get("owner_id"),
            name=hotel["name"],
            description=hotel["description"],
            address=hotel["address"],
            city=hotel["city"],
            phone_number=hotel.get("phone_number", hotel.get("phone", "")),
            whatsapp_number=hotel.get("whatsapp_number"),
            amenities=hotel.get("amenities", []),
            cover_photo=hotel.get("cover_photo"),
            photos=hotel.get("photos", hotel.get("images", [])),
            google_maps_url=hotel.get("google_maps_url"),
            google_rating=hotel.get("google_rating"),
            google_review_count=hotel.get("google_review_count"),
            latitude=hotel.get("latitude"),
            longitude=hotel.get("longitude"),
            status=hotel.get("status", "verified"),
            data_source=hotel.get("data_source", "manual"),
            created_at=hotel["created_at"],
            total_rooms=total_rooms,
            available_rooms=available_rooms,
            bookings_this_month=bookings_count,
            revenue_this_month=revenue,
            average_rating=round(avg_rating, 1),
            review_count=len(reviews),
            min_price=min_price
        ))
    
    # Sort
    if sort_by == "price_low":
        result.sort(key=lambda x: x.min_price or float('inf'))
    elif sort_by == "price_high":
        result.sort(key=lambda x: x.min_price or 0, reverse=True)
    elif sort_by == "rating":
        result.sort(key=lambda x: x.average_rating, reverse=True)
    
    return result

@api_router.get("/hotels/cities")
async def get_cities():
    """Get cities with hotel counts"""
    cities = await db.hotels.aggregate([
        {"$match": {"status": {"$in": ["verified", "imported", "owner_attached"]}}},
        {"$group": {"_id": "$city", "count": {"$sum": 1}}},
        {"$sort": {"_id": 1}}
    ]).to_list(50)
    
    return [{"city": c["_id"], "hotel_count": c["count"]} for c in cities if c["_id"]]

@api_router.get("/hotels/{hotel_id}", response_model=HotelResponse)
async def get_hotel(hotel_id: str):
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    rooms = await db.room_types.find({"hotel_id": hotel["id"]}, {"_id": 0}).to_list(100)
    total_rooms = sum(r.get("total_rooms", 0) for r in rooms)
    available_rooms = sum(r.get("available_rooms", 0) for r in rooms)
    
    reviews = await db.reviews.find({"hotel_id": hotel["id"]}, {"_id": 0}).to_list(100)
    avg_rating = sum(r["rating"] for r in reviews) / len(reviews) if reviews else 0
    
    room_prices = [r.get("price_per_night", 0) for r in rooms if r.get("price_per_night", 0) > 0]
    min_price = min(room_prices) if room_prices else None
    
    return HotelResponse(
        id=hotel["id"],
        hotel_code=hotel.get("hotel_code", ""),
        owner_id=hotel.get("owner_id"),
        name=hotel["name"],
        description=hotel["description"],
        address=hotel["address"],
        city=hotel["city"],
        phone_number=hotel.get("phone_number", hotel.get("phone", "")),
        whatsapp_number=hotel.get("whatsapp_number"),
        amenities=hotel.get("amenities", []),
        cover_photo=hotel.get("cover_photo"),
        photos=hotel.get("photos", hotel.get("images", [])),
        google_maps_url=hotel.get("google_maps_url"),
        google_rating=hotel.get("google_rating"),
        google_review_count=hotel.get("google_review_count"),
        latitude=hotel.get("latitude"),
        longitude=hotel.get("longitude"),
        status=hotel.get("status", "verified"),
        data_source=hotel.get("data_source", "manual"),
        created_at=hotel["created_at"],
        total_rooms=total_rooms,
        available_rooms=available_rooms,
        average_rating=round(avg_rating, 1),
        review_count=len(reviews),
        min_price=min_price
    )

@api_router.get("/hotels/slug/{city}/{slug}")
async def get_hotel_by_slug(city: str, slug: str):
    """Get hotel by SEO slug"""
    # Convert slug to searchable name
    name_pattern = slug.replace("-", " ")
    hotel = await db.hotels.find_one({
        "city": {"$regex": city, "$options": "i"},
        "name": {"$regex": name_pattern, "$options": "i"},
        "status": "verified"
    }, {"_id": 0})
    
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    return await get_hotel(hotel["id"])

@api_router.get("/owner/hotels")
async def get_owner_hotels(current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    query = {"owner_id": current_user["id"]} if current_user["role"] == "owner" else {}
    hotels = await db.hotels.find(query, {"_id": 0}).to_list(100)
    
    result = []
    for hotel in hotels:
        rooms = await db.room_types.find({"hotel_id": hotel["id"]}, {"_id": 0}).to_list(100)
        total_rooms = sum(r.get("total_rooms", 0) for r in rooms)
        available_rooms = sum(r.get("available_rooms", 0) for r in rooms)
        
        start_of_month = datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        bookings = await db.bookings.count_documents({
            "hotel_id": hotel["id"],
            "created_at": {"$gte": start_of_month.isoformat()}
        })
        
        revenue_pipeline = [
            {"$match": {"hotel_id": hotel["id"], "created_at": {"$gte": start_of_month.isoformat()}, "payment_status": "paid"}},
            {"$group": {"_id": None, "total": {"$sum": "$total_amount_tzs"}}}
        ]
        revenue_result = await db.bookings.aggregate(revenue_pipeline).to_list(1)
        revenue = revenue_result[0]["total"] if revenue_result else 0
        
        reviews = await db.reviews.find({"hotel_id": hotel["id"]}, {"_id": 0}).to_list(100)
        avg_rating = sum(r["rating"] for r in reviews) / len(reviews) if reviews else 0
        
        hotel_data = {
            "id": hotel["id"],
            "hotel_code": hotel.get("hotel_code", ""),
            "owner_id": hotel.get("owner_id"),
            "name": hotel["name"],
            "description": hotel.get("description", ""),
            "address": hotel.get("address", ""),
            "city": hotel.get("city", ""),
            "phone_number": hotel.get("phone_number", hotel.get("phone", "")),
            "whatsapp_number": hotel.get("whatsapp_number"),
            "amenities": hotel.get("amenities", []),
            "cover_photo": hotel.get("cover_photo"),
            "photos": hotel.get("photos", hotel.get("images", [])),
            "google_maps_url": hotel.get("google_maps_url"),
            "google_rating": hotel.get("google_rating"),
            "google_review_count": hotel.get("google_review_count"),
            "latitude": hotel.get("latitude"),
            "longitude": hotel.get("longitude"),
            "status": hotel.get("status", "verified"),
            "data_source": hotel.get("data_source", "manual"),
            "created_at": hotel["created_at"],
            "total_rooms": total_rooms,
            "available_rooms": available_rooms,
            "bookings_this_month": bookings,
            "revenue_this_month": revenue,
            "average_rating": round(avg_rating, 1),
            "review_count": len(reviews),
            "room_type_count": len(rooms),
            "has_default_rooms": any(r.get("is_default") for r in rooms),
        }
        result.append(hotel_data)
    return result

@api_router.put("/hotels/{hotel_id}", response_model=HotelResponse)
async def update_hotel(hotel_id: str, hotel_data: HotelCreate, current_user: dict = Depends(get_current_user)):
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] != "admin" and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi kuhariri hotel hii")
    
    await db.hotels.update_one({"id": hotel_id}, {"$set": hotel_data.model_dump()})
    return await get_hotel(hotel_id)

@api_router.get("/hotels/{hotel_id}/full")
async def get_hotel_full(hotel_id: str):
    """Full hotel detail with room types, photos, amenities - used by admin/owner edit pages"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    rooms = await db.room_types.find({"hotel_id": hotel_id}, {"_id": 0}).to_list(100)
    reviews = await db.reviews.find({"hotel_id": hotel_id}, {"_id": 0}).to_list(100)
    
    owner = None
    if hotel.get("owner_id"):
        owner = await db.users.find_one({"id": hotel["owner_id"]}, {"_id": 0, "password_hash": 0})
    
    start_of_month = datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    bookings_count = await db.bookings.count_documents({
        "hotel_id": hotel_id, "created_at": {"$gte": start_of_month.isoformat()}
    })
    
    return {
        **hotel,
        "room_types": rooms,
        "reviews": reviews,
        "owner": owner,
        "total_rooms": sum(r.get("total_rooms", 0) for r in rooms),
        "available_rooms": sum(r.get("available_rooms", 0) for r in rooms),
        "room_type_count": len(rooms),
        "has_default_rooms": any(r.get("is_default") for r in rooms),
        "bookings_this_month": bookings_count,
        "average_rating": round(sum(r["rating"] for r in reviews) / len(reviews), 1) if reviews else 0,
        "review_count": len(reviews)
    }

@api_router.patch("/hotels/{hotel_id}")
async def patch_hotel(hotel_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    """Partial update of hotel - respects owner vs admin permissions"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    is_admin = current_user["role"] == "admin"
    is_owner = hotel.get("owner_id") == current_user["id"]
    if not is_admin and not is_owner:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    body = await request.json()
    
    # Owner can edit these fields
    owner_editable = ["description", "phone_number", "whatsapp_number", "website", "amenities", "address", "google_maps_url"]
    # Admin can edit everything
    admin_editable = owner_editable + ["name", "city", "status", "google_rating", "google_review_count", "latitude", "longitude"]
    
    allowed = admin_editable if is_admin else owner_editable
    update = {k: v for k, v in body.items() if k in allowed}
    
    if not update:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    # Validate amenities if provided
    if "amenities" in update:
        update["amenities"] = [a for a in update["amenities"] if a in HOTEL_AMENITIES]
    
    await db.hotels.update_one({"id": hotel_id}, {"$set": update})
    updated = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    return updated

@api_router.get("/config/amenities")
async def get_amenity_options():
    """Return available amenity options for hotel and room"""
    return {
        "hotel_amenities": HOTEL_AMENITIES,
        "room_amenities": ROOM_AMENITIES
    }

# ===================== ROOM TYPE ENDPOINTS =====================

@api_router.post("/hotels/{hotel_id}/room-types")
async def create_hotel_room_type(hotel_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    """Create a new room type for a specific hotel"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    if current_user["role"] != "admin" and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    body = await request.json()
    room_id = str(uuid.uuid4())
    room = {
        "id": room_id,
        "hotel_id": hotel_id,
        "name": body.get("name", ""),
        "description": body.get("description", ""),
        "price_per_night": body.get("price_per_night", 0),
        "capacity": body.get("capacity", 2),
        "total_rooms": body.get("total_rooms", 1),
        "available_rooms": body.get("available_rooms", 1),
        "amenities": [a for a in body.get("amenities", []) if a in ROOM_AMENITIES],
        "photos": [],
        "is_default": False,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.room_types.insert_one(room)
    room.pop("_id", None)
    return room

@api_router.get("/room-types/{room_id}")
async def get_room_type_detail(room_id: str):
    """Full room type detail"""
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    return room

@api_router.patch("/room-types/{room_id}")
async def patch_room_type(room_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    """Update room type - flips is_default to false on first edit"""
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    
    hotel = await db.hotels.find_one({"id": room["hotel_id"]}, {"_id": 0})
    if current_user["role"] != "admin" and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    body = await request.json()
    updatable = ["name", "description", "price_per_night", "capacity", "total_rooms", "available_rooms", "amenities"]
    update = {}
    for field in updatable:
        if field in body:
            update[field] = body[field]
    
    if "amenities" in update:
        update["amenities"] = [a for a in update["amenities"] if a in ROOM_AMENITIES]
    
    if "available_rooms" in update and "total_rooms" in update:
        if update["available_rooms"] > update["total_rooms"]:
            update["available_rooms"] = update["total_rooms"]
    
    # Flip is_default on first edit
    if room.get("is_default"):
        update["is_default"] = False
    
    if not update:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    await db.room_types.update_one({"id": room_id}, {"$set": update})
    updated = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    return updated

@api_router.get("/room-types", response_model=List[RoomTypeResponse])
async def get_room_types(hotel_id: str):
    rooms = await db.room_types.find({"hotel_id": hotel_id}, {"_id": 0}).to_list(100)
    return [RoomTypeResponse(**r) for r in rooms]

@api_router.put("/room-types/{room_id}", response_model=RoomTypeResponse)
async def update_room_type(room_id: str, room_data: RoomTypeBase, current_user: dict = Depends(get_current_user)):
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    
    hotel = await db.hotels.find_one({"id": room["hotel_id"]}, {"_id": 0})
    if current_user["role"] != "admin" and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    await db.room_types.update_one({"id": room_id}, {"$set": room_data.model_dump()})
    updated = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    return RoomTypeResponse(**updated)

@api_router.put("/room-types/{room_id}/availability")
async def update_room_availability(room_id: str, change: int, current_user: dict = Depends(get_current_user)):
    """Quick +/- availability update"""
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Room not found")
    
    new_available = max(0, min(room["total_rooms"], room["available_rooms"] + change))
    await db.room_types.update_one({"id": room_id}, {"$set": {"available_rooms": new_available}})
    return {"available_rooms": new_available}

@api_router.delete("/room-types/{room_id}")
async def delete_room_type(room_id: str, current_user: dict = Depends(get_current_user)):
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    
    hotel = await db.hotels.find_one({"id": room["hotel_id"]}, {"_id": 0})
    if current_user["role"] != "admin" and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    # Check for active bookings
    active = await db.bookings.count_documents({
        "room_type_id": room_id,
        "status": {"$in": ["confirmed", "pending"]},
        "checkin_status": {"$ne": "checked_out"}
    })
    if active > 0:
        raise HTTPException(status_code=400, detail="Huwezi kufuta chumba chenye bookings hai.")
    
    await db.room_types.delete_one({"id": room_id})
    return {"message": "Aina ya chumba imefutwa"}

# ===================== BOOKING ENDPOINTS =====================

@api_router.post("/bookings", response_model=BookingResponse)
async def create_booking(
    booking_data: BookingCreate,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user_optional)
):
    hotel = await db.hotels.find_one({"id": booking_data.hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    room_type = await db.room_types.find_one({"id": booking_data.room_type_id}, {"_id": 0})
    if not room_type:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    
    if room_type.get("available_rooms", 0) < 1:
        raise HTTPException(status_code=400, detail="Hakuna vyumba vilivyopo / No rooms available")
    
    # Calculate nights and amount
    checkin = datetime.fromisoformat(booking_data.checkin_date)
    checkout = datetime.fromisoformat(booking_data.checkout_date)
    nights = (checkout - checkin).days
    if nights < 1:
        raise HTTPException(status_code=400, detail="Tarehe si sahihi / Invalid dates")
    
    total_amount = nights * room_type["price_per_night"]
    
    booking_id = str(uuid.uuid4())
    booking_ref = generate_booking_ref()
    
    is_walkin = booking_data.booking_type == "walkin"
    now = datetime.now(timezone.utc)
    expires_at = None if is_walkin else (now + timedelta(seconds=BOOKING_EXPIRY_SECONDS)).isoformat()
    
    booking = {
        "id": booking_id,
        "booking_ref": booking_ref,
        "hotel_id": booking_data.hotel_id,
        "room_type_id": booking_data.room_type_id,
        "guest_name": booking_data.guest_name,
        "guest_phone": format_phone(booking_data.guest_phone),
        "guest_email": booking_data.guest_email,
        "checkin_date": booking_data.checkin_date,
        "checkout_date": booking_data.checkout_date,
        "nights": nights,
        "num_guests": booking_data.num_guests,
        "total_amount_tzs": total_amount,
        "booking_type": booking_data.booking_type,
        "payment_method": booking_data.payment_method,
        "payment_status": "paid" if is_walkin else "unpaid",
        "payment_reference": None,
        "status": "confirmed" if is_walkin else "pending",
        "checkin_status": "checked_in" if is_walkin else "not_checked_in",
        "actual_checkin_time": now.isoformat() if is_walkin else None,
        "actual_checkout_time": None,
        "created_by": current_user["id"] if current_user else None,
        "notes": booking_data.notes,
        "created_at": now.isoformat(),
        "expires_at": expires_at
    }
    
    await db.bookings.insert_one(booking)
    
    # Decrement available rooms
    await db.room_types.update_one(
        {"id": booking_data.room_type_id},
        {"$inc": {"available_rooms": -1}}
    )
    
    # Send SMS notifications
    if is_walkin and current_user:
        # Notify owner about walk-in
        owner = await db.users.find_one({"id": hotel.get("owner_id")}, {"_id": 0})
        if owner:
            cashier_name = current_user["full_name"]
            sms_msg = f"HABARI STAYS: Walk-in Mpya! Mgeni: {booking_data.guest_name}, Chumba: {room_type['name']}, Usiku: {nights}, Kiasi: {format_tzs(total_amount)}. Mweka Hazina: {cashier_name}. Ref: {booking_ref}"
            background_tasks.add_task(send_sms, owner["phone"], sms_msg)
    elif not is_walkin:
        # Notify owner about online booking
        owner = await db.users.find_one({"id": hotel.get("owner_id")}, {"_id": 0})
        if owner:
            sms_msg = f"HABARI STAYS: Booking Mpya! Mgeni: {booking_data.guest_name}, Chumba: {room_type['name']}, Check-in: {booking_data.checkin_date}, Kiasi: {format_tzs(total_amount)}. Ref: {booking_ref}"
            background_tasks.add_task(send_sms, owner["phone"], sms_msg)
    
    creator_name = current_user["full_name"] if current_user and is_walkin else "Online System"
    
    return BookingResponse(
        id=booking["id"],
        booking_ref=booking["booking_ref"],
        hotel_id=booking["hotel_id"],
        hotel_name=hotel["name"],
        room_type_id=booking["room_type_id"],
        room_type_name=room_type["name"],
        guest_name=booking["guest_name"],
        guest_phone=booking["guest_phone"],
        guest_email=booking.get("guest_email"),
        checkin_date=booking["checkin_date"],
        checkout_date=booking["checkout_date"],
        nights=booking["nights"],
        num_guests=booking["num_guests"],
        total_amount_tzs=booking["total_amount_tzs"],
        booking_type=booking["booking_type"],
        payment_method=booking["payment_method"],
        payment_status=booking["payment_status"],
        payment_reference=booking.get("payment_reference"),
        status=booking["status"],
        checkin_status=booking["checkin_status"],
        actual_checkin_time=booking.get("actual_checkin_time"),
        actual_checkout_time=booking.get("actual_checkout_time"),
        created_by=booking.get("created_by"),
        created_by_name=creator_name,
        notes=booking.get("notes"),
        created_at=booking["created_at"],
        expires_at=booking.get("expires_at")
    )

@api_router.post("/bookings/{booking_id}/confirm-payment")
async def confirm_payment(
    booking_id: str,
    payment: PaymentConfirmation,
    background_tasks: BackgroundTasks
):
    """Confirm payment for online booking"""
    booking = await db.bookings.find_one({"id": booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    if booking["status"] == "expired":
        raise HTTPException(status_code=400, detail="Booking imeisha muda / Booking expired")
    
    if booking["payment_status"] == "paid":
        raise HTTPException(status_code=400, detail="Malipo yamepokelewa tayari / Already paid")
    
    # Update booking
    await db.bookings.update_one(
        {"id": booking_id},
        {"$set": {
            "payment_status": "paid",
            "payment_reference": payment.payment_reference,
            "status": "confirmed"
        }}
    )
    
    # Send confirmation SMS to guest
    hotel = await db.hotels.find_one({"id": booking["hotel_id"]}, {"_id": 0})
    sms_msg = f"HABARI STAYS: Booking imethibitishwa! Hotel: {hotel['name']}, Ref: {booking['booking_ref']}. Asante! habaristays.com"
    background_tasks.add_task(send_sms, booking["guest_phone"], sms_msg)
    
    return {"message": "Payment confirmed", "status": "confirmed"}

@api_router.get("/bookings", response_model=List[BookingResponse])
async def get_bookings(
    hotel_id: Optional[str] = None,
    booking_type: Optional[str] = None,
    checkin_status: Optional[str] = None,
    status: Optional[str] = None,
    date: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    
    if current_user["role"] == "cashier":
        query["hotel_id"] = current_user["assigned_hotel_id"]
    elif current_user["role"] == "owner":
        owner_hotels = await db.hotels.find({"owner_id": current_user["id"]}, {"id": 1, "_id": 0}).to_list(100)
        hotel_ids = [h["id"] for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            query["hotel_id"] = hotel_id
        else:
            query["hotel_id"] = {"$in": hotel_ids}
    elif hotel_id:
        query["hotel_id"] = hotel_id
    
    if booking_type:
        query["booking_type"] = booking_type
    if checkin_status:
        query["checkin_status"] = checkin_status
    if status:
        query["status"] = status
    if date:
        query["checkin_date"] = date
    
    bookings = await db.bookings.find(query, {"_id": 0}).sort("created_at", -1).to_list(500)
    
    result = []
    for b in bookings:
        hotel = await db.hotels.find_one({"id": b["hotel_id"]}, {"_id": 0})
        room_type = await db.room_types.find_one({"id": b["room_type_id"]}, {"_id": 0})
        creator = await db.users.find_one({"id": b.get("created_by")}, {"_id": 0}) if b.get("created_by") else None
        
        result.append(BookingResponse(
            id=b["id"],
            booking_ref=b["booking_ref"],
            hotel_id=b["hotel_id"],
            hotel_name=hotel["name"] if hotel else "Unknown",
            room_type_id=b["room_type_id"],
            room_type_name=room_type["name"] if room_type else "Unknown",
            guest_name=b["guest_name"],
            guest_phone=b["guest_phone"],
            guest_email=b.get("guest_email"),
            checkin_date=b["checkin_date"],
            checkout_date=b["checkout_date"],
            nights=b.get("nights", b.get("num_nights", 1)),
            num_guests=b["num_guests"],
            total_amount_tzs=b.get("total_amount_tzs", b.get("total_amount", 0)),
            booking_type=b["booking_type"],
            payment_method=b["payment_method"],
            payment_status=b["payment_status"],
            payment_reference=b.get("payment_reference"),
            status=b["status"],
            checkin_status=b["checkin_status"],
            actual_checkin_time=b.get("actual_checkin_time"),
            actual_checkout_time=b.get("actual_checkout_time"),
            created_by=b.get("created_by"),
            created_by_name=creator["full_name"] if creator else "Online System",
            notes=b.get("notes"),
            created_at=b["created_at"],
            expires_at=b.get("expires_at")
        ))
    return result

@api_router.get("/bookings/recent")
async def get_recent_bookings(hotel_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    """Get recent bookings for toast notifications (polling)"""
    query = {"status": "confirmed"}
    
    if current_user["role"] == "owner":
        owner_hotels = await db.hotels.find({"owner_id": current_user["id"]}, {"id": 1, "_id": 0}).to_list(100)
        hotel_ids = [h["id"] for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            query["hotel_id"] = hotel_id
        else:
            query["hotel_id"] = {"$in": hotel_ids}
    elif hotel_id:
        query["hotel_id"] = hotel_id
    
    # Last 5 minutes
    five_min_ago = (datetime.now(timezone.utc) - timedelta(minutes=5)).isoformat()
    query["created_at"] = {"$gte": five_min_ago}
    
    bookings = await db.bookings.find(query, {"_id": 0}).sort("created_at", -1).to_list(10)
    return bookings

@api_router.get("/bookings/{booking_id}", response_model=BookingResponse)
async def get_booking(booking_id: str):
    booking = await db.bookings.find_one({"id": booking_id}, {"_id": 0})
    if not booking:
        # Try by booking_ref
        booking = await db.bookings.find_one({"booking_ref": booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    hotel = await db.hotels.find_one({"id": booking["hotel_id"]}, {"_id": 0})
    room_type = await db.room_types.find_one({"id": booking["room_type_id"]}, {"_id": 0})
    creator = await db.users.find_one({"id": booking.get("created_by")}, {"_id": 0}) if booking.get("created_by") else None
    
    return BookingResponse(
        id=booking["id"],
        booking_ref=booking["booking_ref"],
        hotel_id=booking["hotel_id"],
        hotel_name=hotel["name"] if hotel else "Unknown",
        room_type_id=booking["room_type_id"],
        room_type_name=room_type["name"] if room_type else "Unknown",
        guest_name=booking["guest_name"],
        guest_phone=booking["guest_phone"],
        guest_email=booking.get("guest_email"),
        checkin_date=booking["checkin_date"],
        checkout_date=booking["checkout_date"],
        nights=booking.get("nights", booking.get("num_nights", 1)),
        num_guests=booking["num_guests"],
        total_amount_tzs=booking.get("total_amount_tzs", booking.get("total_amount", 0)),
        booking_type=booking["booking_type"],
        payment_method=booking["payment_method"],
        payment_status=booking["payment_status"],
        payment_reference=booking.get("payment_reference"),
        status=booking["status"],
        checkin_status=booking["checkin_status"],
        actual_checkin_time=booking.get("actual_checkin_time"),
        actual_checkout_time=booking.get("actual_checkout_time"),
        created_by=booking.get("created_by"),
        created_by_name=creator["full_name"] if creator else "Online System",
        notes=booking.get("notes"),
        created_at=booking["created_at"],
        expires_at=booking.get("expires_at")
    )

@api_router.post("/bookings/{booking_id}/checkin")
async def confirm_checkin(
    booking_id: str,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["cashier", "owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    booking = await db.bookings.find_one({"id": booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    if booking["checkin_status"] != "not_checked_in":
        raise HTTPException(status_code=400, detail="Mgeni tayari amefika au ameondoka")
    
    await db.bookings.update_one(
        {"id": booking_id},
        {"$set": {
            "checkin_status": "checked_in",
            "actual_checkin_time": datetime.now(timezone.utc).isoformat()
        }}
    )
    
    # Notify owner
    hotel = await db.hotels.find_one({"id": booking["hotel_id"]}, {"_id": 0})
    owner = await db.users.find_one({"id": hotel.get("owner_id")}, {"_id": 0})
    if owner:
        cashier_name = current_user["full_name"]
        sms_msg = f"HABARI STAYS: {booking['guest_name']} amefika. Ref: {booking['booking_ref']} imethibitishwa na {cashier_name}."
        background_tasks.add_task(send_sms, owner["phone"], sms_msg)
    
    return {"message": "Check-in imethibitishwa", "booking_id": booking_id}

@api_router.post("/bookings/{booking_id}/checkout")
async def confirm_checkout(booking_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["cashier", "owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    booking = await db.bookings.find_one({"id": booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    if booking["checkin_status"] != "checked_in":
        raise HTTPException(status_code=400, detail="Mgeni hajaingizwa au tayari ameondoka")
    
    await db.bookings.update_one(
        {"id": booking_id},
        {"$set": {
            "checkin_status": "checked_out",
            "actual_checkout_time": datetime.now(timezone.utc).isoformat()
        }}
    )
    
    # Increment available rooms
    await db.room_types.update_one(
        {"id": booking["room_type_id"]},
        {"$inc": {"available_rooms": 1}}
    )
    
    return {"message": "Check-out imethibitishwa", "booking_id": booking_id}

# ===================== REVIEW ENDPOINTS =====================

@api_router.post("/reviews", response_model=ReviewResponse)
async def create_review(review_data: ReviewCreate):
    booking = await db.bookings.find_one({"id": review_data.booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    review_id = str(uuid.uuid4())
    review = {
        "id": review_id,
        "booking_id": review_data.booking_id,
        "hotel_id": review_data.hotel_id,
        "guest_name": booking["guest_name"],
        "guest_phone": booking["guest_phone"],
        "rating": review_data.rating,
        "comment": review_data.comment,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.reviews.insert_one(review)
    
    return ReviewResponse(**{k: v for k, v in review.items() if k != "_id"})

@api_router.get("/reviews", response_model=List[ReviewResponse])
async def get_reviews(hotel_id: str):
    reviews = await db.reviews.find({"hotel_id": hotel_id}, {"_id": 0}).sort("created_at", -1).to_list(100)
    return [ReviewResponse(**r) for r in reviews]

# ===================== CASHIER MANAGEMENT =====================

@api_router.post("/cashiers", response_model=CashierResponse)
async def create_cashier(
    cashier_data: CashierCreate,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi kuunda cashier")
    
    hotel = await db.hotels.find_one({"id": cashier_data.hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] == "owner" and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    existing = await db.users.find_one({"email": cashier_data.email})
    if existing:
        raise HTTPException(status_code=400, detail="Email tayari imetumika")
    
    temp_password = generate_temp_password()
    user_id = str(uuid.uuid4())
    
    cashier = {
        "id": user_id,
        "email": cashier_data.email,
        "full_name": cashier_data.full_name,
        "phone": format_phone(cashier_data.phone),
        "role": "cashier",
        "assigned_hotel_id": cashier_data.hotel_id,
        "password_hash": get_password_hash(temp_password),
        "is_active": True,
        "is_verified": True,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "created_by": current_user["id"],
        "last_login": None
    }
    await db.users.insert_one(cashier)
    
    # Create assignment record
    assignment = {
        "id": str(uuid.uuid4()),
        "cashier_id": user_id,
        "hotel_id": cashier_data.hotel_id,
        "assigned_by": current_user["id"],
        "is_active": True,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.cashier_assignments.insert_one(assignment)
    
    # Send SMS with credentials
    sms_msg = f"Habari {cashier_data.full_name}! Umewekwa kama Mweka Hazina wa {hotel['name']} kwenye Habari Stays. Ingia: habaristays.com/login Email: {cashier_data.email} Neno la siri: {temp_password}"
    background_tasks.add_task(send_sms, cashier_data.phone, sms_msg)
    
    return CashierResponse(
        id=cashier["id"],
        full_name=cashier["full_name"],
        phone=cashier["phone"],
        email=cashier["email"],
        role=cashier["role"],
        assigned_hotel_id=cashier["assigned_hotel_id"],
        assigned_hotel_name=hotel["name"],
        is_active=cashier["is_active"],
        created_at=cashier["created_at"],
        last_login=cashier.get("last_login")
    )

@api_router.get("/cashiers", response_model=List[CashierResponse])
async def get_cashiers(hotel_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    query = {"role": "cashier"}
    
    if current_user["role"] == "owner":
        owner_hotels = await db.hotels.find({"owner_id": current_user["id"]}, {"id": 1, "_id": 0}).to_list(100)
        hotel_ids = [h["id"] for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            query["assigned_hotel_id"] = hotel_id
        else:
            query["assigned_hotel_id"] = {"$in": hotel_ids}
    elif hotel_id:
        query["assigned_hotel_id"] = hotel_id
    
    cashiers = await db.users.find(query, {"_id": 0, "password_hash": 0}).to_list(100)
    
    result = []
    for c in cashiers:
        hotel = await db.hotels.find_one({"id": c["assigned_hotel_id"]}, {"_id": 0})
        result.append(CashierResponse(
            id=c["id"],
            full_name=c["full_name"],
            phone=c["phone"],
            email=c["email"],
            role=c["role"],
            assigned_hotel_id=c["assigned_hotel_id"],
            assigned_hotel_name=hotel["name"] if hotel else "Unknown",
            is_active=c.get("is_active", True),
            created_at=c["created_at"],
            last_login=c.get("last_login")
        ))
    return result

@api_router.put("/cashiers/{cashier_id}/toggle-status")
async def toggle_cashier_status(cashier_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    cashier = await db.users.find_one({"id": cashier_id, "role": "cashier"}, {"_id": 0})
    if not cashier:
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    
    if current_user["role"] == "owner":
        hotel = await db.hotels.find_one({"id": cashier["assigned_hotel_id"]}, {"_id": 0})
        if hotel.get("owner_id") != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    new_status = not cashier.get("is_active", True)
    await db.users.update_one({"id": cashier_id}, {"$set": {"is_active": new_status}})
    await db.cashier_assignments.update_one(
        {"cashier_id": cashier_id, "is_active": True},
        {"$set": {"is_active": new_status}}
    )
    
    return {"message": "Hali imebadilishwa", "is_active": new_status}

@api_router.put("/cashiers/{cashier_id}/reassign")
async def reassign_cashier(cashier_id: str, hotel_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    cashier = await db.users.find_one({"id": cashier_id, "role": "cashier"}, {"_id": 0})
    if not cashier:
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    
    new_hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not new_hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] == "owner" and new_hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    await db.users.update_one({"id": cashier_id}, {"$set": {"assigned_hotel_id": hotel_id}})
    
    # Deactivate old assignment
    await db.cashier_assignments.update_many(
        {"cashier_id": cashier_id, "is_active": True},
        {"$set": {"is_active": False}}
    )
    
    # Create new assignment
    assignment = {
        "id": str(uuid.uuid4()),
        "cashier_id": cashier_id,
        "hotel_id": hotel_id,
        "assigned_by": current_user["id"],
        "is_active": True,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.cashier_assignments.insert_one(assignment)
    
    return {"message": "Cashier amehamishwa", "hotel_id": hotel_id}

# ===================== CASHIER SHIFT SUMMARY =====================

@api_router.get("/cashier/shift-summary", response_model=ShiftSummary)
async def get_shift_summary(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Hii ni kwa cashier tu")
    
    today = datetime.now().strftime("%Y-%m-%d")
    start_of_day = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0).isoformat()
    
    # Walk-ins today
    walkin_bookings = await db.bookings.find({
        "hotel_id": current_user["assigned_hotel_id"],
        "booking_type": "walkin",
        "created_by": current_user["id"],
        "created_at": {"$gte": start_of_day}
    }, {"_id": 0}).to_list(100)
    
    walkins_count = len(walkin_bookings)
    walkins_total = sum(b.get("total_amount_tzs", b.get("total_amount", 0)) for b in walkin_bookings)
    cash_collected = sum(b.get("total_amount_tzs", b.get("total_amount", 0)) for b in walkin_bookings if b["payment_method"] == "cash")
    
    # Online check-ins confirmed today
    online_checkins = await db.bookings.count_documents({
        "hotel_id": current_user["assigned_hotel_id"],
        "booking_type": "online",
        "checkin_status": "checked_in",
        "actual_checkin_time": {"$gte": start_of_day}
    })
    
    # Checkouts today
    checkouts = await db.bookings.count_documents({
        "hotel_id": current_user["assigned_hotel_id"],
        "checkin_status": "checked_out",
        "actual_checkout_time": {"$gte": start_of_day}
    })
    
    return ShiftSummary(
        cashier_id=current_user["id"],
        cashier_name=current_user["full_name"],
        date=today,
        walkins_count=walkins_count,
        walkins_total=walkins_total,
        online_checkins_count=online_checkins,
        checkouts_count=checkouts,
        cash_collected=cash_collected
    )

# ===================== CASHIER ACTIVITY LOG =====================

async def log_cashier_activity(cashier_id: str, hotel_id: str, booking_id: str, action_type: str,
                                payment_method: str = None, amount_tzs: int = 0,
                                guest_name: str = "", room_type_name: str = ""):
    entry = {
        "id": str(uuid.uuid4()),
        "cashier_id": cashier_id,
        "hotel_id": hotel_id,
        "booking_id": booking_id,
        "action_type": action_type,
        "payment_method": payment_method,
        "amount_tzs": amount_tzs,
        "guest_name": guest_name,
        "room_type_name": room_type_name,
        "timestamp": datetime.now(timezone.utc).isoformat()
    }
    await db.cashier_activity_log.insert_one(entry)

# ── Cashier: Today's Check-ins ──
@api_router.get("/cashier/todays-checkins")
async def cashier_todays_checkins(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    hotel_id = current_user["assigned_hotel_id"]
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    bookings = await db.bookings.find({
        "hotel_id": hotel_id,
        "checkin_date": today,
        "booking_type": "online",
        "status": {"$in": ["confirmed", "pending"]},
    }, {"_id": 0}).to_list(200)
    result = []
    for b in bookings:
        room = await db.room_types.find_one({"id": b["room_type_id"]}, {"_id": 0})
        result.append({
            "id": b["id"],
            "booking_ref": b["booking_ref"],
            "guest_name": b["guest_name"],
            "guest_phone": b["guest_phone"],
            "room_type_name": room["name"] if room else "Unknown",
            "nights": b.get("nights", 1),
            "total_amount_tzs": b.get("total_amount_tzs", 0),
            "payment_status": b.get("payment_status", "unpaid"),
            "checkin_status": b.get("checkin_status", "not_checked_in"),
            "checkin_date": b["checkin_date"],
            "checkout_date": b["checkout_date"],
            "actual_checkin_time": b.get("actual_checkin_time"),
            "actual_checkout_time": b.get("actual_checkout_time"),
            "status": b.get("status", "pending"),
        })
    return result

# ── Cashier: Active Guests ──
@api_router.get("/cashier/active-guests")
async def cashier_active_guests(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    hotel_id = current_user["assigned_hotel_id"]
    bookings = await db.bookings.find({
        "hotel_id": hotel_id,
        "checkin_status": "checked_in",
    }, {"_id": 0}).to_list(200)
    result = []
    for b in bookings:
        room = await db.room_types.find_one({"id": b["room_type_id"]}, {"_id": 0})
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        checkout_date = b.get("checkout_date", today)
        try:
            nights_left = max(0, (datetime.fromisoformat(checkout_date) - datetime.fromisoformat(today)).days)
        except Exception:
            nights_left = 0
        result.append({
            "id": b["id"],
            "booking_ref": b["booking_ref"],
            "guest_name": b["guest_name"],
            "guest_phone": b["guest_phone"],
            "room_type_name": room["name"] if room else "Unknown",
            "checkin_date": b["checkin_date"],
            "checkout_date": b["checkout_date"],
            "actual_checkin_time": b.get("actual_checkin_time"),
            "nights_left": nights_left,
            "booking_type": b["booking_type"],
            "total_amount_tzs": b.get("total_amount_tzs", 0),
            "payment_method": b.get("payment_method", ""),
        })
    return result

# ── Cashier: Confirm Check-in (with activity log) ──
@api_router.post("/cashier/confirm-checkin/{booking_id}")
async def cashier_confirm_checkin(
    booking_id: str,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    booking = await db.bookings.find_one({"id": booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    if booking["hotel_id"] != current_user["assigned_hotel_id"]:
        raise HTTPException(status_code=403, detail="Booking hii ni kwa hoteli nyingine")
    if booking.get("payment_status") != "paid":
        raise HTTPException(status_code=400, detail="Booking hii haijalipiwa")
    if booking["checkin_status"] != "not_checked_in":
        raise HTTPException(status_code=400, detail="Mgeni tayari amefika au ameondoka")
    now = datetime.now(timezone.utc).isoformat()
    await db.bookings.update_one({"id": booking_id}, {"$set": {
        "checkin_status": "checked_in",
        "actual_checkin_time": now,
        "status": "confirmed"
    }})
    room = await db.room_types.find_one({"id": booking["room_type_id"]}, {"_id": 0})
    await log_cashier_activity(
        current_user["id"], booking["hotel_id"], booking_id, "checkin_confirmed",
        guest_name=booking["guest_name"], room_type_name=room["name"] if room else ""
    )
    hotel = await db.hotels.find_one({"id": booking["hotel_id"]}, {"_id": 0})
    owner = await db.users.find_one({"id": hotel.get("owner_id")}, {"_id": 0}) if hotel else None
    if owner:
        sms_msg = f"HABARI STAYS: {booking['guest_name']} amewasili. Ref: {booking['booking_ref']}. Mweka Hazina: {current_user['full_name']}."
        background_tasks.add_task(send_sms, owner["phone"], sms_msg)
    return {"message": f"{booking['guest_name']} amethibitishwa", "booking_id": booking_id}

# ── Cashier: Confirm Checkout (with room restore + activity log) ──
@api_router.post("/cashier/confirm-checkout/{booking_id}")
async def cashier_confirm_checkout(booking_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    booking = await db.bookings.find_one({"id": booking_id}, {"_id": 0})
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    if booking["hotel_id"] != current_user["assigned_hotel_id"]:
        raise HTTPException(status_code=403, detail="Booking hii ni kwa hoteli nyingine")
    if booking["checkin_status"] != "checked_in":
        raise HTTPException(status_code=400, detail="Mgeni hajaingizwa au tayari ameondoka")
    now = datetime.now(timezone.utc).isoformat()
    await db.bookings.update_one({"id": booking_id}, {"$set": {
        "checkin_status": "checked_out",
        "actual_checkout_time": now
    }})
    await db.room_types.update_one({"id": booking["room_type_id"]}, {"$inc": {"available_rooms": 1}})
    room = await db.room_types.find_one({"id": booking["room_type_id"]}, {"_id": 0})
    await log_cashier_activity(
        current_user["id"], booking["hotel_id"], booking_id, "checkout_processed",
        guest_name=booking["guest_name"], room_type_name=room["name"] if room else ""
    )
    return {"message": f"{booking['guest_name']} ameondoka. Chumba kimerudishwa.", "booking_id": booking_id}

# ── Cashier: Walk-in Booking ──
@api_router.post("/cashier/walkin")
async def cashier_walkin(request: Request, background_tasks: BackgroundTasks, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    body = await request.json()
    hotel_id = current_user["assigned_hotel_id"]
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    room_type_id = body.get("room_type_id")
    room = await db.room_types.find_one({"id": room_type_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    if room.get("available_rooms", 0) < 1:
        raise HTTPException(status_code=400, detail="Hakuna vyumba vilivyopo")
    guest_name = body.get("guest_name", "").strip()
    guest_phone = format_phone(body.get("guest_phone", ""))
    checkin_date = body.get("checkin_date", datetime.now(timezone.utc).strftime("%Y-%m-%d"))
    checkout_date = body.get("checkout_date")
    if not guest_name or not guest_phone or not checkout_date:
        raise HTTPException(status_code=400, detail="Taarifa za mgeni zinahitajika")
    checkin_dt = datetime.fromisoformat(checkin_date)
    checkout_dt = datetime.fromisoformat(checkout_date)
    nights = (checkout_dt - checkin_dt).days
    if nights < 1:
        raise HTTPException(status_code=400, detail="Tarehe si sahihi")
    total_amount = nights * room["price_per_night"]
    payment_method = body.get("payment_method", "cash")
    notes = body.get("notes", "")
    booking_id = str(uuid.uuid4())
    booking_ref = generate_booking_ref()
    now = datetime.now(timezone.utc).isoformat()
    booking = {
        "id": booking_id,
        "booking_ref": booking_ref,
        "hotel_id": hotel_id,
        "room_type_id": room_type_id,
        "guest_name": guest_name,
        "guest_phone": guest_phone,
        "guest_email": body.get("guest_email"),
        "checkin_date": checkin_date,
        "checkout_date": checkout_date,
        "nights": nights,
        "num_guests": body.get("num_guests", 1),
        "total_amount_tzs": total_amount,
        "booking_type": "walkin",
        "payment_method": payment_method,
        "payment_status": "paid",
        "payment_reference": None,
        "status": "confirmed",
        "checkin_status": "checked_in",
        "actual_checkin_time": now,
        "actual_checkout_time": None,
        "created_by": current_user["id"],
        "notes": notes,
        "created_at": now,
        "expires_at": None
    }
    await db.bookings.insert_one(booking)
    await db.room_types.update_one({"id": room_type_id}, {"$inc": {"available_rooms": -1}})
    await log_cashier_activity(
        current_user["id"], hotel_id, booking_id, "walkin_recorded",
        payment_method=payment_method, amount_tzs=total_amount,
        guest_name=guest_name, room_type_name=room["name"]
    )
    owner = await db.users.find_one({"id": hotel.get("owner_id")}, {"_id": 0})
    if owner:
        sms_msg = f"HABARI STAYS: Walk-in Mpya! Mgeni: {guest_name}, Chumba: {room['name']}, Usiku: {nights}, TZS {total_amount:,}. Mweka Hazina: {current_user['full_name']}. Ref: {booking_ref}"
        background_tasks.add_task(send_sms, owner["phone"], sms_msg)
    return {
        "id": booking_id, "booking_ref": booking_ref, "hotel_name": hotel["name"],
        "hotel_address": hotel.get("address", ""), "guest_name": guest_name,
        "guest_phone": guest_phone, "room_type_name": room["name"],
        "checkin_date": checkin_date, "checkout_date": checkout_date,
        "nights": nights, "total_amount_tzs": total_amount,
        "payment_method": payment_method, "cashier_name": current_user["full_name"],
        "created_at": now
    }

# ── Cashier: Activity Log (today) ──
@api_router.get("/cashier/activity-log")
async def cashier_activity_log_today(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    start_of_day = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0).isoformat()
    logs = await db.cashier_activity_log.find({
        "cashier_id": current_user["id"],
        "timestamp": {"$gte": start_of_day}
    }, {"_id": 0}).sort("timestamp", -1).to_list(200)
    return logs

# ── Cashier: Enhanced Shift Summary ──
@api_router.get("/cashier/shift-summary-full")
async def cashier_shift_summary_full(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    hotel_id = current_user["assigned_hotel_id"]
    start_of_day = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0).isoformat()
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    walkin_bookings = await db.bookings.find({
        "hotel_id": hotel_id, "booking_type": "walkin",
        "created_by": current_user["id"], "created_at": {"$gte": start_of_day}
    }, {"_id": 0}).to_list(200)
    walkins_count = len(walkin_bookings)
    cash_total = sum(b.get("total_amount_tzs", 0) for b in walkin_bookings if b.get("payment_method") == "cash")
    mpesa_total = sum(b.get("total_amount_tzs", 0) for b in walkin_bookings if b.get("payment_method") == "mpesa")
    card_total = sum(b.get("total_amount_tzs", 0) for b in walkin_bookings if b.get("payment_method") == "card")
    walkins_total = cash_total + mpesa_total + card_total
    online_checkins = await db.bookings.count_documents({
        "hotel_id": hotel_id, "booking_type": "online",
        "checkin_status": {"$in": ["checked_in", "checked_out"]},
        "actual_checkin_time": {"$gte": start_of_day}
    })
    checkouts = await db.bookings.count_documents({
        "hotel_id": hotel_id, "checkin_status": "checked_out",
        "actual_checkout_time": {"$gte": start_of_day}
    })
    activity_logs = await db.cashier_activity_log.find({
        "cashier_id": current_user["id"], "timestamp": {"$gte": start_of_day}
    }, {"_id": 0}).sort("timestamp", -1).to_list(200)
    return {
        "cashier_name": current_user["full_name"],
        "hotel_name": hotel["name"] if hotel else "",
        "hotel_address": hotel.get("address", "") if hotel else "",
        "date": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        "walkins_count": walkins_count,
        "walkins_total": walkins_total,
        "online_checkins_count": online_checkins,
        "checkouts_count": checkouts,
        "cash_total": cash_total,
        "mpesa_total": mpesa_total,
        "card_total": card_total,
        "total_revenue": walkins_total,
        "activity_log": activity_logs
    }

# ── Cashier: Verify/Lookup Booking ──
@api_router.post("/cashier/verify-booking")
async def cashier_verify_booking(request: Request, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    body = await request.json()
    search_term = body.get("search", "").strip()
    if not search_term:
        raise HTTPException(status_code=400, detail="Tafadhali ingiza namba ya booking, simu, au jina")
    hotel_id = current_user["assigned_hotel_id"]
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    booking = None
    if search_term.upper().startswith("HS-"):
        booking = await db.bookings.find_one({"booking_ref": search_term.upper()}, {"_id": 0})
    if not booking:
        booking = await db.bookings.find_one({"guest_phone": {"$regex": search_term.replace("+", "\\+"), "$options": "i"}}, {"_id": 0})
    if not booking:
        bookings = await db.bookings.find({"guest_name": {"$regex": search_term, "$options": "i"}}, {"_id": 0}).sort("created_at", -1).to_list(10)
        booking = bookings[0] if bookings else None
    if not booking:
        return {"found": False, "error": "not_found", "message": "Booking haikupatikana. Angalia namba na ujaribu tena."}
    room = await db.room_types.find_one({"id": booking["room_type_id"]}, {"_id": 0})
    hotel = await db.hotels.find_one({"id": booking["hotel_id"]}, {"_id": 0})
    result = {
        "found": True,
        "id": booking["id"],
        "booking_ref": booking["booking_ref"],
        "guest_name": booking["guest_name"],
        "guest_phone": booking["guest_phone"],
        "room_type_name": room["name"] if room else "Unknown",
        "checkin_date": booking["checkin_date"],
        "checkout_date": booking["checkout_date"],
        "nights": booking.get("nights", 1),
        "total_amount_tzs": booking.get("total_amount_tzs", 0),
        "payment_status": booking.get("payment_status", "unpaid"),
        "payment_method": booking.get("payment_method", ""),
        "booking_type": booking["booking_type"],
        "status": booking.get("status", "pending"),
        "checkin_status": booking.get("checkin_status", "not_checked_in"),
        "actual_checkin_time": booking.get("actual_checkin_time"),
        "hotel_name": hotel["name"] if hotel else "Unknown",
        "hotel_id": booking["hotel_id"],
        "can_checkin": False,
        "errors": []
    }
    if booking["hotel_id"] != hotel_id:
        result["errors"].append({"type": "wrong_hotel", "message": f"Booking hii ni kwa hoteli nyingine: {hotel['name'] if hotel else 'Unknown'}."})
    elif booking.get("status") == "expired":
        result["errors"].append({"type": "expired", "message": "Booking hii imeisha muda na haijalipiwa."})
    elif booking.get("payment_status") != "paid":
        result["errors"].append({"type": "unpaid", "message": "Booking hii haijalipiwa. Mwambie mgeni alipe kwanza."})
    elif booking.get("checkin_status") == "checked_in":
        checkin_time = booking.get("actual_checkin_time", "")
        result["errors"].append({"type": "already_checked_in", "message": f"Mgeni huyu tayari amethibitishwa saa {checkin_time[:16]}."})
    elif booking.get("checkin_status") == "checked_out":
        result["errors"].append({"type": "checked_out", "message": "Mgeni huyu tayari ameondoka."})
    elif booking.get("checkin_date") != today:
        result["errors"].append({"type": "wrong_date", "message": f"Check-in ya booking hii ni {booking['checkin_date']}, si leo."})
    else:
        result["can_checkin"] = True
    return result

# ── Admin: Reset Cashier Password ──
@api_router.post("/cashiers/{cashier_id}/reset-password")
async def reset_cashier_password(cashier_id: str, background_tasks: BackgroundTasks, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    cashier = await db.users.find_one({"id": cashier_id, "role": "cashier"}, {"_id": 0})
    if not cashier:
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    if current_user["role"] == "owner":
        hotel = await db.hotels.find_one({"id": cashier["assigned_hotel_id"]}, {"_id": 0})
        if not hotel or hotel.get("owner_id") != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    new_password = generate_temp_password()
    await db.users.update_one({"id": cashier_id}, {"$set": {"password_hash": get_password_hash(new_password)}})
    sms_msg = f"Habari {cashier['full_name']}! Neno lako la siri la Habari Stays limewekwa upya. Neno jipya: {new_password} Ingia: habaristays.com/login Badilisha neno la siri haraka."
    background_tasks.add_task(send_sms, cashier["phone"], sms_msg)
    return {"message": "Neno la siri limewekwa upya", "new_password": new_password}

# ── Admin: Edit Cashier ──
@api_router.patch("/cashiers/{cashier_id}")
async def edit_cashier(cashier_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    cashier = await db.users.find_one({"id": cashier_id, "role": "cashier"}, {"_id": 0})
    if not cashier:
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    if current_user["role"] == "owner":
        hotel = await db.hotels.find_one({"id": cashier["assigned_hotel_id"]}, {"_id": 0})
        if not hotel or hotel.get("owner_id") != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    body = await request.json()
    update = {}
    if "full_name" in body:
        update["full_name"] = body["full_name"]
    if "phone" in body:
        update["phone"] = format_phone(body["phone"])
    if "assigned_hotel_id" in body:
        new_hotel = await db.hotels.find_one({"id": body["assigned_hotel_id"]}, {"_id": 0})
        if not new_hotel:
            raise HTTPException(status_code=404, detail="Hotel haipatikani")
        if current_user["role"] == "owner" and new_hotel.get("owner_id") != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hii si hotel yako")
        update["assigned_hotel_id"] = body["assigned_hotel_id"]
        await db.cashier_assignments.update_many({"cashier_id": cashier_id, "is_active": True}, {"$set": {"is_active": False}})
        await db.cashier_assignments.insert_one({
            "id": str(uuid.uuid4()), "cashier_id": cashier_id,
            "hotel_id": body["assigned_hotel_id"], "assigned_by": current_user["id"],
            "is_active": True, "created_at": datetime.now(timezone.utc).isoformat()
        })
    if not update:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    await db.users.update_one({"id": cashier_id}, {"$set": update})
    return {"message": "Cashier amesasishwa"}

# ── Admin: Toggle Cashier with SMS ──
@api_router.put("/cashiers/{cashier_id}/toggle-status-sms")
async def toggle_cashier_status_sms(cashier_id: str, background_tasks: BackgroundTasks, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    cashier = await db.users.find_one({"id": cashier_id, "role": "cashier"}, {"_id": 0})
    if not cashier:
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    if current_user["role"] == "owner":
        hotel = await db.hotels.find_one({"id": cashier["assigned_hotel_id"]}, {"_id": 0})
        if not hotel or hotel.get("owner_id") != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    new_status = not cashier.get("is_active", True)
    await db.users.update_one({"id": cashier_id}, {"$set": {"is_active": new_status}})
    await db.cashier_assignments.update_one({"cashier_id": cashier_id, "is_active": True}, {"$set": {"is_active": new_status}})
    if new_status:
        sms_msg = f"Habari {cashier['full_name']}! Akaunti yako ya Habari Stays imewashwa tena. Unaweza kuingia sasa."
    else:
        sms_msg = f"Habari {cashier['full_name']}. Akaunti yako ya Habari Stays imezimwa. Wasiliana na msimamizi wako."
    background_tasks.add_task(send_sms, cashier["phone"], sms_msg)
    return {"message": "Hali imebadilishwa", "is_active": new_status}

# ── Admin: Cashier Performance Detail ──
@api_router.get("/cashiers/{cashier_id}/performance")
async def cashier_performance_detail(cashier_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    cashier = await db.users.find_one({"id": cashier_id, "role": "cashier"}, {"_id": 0, "password_hash": 0})
    if not cashier:
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    hotel = await db.hotels.find_one({"id": cashier["assigned_hotel_id"]}, {"_id": 0})
    start_of_month = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat()
    logs = await db.cashier_activity_log.find({
        "cashier_id": cashier_id, "timestamp": {"$gte": start_of_month}
    }, {"_id": 0}).sort("timestamp", -1).to_list(500)
    walkins = [l for l in logs if l["action_type"] == "walkin_recorded"]
    checkins = [l for l in logs if l["action_type"] == "checkin_confirmed"]
    checkouts = [l for l in logs if l["action_type"] == "checkout_processed"]
    return {
        "cashier": {**cashier, "assigned_hotel_name": hotel["name"] if hotel else "Unknown"},
        "this_month": {
            "walkins_recorded": len(walkins),
            "checkins_confirmed": len(checkins),
            "checkouts_processed": len(checkouts),
            "total_revenue": sum(l.get("amount_tzs", 0) for l in walkins),
        },
        "activity_log": logs[:100]
    }

# ===================== ANALYTICS ENDPOINTS =====================

@api_router.get("/analytics/revenue")
async def get_revenue_analytics(
    hotel_id: Optional[str] = None,
    period: str = "month",
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    now = datetime.now()
    if period == "week":
        start_date = (now - timedelta(days=7)).isoformat()
    elif period == "month":
        start_date = now.replace(day=1).isoformat()
    else:
        start_date = now.replace(month=1, day=1).isoformat()
    
    query = {"created_at": {"$gte": start_date}, "payment_status": "paid"}
    
    if current_user["role"] == "owner":
        owner_hotels = await db.hotels.find({"owner_id": current_user["id"]}, {"id": 1, "_id": 0}).to_list(100)
        hotel_ids = [h["id"] for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            query["hotel_id"] = hotel_id
        else:
            query["hotel_id"] = {"$in": hotel_ids}
    elif hotel_id:
        query["hotel_id"] = hotel_id
    
    bookings = await db.bookings.find(query, {"_id": 0}).to_list(1000)
    
    total_revenue = sum(b.get("total_amount_tzs", b.get("total_amount", 0)) for b in bookings)
    online_revenue = sum(b.get("total_amount_tzs", b.get("total_amount", 0)) for b in bookings if b["booking_type"] == "online")
    walkin_revenue = sum(b.get("total_amount_tzs", b.get("total_amount", 0)) for b in bookings if b["booking_type"] == "walkin")
    
    daily_revenue = {}
    for b in bookings:
        date = b["created_at"][:10]
        if date not in daily_revenue:
            daily_revenue[date] = {"date": date, "online": 0, "walkin": 0, "total": 0}
        amount = b.get("total_amount_tzs", b.get("total_amount", 0))
        daily_revenue[date]["total"] += amount
        if b["booking_type"] == "online":
            daily_revenue[date]["online"] += amount
        else:
            daily_revenue[date]["walkin"] += amount
    
    return {
        "total_revenue": total_revenue,
        "online_revenue": online_revenue,
        "walkin_revenue": walkin_revenue,
        "bookings_count": len(bookings),
        "daily_breakdown": sorted(daily_revenue.values(), key=lambda x: x["date"])
    }

@api_router.get("/analytics/cashier-performance")
async def get_cashier_performance(
    hotel_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    start_of_day = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0).isoformat()
    start_of_month = datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat()
    
    cashier_query = {"role": "cashier"}
    if current_user["role"] == "owner":
        owner_hotels = await db.hotels.find({"owner_id": current_user["id"]}, {"id": 1, "_id": 0}).to_list(100)
        hotel_ids = [h["id"] for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            cashier_query["assigned_hotel_id"] = hotel_id
        else:
            cashier_query["assigned_hotel_id"] = {"$in": hotel_ids}
    elif hotel_id:
        cashier_query["assigned_hotel_id"] = hotel_id
    
    cashiers = await db.users.find(cashier_query, {"_id": 0, "password_hash": 0}).to_list(100)
    
    performance = []
    for cashier in cashiers:
        walkins_today = await db.bookings.find({
            "created_by": cashier["id"],
            "booking_type": "walkin",
            "created_at": {"$gte": start_of_day}
        }, {"_id": 0}).to_list(100)
        
        walkins_month = await db.bookings.find({
            "created_by": cashier["id"],
            "booking_type": "walkin",
            "created_at": {"$gte": start_of_month}
        }, {"_id": 0}).to_list(500)
        
        cash_today = sum(b.get("total_amount_tzs", b.get("total_amount", 0)) for b in walkins_today if b["payment_method"] == "cash")
        
        performance.append({
            "cashier_id": cashier["id"],
            "cashier_name": cashier["full_name"],
            "walkins_today": len(walkins_today),
            "walkins_month": len(walkins_month),
            "cash_collected_today": cash_today,
            "last_active": cashier.get("last_login")
        })
    
    return performance

@api_router.get("/analytics/occupancy")
async def get_occupancy(hotel_id: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    """Calculate occupancy rate"""
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    query = {}
    if current_user["role"] == "owner":
        owner_hotels = await db.hotels.find({"owner_id": current_user["id"]}, {"id": 1, "_id": 0}).to_list(100)
        hotel_ids = [h["id"] for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            query["hotel_id"] = hotel_id
        else:
            query["hotel_id"] = {"$in": hotel_ids}
    elif hotel_id:
        query["hotel_id"] = hotel_id
    
    rooms = await db.room_types.find(query if query else {}, {"_id": 0}).to_list(100)
    total = sum(r["total_rooms"] for r in rooms)
    available = sum(r["available_rooms"] for r in rooms)
    occupied = total - available
    
    return {
        "total_rooms": total,
        "occupied_rooms": occupied,
        "available_rooms": available,
        "occupancy_rate": round((occupied / total * 100) if total > 0 else 0, 1)
    }

# ===================== ADMIN ENDPOINTS =====================

@api_router.get("/admin/stats")
async def get_admin_stats(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    total_hotels = await db.hotels.count_documents({})
    verified_hotels = await db.hotels.count_documents({"status": "verified"})
    pending_hotels = await db.hotels.count_documents({"status": "pending"})
    total_bookings = await db.bookings.count_documents({})
    online_bookings = await db.bookings.count_documents({"booking_type": "online"})
    walkin_bookings = await db.bookings.count_documents({"booking_type": "walkin"})
    total_cashiers = await db.users.count_documents({"role": "cashier"})
    total_owners = await db.users.count_documents({"role": "owner"})
    pending_owners = await db.users.count_documents({"role": "owner", "is_verified": False})
    
    revenue_pipeline = [
        {"$match": {"payment_status": "paid"}},
        {"$group": {"_id": None, "total": {"$sum": "$total_amount_tzs"}}}
    ]
    revenue_result = await db.bookings.aggregate(revenue_pipeline).to_list(1)
    total_revenue = revenue_result[0]["total"] if revenue_result else 0
    
    # Revenue by city
    city_revenue = await db.bookings.aggregate([
        {"$match": {"payment_status": "paid"}},
        {"$lookup": {"from": "hotels", "localField": "hotel_id", "foreignField": "id", "as": "hotel"}},
        {"$unwind": "$hotel"},
        {"$group": {"_id": "$hotel.city", "revenue": {"$sum": "$total_amount_tzs"}, "count": {"$sum": 1}}},
        {"$sort": {"revenue": -1}}
    ]).to_list(10)
    
    return {
        "total_hotels": total_hotels,
        "verified_hotels": verified_hotels,
        "pending_hotels": pending_hotels,
        "total_bookings": total_bookings,
        "online_bookings": online_bookings,
        "walkin_bookings": walkin_bookings,
        "total_cashiers": total_cashiers,
        "total_owners": total_owners,
        "pending_owners": pending_owners,
        "total_revenue": total_revenue,
        "commission_10_percent": int(total_revenue * 0.1),
        "revenue_by_city": [{"city": c["_id"], "revenue": c["revenue"], "bookings": c["count"]} for c in city_revenue]
    }

@api_router.put("/admin/hotels/{hotel_id}/verify")
async def verify_hotel(hotel_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    await db.hotels.update_one({"id": hotel_id}, {"$set": {"status": "verified"}})
    return {"message": "Hotel imethibitishwa"}

@api_router.put("/admin/hotels/{hotel_id}/suspend")
async def suspend_hotel(hotel_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    await db.hotels.update_one({"id": hotel_id}, {"$set": {"status": "suspended"}})
    return {"message": "Hotel imesimamishwa"}

@api_router.put("/admin/users/{user_id}/verify")
async def verify_owner(user_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    user = await db.users.find_one({"id": user_id}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=404, detail="User hapatikani")
    
    await db.users.update_one({"id": user_id}, {"$set": {"is_verified": True}})
    return {"message": "Mmiliki amethibitishwa"}

@api_router.get("/admin/pending-owners")
async def get_pending_owners(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    owners = await db.users.find(
        {"role": "owner", "is_verified": False},
        {"_id": 0, "password_hash": 0}
    ).to_list(100)
    return owners

@api_router.get("/admin/all-bookings")
async def get_all_bookings(
    hotel_id: Optional[str] = None,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    query = {}
    if hotel_id:
        query["hotel_id"] = hotel_id
    if status:
        query["status"] = status
    if date_from:
        query["created_at"] = {"$gte": date_from}
    if date_to:
        query.setdefault("created_at", {})["$lte"] = date_to
    
    bookings = await db.bookings.find(query, {"_id": 0}).sort("created_at", -1).to_list(500)
    return bookings

# ===================== ADMIN BULK IMPORT =====================

UPLOAD_DIR = Path("/app/backend/uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

def parse_cloudinary_photos_from_row(row, headers):
    """Parse P1/P2/P3 photo columns from an Excel row into structured photo objects."""
    photos = []
    header_lower = [str(h or "").strip().lower() for h in headers]
    
    for pn in [1, 2, 3]:
        prefix = f"p{pn}"
        local_path = original = thumb = mobile = web = hd = None
        
        for i, h in enumerate(header_lower):
            val = str(row[i] or "").strip() if i < len(row) and row[i] else ""
            if not val:
                continue
            if h.startswith(f"p{pn}") and "local" in h:
                local_path = val
            elif h.startswith(f"p{pn}") and "original" in h:
                original = val
            elif h.startswith(f"p{pn}") and "thumb" in h:
                thumb = val
            elif h.startswith(f"p{pn}") and "mobile" in h:
                mobile = val
            elif h.startswith(f"p{pn}") and ("web" in h or "1280" in h):
                web = val
            elif h.startswith(f"p{pn}") and ("hd" in h or "1920" in h):
                hd = val
        
        if original or thumb or mobile or web or hd:
            photos.append({
                "id": str(uuid.uuid4()),
                "is_primary": pn == 1,
                "sort_order": pn - 1,
                "local_path": local_path or "",
                "cloudinary_original": original or "",
                "cloudinary_thumb": thumb or "",
                "cloudinary_mobile": mobile or "",
                "cloudinary_web": web or "",
                "cloudinary_hd": hd or "",
                "source": "imported",
                "uploaded_at": datetime.now(timezone.utc).isoformat()
            })
    
    return photos

def count_valid_photos(photos):
    """Count photos that have at least one valid Cloudinary URL."""
    count = 0
    for p in photos:
        if isinstance(p, dict):
            has_url = any(
                p.get(k, "").startswith("https://res.cloudinary.com/")
                for k in ["cloudinary_original", "cloudinary_thumb", "cloudinary_mobile", "cloudinary_web", "cloudinary_hd"]
            )
            if has_url:
                count += 1
    return count

def extract_city_from_address(address):
    """Extract city from address - take second-to-last segment before Tanzania."""
    if not address:
        return ""
    parts = [p.strip() for p in address.split(",")]
    tz_cities = ["Dodoma", "Dar es Salaam", "Arusha", "Zanzibar", "Mwanza", "Mbeya", 
                 "Morogoro", "Tanga", "Iringa", "Moshi", "Musoma", "Bukoba", "Kigoma",
                 "Songea", "Lindi", "Mtwara", "Tabora", "Singida", "Shinyanga", "Kagera"]
    for c in tz_cities:
        if c.lower() in address.lower():
            return c
    # Try second-to-last segment
    if len(parts) >= 2:
        candidate = parts[-2].strip() if parts[-1].strip().lower() == "tanzania" else parts[-1].strip()
        if candidate and candidate.lower() != "tanzania":
            return candidate
    return ""

# Store latest preview for confirm step
_import_previews = {}

@api_router.post("/admin/import/preview")
async def preview_import(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user)
):
    """Parse Excel file with Cloudinary photo URLs and return preview data"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Excel file (.xlsx) required")
    
    try:
        import openpyxl
        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents))
        ws = wb.active
        
        # Auto-detect header row: find the row containing "Name" column
        header_row = 1
        data_start_row = 2
        for r in range(1, min(6, ws.max_row + 1)):
            row_vals = [str(ws.cell(row=r, column=c+1).value or "").strip().lower() for c in range(min(15, ws.max_column))]
            if "name" in row_vals:
                header_row = r
                data_start_row = r + 1
                break
        
        headers_raw = [ws.cell(row=header_row, column=c+1).value for c in range(ws.max_column)]
        headers = [str(h or "").strip().lower() for h in headers_raw]
        
        # Map base columns
        col_map = {}
        for i, h in enumerate(headers):
            if h in ['name', 'hotel name', 'jina']:
                col_map['name'] = i
            elif h in ['address', 'anwani']:
                col_map['address'] = i
            elif h in ['phone', 'simu', 'phone number']:
                col_map['phone'] = i
            elif h in ['website']:
                col_map['website'] = i
            elif h in ['rating', 'google rating']:
                col_map['rating'] = i
            elif h in ['# reviews', 'reviews', 'review count']:
                col_map['review_count'] = i
            elif h in ['price level']:
                col_map['price_level'] = i
            elif h in ['status', 'business status']:
                col_map['business_status'] = i
            elif h in ['latitude', 'lat']:
                col_map['latitude'] = i
            elif h in ['longitude', 'lng', 'long']:
                col_map['longitude'] = i
            elif h in ['types']:
                col_map['types'] = i
            elif h in ['summary']:
                col_map['summary'] = i
            elif h in ['google maps url', 'maps url', 'google_maps_url']:
                col_map['google_maps_url'] = i
            elif h in ['city', 'mji']:
                col_map['city'] = i
            elif h in ['#', 'no', 'number']:
                col_map['number'] = i
        
        preview = []
        existing_names = set()
        
        existing_hotels = await db.hotels.find({}, {"name": 1, "_id": 0}).to_list(1000)
        for eh in existing_hotels:
            existing_names.add(eh["name"].strip().lower())
        
        seen_names = set()
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=data_start_row, values_only=True), start=data_start_row):
            if not row or not any(row):
                continue
            
            name = str(row[col_map['name']] or "").strip() if col_map.get('name') is not None else ""
            if not name:
                continue
            
            address = str(row[col_map['address']] or "").strip() if col_map.get('address') is not None else ""
            phone = str(row[col_map['phone']] or "").strip() if col_map.get('phone') is not None else ""
            website = str(row[col_map['website']] or "").strip() if col_map.get('website') is not None else ""
            
            rating_val = row[col_map['rating']] if col_map.get('rating') is not None else None
            rating = float(rating_val) if rating_val else None
            
            review_val = row[col_map['review_count']] if col_map.get('review_count') is not None else None
            review_count = int(float(review_val)) if review_val else None
            
            business_status = str(row[col_map['business_status']] or "").strip() if col_map.get('business_status') is not None else ""
            
            lat_val = row[col_map['latitude']] if col_map.get('latitude') is not None else None
            latitude = float(lat_val) if lat_val else None
            
            lng_val = row[col_map['longitude']] if col_map.get('longitude') is not None else None
            longitude = float(lng_val) if lng_val else None
            
            maps_url = str(row[col_map['google_maps_url']] or "").strip() if col_map.get('google_maps_url') is not None else ""
            summary = str(row[col_map['summary']] or "").strip() if col_map.get('summary') is not None else ""
            types = str(row[col_map['types']] or "").strip() if col_map.get('types') is not None else ""
            
            # Parse photos from P1/P2/P3 columns
            photos = parse_cloudinary_photos_from_row(row, headers_raw)
            photo_count = count_valid_photos(photos)
            
            # Extract or infer city
            city = str(row[col_map['city']] or "").strip() if col_map.get('city') is not None else ""
            if not city:
                city = extract_city_from_address(address)
            
            # Determine row status
            name_lower = name.lower()
            is_duplicate = name_lower in existing_names or name_lower in seen_names
            has_missing_critical = not name or not address
            has_missing_optional = not phone or photo_count == 0
            
            if is_duplicate:
                row_status = "duplicate"
            elif has_missing_critical:
                row_status = "error"
            elif has_missing_optional:
                row_status = "warning"
            else:
                row_status = "valid"
            
            seen_names.add(name_lower)
            
            issues = []
            if is_duplicate:
                issues.append("Duplicate hotel name")
            if not name:
                issues.append("Missing name")
            if not address:
                issues.append("Missing address")
            if not phone:
                issues.append("Missing phone")
            if not city:
                issues.append("City not detected")
            if photo_count == 0:
                issues.append("No photos found")
            
            preview.append({
                "row": row_idx,
                "name": name,
                "address": address,
                "phone": phone,
                "website": website,
                "rating": rating,
                "review_count": review_count,
                "business_status": business_status,
                "latitude": latitude,
                "longitude": longitude,
                "google_maps_url": maps_url,
                "summary": summary,
                "types": types,
                "city": city,
                "photos": photos,
                "photo_count": photo_count,
                "status": row_status,
                "issues": issues
            })
        
        # Store preview for confirm step
        preview_id = str(uuid.uuid4())
        _import_previews[preview_id] = preview
        
        return {
            "preview_id": preview_id,
            "total_rows": len(preview),
            "valid": sum(1 for p in preview if p["status"] == "valid"),
            "warnings": sum(1 for p in preview if p["status"] == "warning"),
            "duplicates": sum(1 for p in preview if p["status"] == "duplicate"),
            "errors": sum(1 for p in preview if p["status"] == "error"),
            "columns_found": list(col_map.keys()),
            "has_photo_columns": any("p1" in h.lower() or "p2" in h.lower() for h in headers),
            "preview": preview
        }
    except Exception as e:
        logging.error(f"Import preview error: {str(e)}")
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel: {str(e)}")

@api_router.post("/admin/import/execute")
async def execute_import(
    request: Request,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user)
):
    """Execute the import from previously previewed data"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    body = await request.json()
    preview_id = body.get("preview_id")
    batch_name = body.get("batch_name", "")
    city_override = body.get("city_override", "")
    skip_duplicates = body.get("skip_duplicates", True)
    hotels_data = body.get("hotels") or (_import_previews.get(preview_id) if preview_id else None)
    
    if not hotels_data:
        raise HTTPException(status_code=400, detail="No preview data found")
    
    batch_id = str(uuid.uuid4())
    batch_name = batch_name or f"Import {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')}"
    
    imported = 0
    skipped = 0
    errors_list = []
    
    for hotel_data in hotels_data:
        if hotel_data.get("status") == "duplicate" and skip_duplicates:
            skipped += 1
            continue
        if hotel_data.get("status") == "error":
            errors_list.append({"name": hotel_data.get("name", ""), "reason": "Missing critical data"})
            skipped += 1
            continue
        
        name = hotel_data.get("name", "").strip()
        if not name:
            skipped += 1
            continue
        
        existing = await db.hotels.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}})
        if existing:
            skipped += 1
            continue
        
        city = city_override or hotel_data.get("city", "")
        if not city:
            city = extract_city_from_address(hotel_data.get("address", ""))
        if not city:
            city = "Tanzania"
        
        phone = hotel_data.get("phone", "")
        photos = hotel_data.get("photos", [])
        cover = get_cover_url(photos, "cloudinary_web")
        summary = hotel_data.get("summary", "")
        
        hotel_id = str(uuid.uuid4())
        hotel = {
            "id": hotel_id,
            "hotel_code": generate_hotel_code(),
            "owner_id": None,
            "name": name,
            "description": summary or f"Hotel in {city}",
            "address": hotel_data.get("address", ""),
            "city": city,
            "phone_number": phone,
            "whatsapp_number": phone,
            "amenities": [],
            "cover_photo": cover,
            "photos": photos,
            "google_maps_url": hotel_data.get("google_maps_url", ""),
            "google_rating": hotel_data.get("rating"),
            "google_review_count": hotel_data.get("review_count"),
            "latitude": hotel_data.get("latitude"),
            "longitude": hotel_data.get("longitude"),
            "website": hotel_data.get("website", ""),
            "status": "imported",
            "data_source": "bulk_import",
            "import_batch_id": batch_id,
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.hotels.insert_one(hotel)
        await create_default_room_type(hotel_id)
        imported += 1
    
    # Save batch record
    batch = {
        "id": batch_id,
        "name": batch_name,
        "city_override": city_override,
        "total": len(hotels_data),
        "imported": imported,
        "skipped": skipped,
        "errors": errors_list,
        "created_by": current_user["id"],
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.import_batches.insert_one(batch)
    
    # Clean up stored preview
    if preview_id and preview_id in _import_previews:
        del _import_previews[preview_id]
    
    return {
        "batch_id": batch_id,
        "batch_name": batch_name,
        "imported": imported,
        "skipped": skipped,
        "errors": errors_list,
        "message": f"Hoteli {imported} zimeingizwa, {skipped} zimepitishwa"
    }

@api_router.get("/admin/import/batches")
async def get_import_batches(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    batches = await db.import_batches.find({}, {"_id": 0}).sort("created_at", -1).to_list(50)
    return batches

@api_router.get("/admin/import/batches/{batch_id}")
async def get_import_batch(batch_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    batch = await db.import_batches.find_one({"id": batch_id}, {"_id": 0})
    if not batch:
        raise HTTPException(status_code=404, detail="Batch haipatikani")
    hotels = await db.hotels.find({"import_batch_id": batch_id}, {"_id": 0}).to_list(500)
    batch["hotels"] = hotels
    return batch

@api_router.get("/admin/import/batches/{batch_id}/errors")
async def get_import_batch_errors(batch_id: str, current_user: dict = Depends(get_current_user)):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    batch = await db.import_batches.find_one({"id": batch_id}, {"_id": 0})
    if not batch:
        raise HTTPException(status_code=404, detail="Batch haipatikani")
    return batch.get("errors", [])

# Keep old confirm endpoint for backward compatibility
@api_router.post("/admin/import/confirm")
async def confirm_import(
    request: Request,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user)
):
    """Backward-compatible import confirm - delegates to execute"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    body = await request.json()
    hotels_to_import = body.get("hotels", [])
    if not hotels_to_import:
        raise HTTPException(status_code=400, detail="No hotels to import")
    
    # Reuse execute logic with inline data
    request._body = json.dumps({"hotels": hotels_to_import, "batch_name": "Quick Import"}).encode()
    return await execute_import(request, background_tasks, current_user)

# ===================== ADMIN HOTEL CRUD =====================

@api_router.post("/admin/hotels")
async def admin_create_hotel(
    request: Request,
    current_user: dict = Depends(get_current_user)
):
    """Admin creates a new hotel manually"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    body = await request.json()
    name = body.get("name", "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Jina la hotel linahitajika")
    
    existing = await db.hotels.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}})
    if existing:
        raise HTTPException(status_code=400, detail="Hotel yenye jina hili ipo tayari")
    
    hotel_id = str(uuid.uuid4())
    hotel = {
        "id": hotel_id,
        "hotel_code": generate_hotel_code(),
        "owner_id": body.get("owner_id"),
        "name": name,
        "description": body.get("description", ""),
        "address": body.get("address", ""),
        "city": body.get("city", ""),
        "phone_number": body.get("phone_number", ""),
        "whatsapp_number": body.get("whatsapp_number", body.get("phone_number", "")),
        "amenities": body.get("amenities", []),
        "cover_photo": None,
        "photos": [],
        "google_maps_url": body.get("google_maps_url", ""),
        "google_rating": body.get("google_rating"),
        "google_review_count": body.get("google_review_count"),
        "latitude": body.get("latitude"),
        "longitude": body.get("longitude"),
        "website": body.get("website", ""),
        "status": body.get("status", "verified"),
        "data_source": "admin_created",
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.hotels.insert_one(hotel)
    hotel.pop("_id", None)
    return hotel

@api_router.put("/admin/hotels/{hotel_id}")
async def admin_update_hotel(
    hotel_id: str,
    request: Request,
    current_user: dict = Depends(get_current_user)
):
    """Admin updates hotel details"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    body = await request.json()
    
    # Fields that can be updated
    updatable = ["name", "description", "address", "city", "phone_number", "whatsapp_number",
                 "amenities", "google_maps_url", "google_rating", "google_review_count",
                 "latitude", "longitude", "website", "status"]
    
    update = {}
    for field in updatable:
        if field in body:
            update[field] = body[field]
    
    if not update:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    # Check duplicate name if name is being changed
    if "name" in update and update["name"].lower() != hotel["name"].lower():
        existing = await db.hotels.find_one({
            "name": {"$regex": f"^{re.escape(update['name'])}$", "$options": "i"},
            "id": {"$ne": hotel_id}
        })
        if existing:
            raise HTTPException(status_code=400, detail="Hotel yenye jina hili ipo tayari")
    
    await db.hotels.update_one({"id": hotel_id}, {"$set": update})
    updated = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    return updated

@api_router.delete("/admin/hotels/{hotel_id}")
async def admin_delete_hotel(
    hotel_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Admin deletes a hotel and its room types"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    # Check for active bookings
    active_bookings = await db.bookings.count_documents({
        "hotel_id": hotel_id,
        "status": {"$in": ["confirmed", "pending"]},
        "checkin_status": {"$ne": "checked_out"}
    })
    if active_bookings > 0:
        raise HTTPException(status_code=400, detail=f"Hotel ina bukini {active_bookings} hai. Haziwezi kufutwa.")
    
    # Delete room types
    await db.room_types.delete_many({"hotel_id": hotel_id})
    # Unassign any cashiers
    await db.users.update_many(
        {"assigned_hotel_id": hotel_id, "role": "cashier"},
        {"$set": {"assigned_hotel_id": None}}
    )
    # Delete hotel
    await db.hotels.delete_one({"id": hotel_id})
    
    return {"message": f"Hotel '{hotel['name']}' imefutwa"}

# ===================== PHOTO MANAGEMENT =====================

@api_router.get("/hotels/{hotel_id}/photos")
async def get_hotel_photos(hotel_id: str):
    """Get all photos for a hotel, ordered by sort_order"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0, "photos": 1})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    photos = hotel.get("photos", [])
    # Sort by sort_order
    if photos and isinstance(photos[0], dict):
        photos.sort(key=lambda p: p.get("sort_order", 0))
    return photos

@api_router.post("/hotels/{hotel_id}/photos")
async def add_hotel_photo(
    hotel_id: str,
    request: Request,
    current_user: dict = Depends(get_current_user)
):
    """Add a new photo to a hotel (from Cloudinary upload)"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    # Check permission: admin or hotel owner
    if current_user["role"] not in ["admin"] and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    body = await request.json()
    photos = hotel.get("photos", [])
    
    # Max 10 photos per hotel
    current_count = len([p for p in photos if isinstance(p, dict)])
    if current_count >= 10:
        raise HTTPException(status_code=400, detail="Kiwango cha juu cha picha 10 kimefikia")
    
    is_first = current_count == 0
    photo_obj = {
        "id": str(uuid.uuid4()),
        "is_primary": is_first or body.get("is_primary", False),
        "sort_order": current_count,
        "local_path": body.get("local_path", ""),
        "cloudinary_original": body.get("cloudinary_original", ""),
        "cloudinary_thumb": body.get("cloudinary_thumb", ""),
        "cloudinary_mobile": body.get("cloudinary_mobile", ""),
        "cloudinary_web": body.get("cloudinary_web", ""),
        "cloudinary_hd": body.get("cloudinary_hd", ""),
        "source": body.get("source", "uploaded"),
        "uploaded_at": datetime.now(timezone.utc).isoformat()
    }
    
    # If this is set as primary, unset others
    if photo_obj["is_primary"]:
        for p in photos:
            if isinstance(p, dict):
                p["is_primary"] = False
    
    photos.append(photo_obj)
    cover = get_cover_url(photos, "cloudinary_web")
    
    await db.hotels.update_one(
        {"id": hotel_id},
        {"$set": {"photos": photos, "cover_photo": cover}}
    )
    return photo_obj

@api_router.patch("/hotels/{hotel_id}/photos/{photo_id}/set-primary")
async def set_primary_photo(
    hotel_id: str,
    photo_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Set a photo as the primary photo for a hotel"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] not in ["admin"] and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    photos = hotel.get("photos", [])
    found = False
    for p in photos:
        if isinstance(p, dict):
            if p["id"] == photo_id:
                p["is_primary"] = True
                found = True
            else:
                p["is_primary"] = False
    
    if not found:
        raise HTTPException(status_code=404, detail="Picha haipatikani")
    
    cover = get_cover_url(photos, "cloudinary_web")
    await db.hotels.update_one(
        {"id": hotel_id},
        {"$set": {"photos": photos, "cover_photo": cover}}
    )
    return {"message": "Picha ya kwanza imebadilishwa"}

@api_router.patch("/hotels/{hotel_id}/photos/reorder")
async def reorder_hotel_photos(
    hotel_id: str,
    request: Request,
    current_user: dict = Depends(get_current_user)
):
    """Reorder photos by providing ordered list of photo IDs"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] not in ["admin"] and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    body = await request.json()
    photo_ids = body.get("photo_ids", [])
    
    photos = hotel.get("photos", [])
    photo_map = {p["id"]: p for p in photos if isinstance(p, dict)}
    
    reordered = []
    for idx, pid in enumerate(photo_ids):
        if pid in photo_map:
            photo_map[pid]["sort_order"] = idx
            reordered.append(photo_map[pid])
    
    # Add any photos not in the order list
    for p in photos:
        if isinstance(p, dict) and p["id"] not in [r["id"] for r in reordered]:
            p["sort_order"] = len(reordered)
            reordered.append(p)
    
    cover = get_cover_url(reordered, "cloudinary_web")
    await db.hotels.update_one(
        {"id": hotel_id},
        {"$set": {"photos": reordered, "cover_photo": cover}}
    )
    return {"message": "Mpangilio wa picha umebadilishwa"}

@api_router.delete("/hotels/{hotel_id}/photos/{photo_id}")
async def delete_hotel_photo(
    hotel_id: str,
    photo_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Delete a photo from a hotel (DB only, NOT from Cloudinary)"""
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] not in ["admin"] and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    photos = hotel.get("photos", [])
    photo_dicts = [p for p in photos if isinstance(p, dict)]
    
    if len(photo_dicts) <= 1:
        raise HTTPException(status_code=400, detail="Lazima kuwe na picha moja angalau")
    
    target = None
    for p in photo_dicts:
        if p["id"] == photo_id:
            target = p
            break
    
    if not target:
        raise HTTPException(status_code=404, detail="Picha haipatikani")
    
    was_primary = target.get("is_primary", False)
    new_photos = [p for p in photos if not (isinstance(p, dict) and p["id"] == photo_id)]
    
    # Auto-promote next photo if deleted was primary
    if was_primary and new_photos:
        for p in new_photos:
            if isinstance(p, dict):
                p["is_primary"] = True
                break
    
    # Reindex sort_order
    for idx, p in enumerate(new_photos):
        if isinstance(p, dict):
            p["sort_order"] = idx
    
    cover = get_cover_url(new_photos, "cloudinary_web")
    await db.hotels.update_one(
        {"id": hotel_id},
        {"$set": {"photos": new_photos, "cover_photo": cover}}
    )
    return {"message": "Picha imefutwa"}

# ===================== ROOM TYPE PHOTO MANAGEMENT =====================

@api_router.post("/rooms/{room_id}/photos")
async def add_room_photo(
    room_id: str,
    request: Request,
    current_user: dict = Depends(get_current_user)
):
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    
    hotel = await db.hotels.find_one({"id": room["hotel_id"]}, {"_id": 0})
    if current_user["role"] not in ["admin"] and hotel.get("owner_id") != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    body = await request.json()
    photos = room.get("photos", [])
    current_count = len([p for p in photos if isinstance(p, dict)])
    if current_count >= 5:
        raise HTTPException(status_code=400, detail="Kiwango cha juu cha picha 5 kimefikia")
    
    is_first = current_count == 0
    photo_obj = {
        "id": str(uuid.uuid4()),
        "is_primary": is_first or body.get("is_primary", False),
        "sort_order": current_count,
        "cloudinary_original": body.get("cloudinary_original", ""),
        "cloudinary_thumb": body.get("cloudinary_thumb", ""),
        "cloudinary_mobile": body.get("cloudinary_mobile", ""),
        "cloudinary_web": body.get("cloudinary_web", ""),
        "cloudinary_hd": body.get("cloudinary_hd", ""),
        "source": body.get("source", "uploaded"),
        "uploaded_at": datetime.now(timezone.utc).isoformat()
    }
    
    if photo_obj["is_primary"]:
        for p in photos:
            if isinstance(p, dict):
                p["is_primary"] = False
    
    photos.append(photo_obj)
    await db.room_types.update_one({"id": room_id}, {"$set": {"photos": photos}})
    return photo_obj

@api_router.patch("/rooms/{room_id}/photos/{photo_id}/set-primary")
async def set_primary_room_photo(room_id: str, photo_id: str, current_user: dict = Depends(get_current_user)):
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    photos = room.get("photos", [])
    found = False
    for p in photos:
        if isinstance(p, dict):
            p["is_primary"] = p["id"] == photo_id
            if p["id"] == photo_id:
                found = True
    if not found:
        raise HTTPException(status_code=404, detail="Picha haipatikani")
    await db.room_types.update_one({"id": room_id}, {"$set": {"photos": photos}})
    return {"message": "Picha ya kwanza imebadilishwa"}

@api_router.patch("/rooms/{room_id}/photos/reorder")
async def reorder_room_photos(room_id: str, request: Request, current_user: dict = Depends(get_current_user)):
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    body = await request.json()
    photo_ids = body.get("photo_ids", [])
    photos = room.get("photos", [])
    photo_map = {p["id"]: p for p in photos if isinstance(p, dict)}
    reordered = []
    for idx, pid in enumerate(photo_ids):
        if pid in photo_map:
            photo_map[pid]["sort_order"] = idx
            reordered.append(photo_map[pid])
    for p in photos:
        if isinstance(p, dict) and p["id"] not in [r["id"] for r in reordered]:
            p["sort_order"] = len(reordered)
            reordered.append(p)
    await db.room_types.update_one({"id": room_id}, {"$set": {"photos": reordered}})
    return {"message": "Mpangilio wa picha umebadilishwa"}

@api_router.delete("/rooms/{room_id}/photos/{photo_id}")
async def delete_room_photo(room_id: str, photo_id: str, current_user: dict = Depends(get_current_user)):
    room = await db.room_types.find_one({"id": room_id}, {"_id": 0})
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    photos = room.get("photos", [])
    photo_dicts = [p for p in photos if isinstance(p, dict)]
    if len(photo_dicts) <= 1:
        raise HTTPException(status_code=400, detail="Lazima kuwe na picha moja angalau")
    target = next((p for p in photo_dicts if p["id"] == photo_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="Picha haipatikani")
    was_primary = target.get("is_primary", False)
    new_photos = [p for p in photos if not (isinstance(p, dict) and p["id"] == photo_id)]
    if was_primary and new_photos:
        for p in new_photos:
            if isinstance(p, dict):
                p["is_primary"] = True
                break
    for idx, p in enumerate(new_photos):
        if isinstance(p, dict):
            p["sort_order"] = idx
    await db.room_types.update_one({"id": room_id}, {"$set": {"photos": new_photos}})
    return {"message": "Picha imefutwa"}

# ===================== CLOUDINARY CONFIG =====================

@api_router.get("/config/cloudinary")
async def get_cloudinary_config():
    """Return Cloudinary config for frontend direct uploads"""
    return {
        "cloud_name": CLOUDINARY_CLOUD_NAME,
        "upload_preset": CLOUDINARY_UPLOAD_PRESET,
        "folder": "hotel_research/owner_uploads"
    }

@api_router.post("/admin/hotels/{hotel_id}/attach-owner")
async def attach_owner_to_hotel(
    hotel_id: str,
    request: Request,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user)
):
    """Attach an owner to an imported hotel and send SMS invitation"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    body = await request.json()
    owner_name = body.get("owner_name", "")
    owner_phone = body.get("owner_phone", "")
    owner_email = body.get("owner_email", "")
    
    if not owner_name or not owner_phone or not owner_email:
        raise HTTPException(status_code=400, detail="Owner name, phone and email required")
    
    hotel = await db.hotels.find_one({"id": hotel_id}, {"_id": 0})
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    # Check if owner already exists
    existing_owner = await db.users.find_one({"email": owner_email}, {"_id": 0})
    
    if existing_owner:
        owner_id = existing_owner["id"]
        if existing_owner["role"] != "owner":
            raise HTTPException(status_code=400, detail="User exists but is not an owner")
    else:
        # Create new owner account
        temp_password = generate_temp_password()
        owner_id = str(uuid.uuid4())
        
        owner = {
            "id": owner_id,
            "email": owner_email,
            "full_name": owner_name,
            "phone": format_phone(owner_phone),
            "role": "owner",
            "assigned_hotel_id": None,
            "password_hash": get_password_hash(temp_password),
            "is_active": True,
            "is_verified": True,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "last_login": None
        }
        await db.users.insert_one(owner)
        
        # Send SMS invitation
        sms_msg = f"Habari {owner_name}! Umealikwa kusimamia {hotel['name']} kwenye Habari Stays. Ingia: habaristays.com/login Email: {owner_email} Neno la siri: {temp_password}"
        background_tasks.add_task(send_sms, owner_phone, sms_msg)
    
    # Attach owner to hotel
    await db.hotels.update_one(
        {"id": hotel_id},
        {"$set": {"owner_id": owner_id}}
    )
    
    return {"message": f"Mmiliki amewekwa kwa {hotel['name']}", "owner_id": owner_id}

@api_router.get("/admin/all-hotels")
async def get_all_hotels_admin(
    status_filter: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Get all hotels for admin with full details"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    query = {}
    if status_filter:
        query["status"] = status_filter
    
    hotels = await db.hotels.find(query, {"_id": 0}).sort("created_at", -1).to_list(500)
    
    result = []
    for hotel in hotels:
        rooms = await db.room_types.find({"hotel_id": hotel["id"]}, {"_id": 0}).to_list(100)
        total_rooms = sum(r.get("total_rooms", 0) for r in rooms)
        
        owner = None
        if hotel.get("owner_id"):
            owner = await db.users.find_one({"id": hotel["owner_id"]}, {"_id": 0, "password_hash": 0})
        
        result.append({
            "id": hotel["id"],
            "hotel_code": hotel.get("hotel_code", ""),
            "name": hotel["name"],
            "city": hotel.get("city", ""),
            "address": hotel.get("address", ""),
            "phone_number": hotel.get("phone_number", ""),
            "status": hotel.get("status", "pending"),
            "data_source": hotel.get("data_source", "manual"),
            "owner_id": hotel.get("owner_id"),
            "owner_name": owner["full_name"] if owner else None,
            "owner_email": owner["email"] if owner else None,
            "total_rooms": total_rooms,
            "has_rooms": total_rooms > 0,
            "room_type_count": len(rooms),
            "has_default_rooms": any(r.get("is_default") for r in rooms),
            "google_rating": hotel.get("google_rating"),
            "photo_count": len([p for p in hotel.get("photos", []) if isinstance(p, dict)]) or len(hotel.get("photos", [])),
            "cover_photo": get_cover_url(hotel.get("photos", []), "cloudinary_thumb") or hotel.get("cover_photo"),
            "created_at": hotel["created_at"]
        })
    
    return result

@api_router.put("/admin/users/{user_id}/reject")
async def reject_owner(user_id: str, current_user: dict = Depends(get_current_user)):
    """Reject a pending owner application"""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    user = await db.users.find_one({"id": user_id}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=404, detail="User hapatikani")
    
    await db.users.update_one({"id": user_id}, {"$set": {"is_active": False, "is_verified": False}})
    return {"message": "Mmiliki amekataliwa"}

# ===================== PAYMENT WEBHOOK =====================

@api_router.post("/payments/callback")
async def payment_callback(request: Request, background_tasks: BackgroundTasks):
    """Webhook for payment provider callbacks"""
    body = await request.json()
    logging.info(f"Payment callback received: {body}")
    
    # Extract booking reference from callback
    booking_ref = body.get("reference") or body.get("booking_ref")
    if booking_ref:
        booking = await db.bookings.find_one({"booking_ref": booking_ref}, {"_id": 0})
        if booking and booking["payment_status"] != "paid":
            await db.bookings.update_one(
                {"booking_ref": booking_ref},
                {"$set": {"payment_status": "paid", "status": "confirmed"}}
            )
            
            hotel = await db.hotels.find_one({"id": booking["hotel_id"]}, {"_id": 0})
            sms_msg = f"HABARI STAYS: Booking imethibitishwa! Hotel: {hotel['name']}, Ref: {booking_ref}. Asante! habaristays.com"
            background_tasks.add_task(send_sms, booking["guest_phone"], sms_msg)
    
    return {"status": "received"}

# ===================== CONFIG =====================

@api_router.get("/config/payment")
async def get_payment_config():
    """Get payment configuration for frontend"""
    return {
        "selcom_till": SELCOM_TILL_NUMBER,
        "expiry_seconds": BOOKING_EXPIRY_SECONDS,
        "cloudinary_cloud_name": CLOUDINARY_CLOUD_NAME,
        "cloudinary_upload_preset": CLOUDINARY_UPLOAD_PRESET
    }

# ===================== HEALTH CHECK =====================

@api_router.get("/health")
async def health_check():
    return {"status": "healthy", "service": "Habari Stays API", "version": "2.0.0"}

# Include router
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("startup")
async def startup_event():
    # Start booking expiry background task
    asyncio.create_task(expire_pending_bookings())
    logger.info("Started booking expiry background task")

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
