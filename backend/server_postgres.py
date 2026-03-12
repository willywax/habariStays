"""
Habari Stays API Server
PostgreSQL + SQLAlchemy ORM Version
"""

from fastapi import FastAPI, APIRouter, HTTPException, Depends, status, BackgroundTasks, Response, Request, Cookie, UploadFile, File, Form
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, update, delete, func, and_, or_, desc
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Any
import uuid
from datetime import datetime, timezone, timedelta
import bcrypt
from jose import JWTError, jwt
import re
import httpx
import asyncio
import io
import json
import openpyxl

from database import get_db_session, init_db, close_db, async_session_factory
from models import (
    User, UserSession, Hotel, RoomType, Booking, Review,
    CashierAssignment, CashierActivityLog, ImportBatch, ImportPreview, generate_uuid
)
import crud

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# JWT Settings
SECRET_KEY = os.environ.get('SECRET_KEY', 'habari-stays-secret-key-change-in-production')
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24 * 7  # 7 days

# Security
security = HTTPBearer(auto_error=False)

# Beem SMS Configuration
BEEM_API_KEY = os.environ.get('BEEM_API_KEY', '')
BEEM_SECRET_KEY = os.environ.get('BEEM_SECRET_KEY', '')
BEEM_SENDER_ID = os.environ.get('BEEM_SENDER_ID', 'HABARISTAYS')

# Payment Configuration
SELCOM_TILL_NUMBER = os.environ.get('SELCOM_TILL_NUMBER', '123456')

# Cloudinary Configuration
CLOUDINARY_CLOUD_NAME = os.environ.get('CLOUDINARY_CLOUD_NAME', 'diupey6vs')
CLOUDINARY_UPLOAD_PRESET = os.environ.get('CLOUDINARY_UPLOAD_PRESET', 'habari_stays_upload')

# Booking expiry time in seconds
BOOKING_EXPIRY_SECONDS = 120  # 2 minutes

# Hotel-level and Room-level amenity constants
HOTEL_AMENITIES = ["Breakfast", "Parking", "WiFi", "Hot Water", "Bar"]
ROOM_AMENITIES = ["A/C", "Western Toilet", "Squat Toilet", "En-suite Bathroom", "Balcony", "TV", "Safe", "Mini Fridge", "Hot Shower"]


# Helper: extract cover photo URL from photos list
def get_cover_url(photos, size="cloudinary_mobile"):
    if not photos:
        return None
    for p in photos:
        if isinstance(p, dict) and p.get("is_primary"):
            return p.get(size) or p.get("cloudinary_original") or p.get("cloudinary_web")
    first = photos[0] if photos else None
    if isinstance(first, dict):
        return first.get(size) or first.get("cloudinary_original") or first.get("cloudinary_web")
    return first if isinstance(first, str) else None


async def create_default_room_type(session: AsyncSession, hotel_id: str, price: int = 50000):
    """Create a default Standard room type for a newly imported hotel."""
    room = RoomType(
        id=generate_uuid(),
        hotel_id=hotel_id,
        name="Standard",
        description="",
        price_per_night=price if price > 0 else 50000,
        capacity=2,
        total_rooms=1,
        available_rooms=1,
        amenities=[],
        photos=[],
        is_default=True,
    )
    session.add(room)
    await session.flush()
    return room


app = FastAPI(title="Habari Stays API", version="2.1.0")

# CORS must be configured early - when allow_credentials=True, cannot use allow_origins=["*"]
cors_origins_str = os.getenv("CORS_ORIGINS", "http://localhost:3000,http://localhost:8000,https://habaristays.com")
cors_origins = [origin.strip() for origin in cors_origins_str.split(",") if origin.strip()]

# Add CORS middleware with proper configuration
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS", "PATCH"],
    allow_headers=["*"],
    expose_headers=["*"],
)

api_router = APIRouter(prefix="/api")

# ===================== PYDANTIC MODELS =====================

class UserBase(BaseModel):
    email: EmailStr
    full_name: str
    phone: str
    role: str = "traveler"
    assigned_hotel_id: Optional[str] = None

class UserCreate(UserBase):
    password: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserResponse(BaseModel):
    id: str
    email: str
    full_name: str
    phone: str
    role: str
    assigned_hotel_id: Optional[str] = None
    is_verified: bool = True
    picture: Optional[str] = None
    created_at: str

class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    user: UserResponse

class HotelBase(BaseModel):
    name: str
    description: str
    address: str
    city: str
    phone_number: str
    whatsapp_number: Optional[str] = None
    amenities: List[str] = []
    cover_photo: Optional[str] = None
    photos: list = []
    google_maps_url: Optional[str] = None
    google_rating: Optional[float] = None
    google_review_count: Optional[int] = None
    latitude: Optional[float] = None
    longitude: Optional[float] = None

class HotelCreate(HotelBase):
    pass

class HotelResponse(HotelBase):
    id: str
    hotel_code: str
    owner_id: Optional[str] = None
    status: str = "pending"
    data_source: str = "manual"
    created_at: str
    total_rooms: int = 0
    available_rooms: int = 0
    bookings_this_month: int = 0
    revenue_this_month: float = 0
    average_rating: float = 0
    review_count: int = 0
    min_price: Optional[int] = None

class RoomTypeBase(BaseModel):
    name: str
    description: str = ""
    price_per_night: int
    capacity: int = 2
    total_rooms: int
    available_rooms: int
    amenities: List[str] = []
    photos: list = []
    is_default: bool = False

class RoomTypeCreate(RoomTypeBase):
    hotel_id: str

class RoomTypeResponse(RoomTypeBase):
    id: str
    hotel_id: str
    is_default: bool = False
    created_at: str

class BookingCreate(BaseModel):
    hotel_id: str
    room_type_id: str
    guest_name: str
    guest_phone: str
    guest_email: Optional[EmailStr] = None
    checkin_date: str
    checkout_date: str
    num_guests: int = 1
    booking_type: str = "online"
    payment_method: str = "mpesa"
    notes: Optional[str] = None

class BookingResponse(BaseModel):
    id: str
    booking_ref: str
    hotel_id: str
    hotel_name: str
    room_type_id: str
    room_type_name: str
    guest_name: str
    guest_phone: str
    guest_email: Optional[str] = None
    checkin_date: str
    checkout_date: str
    nights: int
    num_guests: int
    total_amount_tzs: int
    booking_type: str
    payment_method: str
    payment_status: str
    payment_reference: Optional[str] = None
    status: str
    checkin_status: str
    actual_checkin_time: Optional[str] = None
    actual_checkout_time: Optional[str] = None
    created_by: Optional[str] = None
    created_by_name: Optional[str] = None
    notes: Optional[str] = None
    created_at: str
    expires_at: Optional[str] = None

class ReviewCreate(BaseModel):
    booking_id: str
    hotel_id: str
    rating: int
    comment: str

class ReviewResponse(BaseModel):
    id: str
    booking_id: str
    hotel_id: str
    guest_name: str
    guest_phone: str
    rating: int
    comment: str
    created_at: str

class CashierCreate(BaseModel):
    full_name: str
    phone: str
    email: EmailStr
    hotel_id: str

class CashierResponse(BaseModel):
    id: str
    full_name: str
    phone: str
    email: str
    role: str
    assigned_hotel_id: Optional[str] = None
    assigned_hotel_name: str
    is_active: bool
    created_at: str
    last_login: Optional[str] = None

class ShiftSummary(BaseModel):
    cashier_id: str
    cashier_name: str
    date: str
    walkins_count: int
    walkins_total: int
    online_checkins_count: int
    checkouts_count: int
    cash_collected: int

class PaymentConfirmation(BaseModel):
    booking_id: str
    payment_reference: str

# ===================== UTILITIES =====================

def generate_booking_ref():
    year = datetime.now().year
    random_num = str(uuid.uuid4().int)[:4]
    return f"HS-{year}-{random_num}"

def generate_hotel_code():
    """Generate a hotel code with 6 random digits for lower collision probability."""
    random_num = str(uuid.uuid4().int)[:6]
    return f"HTL-{random_num}"

async def generate_unique_hotel_code(session: AsyncSession, max_attempts: int = 10):
    """Generate a unique hotel code, checking database for duplicates."""
    for _ in range(max_attempts):
        code = generate_hotel_code()
        result = await session.execute(
            select(Hotel).where(Hotel.hotel_code == code)
        )
        if not result.scalar_one_or_none():
            return code
    # Fallback: use UUID-based code
    return f"HTL-{str(uuid.uuid4())[:8].upper()}"

def generate_temp_password():
    return str(uuid.uuid4())[:8]

def verify_password(plain_password, hashed_password):
    return bcrypt.checkpw(plain_password.encode('utf-8'), hashed_password.encode('utf-8'))

def get_password_hash(password):
    return bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')

def create_access_token(data: dict):
    to_encode = data.copy()
    expire = datetime.now(timezone.utc) + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)

async def get_current_user(
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    # Check cookie first
    session_token = request.cookies.get("session_token")
    
    if session_token:
        user_session = await crud.get_session_by_token(session, session_token)
        if user_session:
            if user_session.expires_at.tzinfo is None:
                expires_at = user_session.expires_at.replace(tzinfo=timezone.utc)
            else:
                expires_at = user_session.expires_at
            if expires_at > datetime.now(timezone.utc):
                user = await crud.get_user_by_id(session, user_session.user_id)
                if user:
                    return user.to_dict()
    
    # Fallback to JWT token
    if credentials:
        token = credentials.credentials
        try:
            payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
            user_id: str = payload.get("sub")
            if user_id:
                user = await crud.get_user_by_id(session, user_id)
                if user:
                    return user.to_dict()
        except JWTError:
            pass
    
    raise HTTPException(status_code=401, detail="Not authenticated")

async def get_current_user_optional(
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    credentials: HTTPAuthorizationCredentials = Depends(security)
):
    try:
        return await get_current_user(request, session, credentials)
    except HTTPException:
        return None

def validate_tanzania_phone(phone: str) -> bool:
    pattern = r'^\+255[6-8]\d{8}$'
    return bool(re.match(pattern, phone))

def format_phone(phone: str) -> str:
    cleaned = re.sub(r'[\s\-\(\)]', '', phone.strip())
    if cleaned.startswith("+255"):
        return cleaned
    if cleaned.startswith("0"):
        return "+255" + cleaned[1:]
    if cleaned.startswith("255"):
        return "+" + cleaned
    return "+255" + cleaned

def format_tzs(amount: int) -> str:
    return f"TZS {amount:,}"

async def send_sms(phone: str, message: str):
    """Send SMS via Beem Africa API"""
    if not BEEM_API_KEY or not BEEM_SECRET_KEY:
        logging.info(f"SMS (mock): To {phone}: {message}")
        return True
    
    try:
        formatted_phone = format_phone(phone)
        if formatted_phone.startswith("+"):
            formatted_phone = formatted_phone[1:]
        
        url = "https://apisms.beem.africa/v1/send"
        payload = {
            "source_addr": BEEM_SENDER_ID,
            "encoding": 0,
            "message": message,
            "recipients": [{"recipient_id": 1, "dest_addr": formatted_phone}]
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(
                url,
                json=payload,
                auth=(BEEM_API_KEY, BEEM_SECRET_KEY),
                timeout=30
            )
            logging.info(f"SMS sent to {formatted_phone}: {response.status_code}")
            return response.status_code == 200
    except Exception as e:
        logging.error(f"SMS failed to {phone}: {str(e)}")
        logging.info(f"SMS (mock fallback): To {phone}: {message}")
        return False

# ===================== BACKGROUND TASKS =====================

async def expire_pending_bookings():
    """Background task to expire unpaid bookings after 2 minutes"""
    while True:
        try:
            async with async_session_factory() as session:
                now = datetime.now(timezone.utc)
                expired_bookings = await crud.get_expired_bookings(session, now)
                
                for booking in expired_bookings:
                    # Update booking status
                    await crud.update_booking(session, booking.id, {"status": "expired"})
                    
                    # Restore room availability
                    await crud.increment_room_availability(session, booking.room_type_id, 1)
                    
                    # Send expiry SMS
                    sms_msg = f"HABARI STAYS: Booking {booking.booking_ref} imeisha muda. Jaribu tena: habaristays.com"
                    await send_sms(booking.guest_phone, sms_msg)
                    
                    logging.info(f"Expired booking: {booking.booking_ref}")
                
                await session.commit()
        except Exception as e:
            logging.error(f"Expiry task error: {str(e)}")
        
        await asyncio.sleep(60)

# ===================== CASHIER ACTIVITY LOG =====================

async def log_cashier_activity(
    session: AsyncSession,
    cashier_id: str,
    hotel_id: str,
    booking_id: str,
    action_type: str,
    payment_method: str = None,
    amount_tzs: int = 0,
    guest_name: str = "",
    room_type_name: str = ""
):
    await crud.create_activity_log(session, {
        "id": generate_uuid(),
        "cashier_id": cashier_id,
        "hotel_id": hotel_id,
        "booking_id": booking_id,
        "action_type": action_type,
        "payment_method": payment_method,
        "amount_tzs": amount_tzs,
        "guest_name": guest_name,
        "room_type_name": room_type_name,
    })

# ===================== AUTH ENDPOINTS =====================

@api_router.post("/auth/register", response_model=TokenResponse)
async def register(user_data: UserCreate, session: AsyncSession = Depends(get_db_session)):
    existing = await crud.get_user_by_email(session, user_data.email)
    if existing:
        raise HTTPException(status_code=400, detail="Email tayari imetumika / Email already exists")
    
    user_id = generate_uuid()
    is_verified = user_data.role not in ["owner"]
    
    user = await crud.create_user(session, {
        "id": user_id,
        "email": user_data.email,
        "full_name": user_data.full_name,
        "phone": format_phone(user_data.phone),
        "role": user_data.role,
        "assigned_hotel_id": user_data.assigned_hotel_id,
        "password_hash": get_password_hash(user_data.password),
        "is_active": True,
        "is_verified": is_verified,
        "picture": None,
    })
    
    if user_data.role == "owner":
        return TokenResponse(
            access_token="",
            user=UserResponse(
                id=user.id,
                email=user.email,
                full_name=user.full_name,
                phone=user.phone,
                role=user.role,
                assigned_hotel_id=user.assigned_hotel_id,
                is_verified=False,
                picture=None,
                created_at=user.created_at.isoformat()
            )
        )
    
    token = create_access_token({"sub": user_id, "role": user.role})
    return TokenResponse(
        access_token=token,
        user=UserResponse(
            id=user.id,
            email=user.email,
            full_name=user.full_name,
            phone=user.phone,
            role=user.role,
            assigned_hotel_id=user.assigned_hotel_id,
            is_verified=user.is_verified,
            picture=None,
            created_at=user.created_at.isoformat()
        )
    )

@api_router.post("/auth/login", response_model=TokenResponse)
async def login(credentials: UserLogin, response: Response, session: AsyncSession = Depends(get_db_session)):
    user = await crud.get_user_by_email(session, credentials.email)
    if not user or not verify_password(credentials.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Email au neno la siri si sahihi / Invalid credentials")
    
    if not user.is_active:
        raise HTTPException(status_code=401, detail="Akaunti imezimwa / Account disabled")
    
    if user.role == "owner" and not user.is_verified:
        raise HTTPException(status_code=401, detail="Akaunti haijathibitishwa / Account pending verification")
    
    await crud.update_user(session, user.id, {"last_login": datetime.now(timezone.utc)})
    
    token = create_access_token({"sub": user.id, "role": user.role})
    
    # Set session cookie
    session_token = str(uuid.uuid4())
    expires = datetime.now(timezone.utc) + timedelta(days=7)
    await crud.create_user_session(session, user.id, session_token, expires)
    
    response.set_cookie(
        key="session_token",
        value=session_token,
        httponly=True,
        max_age=60*60*24*7,
        samesite="lax"
    )
    
    return TokenResponse(
        access_token=token,
        user=UserResponse(
            id=user.id,
            email=user.email,
            full_name=user.full_name,
            phone=user.phone,
            role=user.role,
            assigned_hotel_id=user.assigned_hotel_id,
            is_verified=user.is_verified,
            picture=user.picture,
            created_at=user.created_at.isoformat()
        )
    )

@api_router.post("/auth/logout")
async def logout(response: Response, request: Request, session: AsyncSession = Depends(get_db_session)):
    session_token = request.cookies.get("session_token")
    if session_token:
        user_session = await crud.get_session_by_token(session, session_token)
        if user_session:
            await session.delete(user_session)
    response.delete_cookie("session_token")
    return {"message": "Umeondoka"}

@api_router.get("/auth/me", response_model=UserResponse)
async def get_me(current_user: dict = Depends(get_current_user)):
    return UserResponse(**{k: v for k, v in current_user.items() if k in UserResponse.model_fields})

# ===================== HOTEL ENDPOINTS =====================

@api_router.post("/hotels", response_model=dict)
async def create_hotel(
    hotel_data: HotelCreate,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    hotel_id = generate_uuid()
    hotel_code = await generate_unique_hotel_code(session)
    hotel = await crud.create_hotel(session, {
        "id": hotel_id,
        "hotel_code": hotel_code,
        "owner_id": current_user["id"] if current_user["role"] == "owner" else None,
        "name": hotel_data.name,
        "description": hotel_data.description,
        "address": hotel_data.address,
        "city": hotel_data.city,
        "phone_number": hotel_data.phone_number,
        "whatsapp_number": hotel_data.whatsapp_number or hotel_data.phone_number,
        "amenities": hotel_data.amenities,
        "cover_photo": hotel_data.cover_photo,
        "photos": hotel_data.photos,
        "google_maps_url": hotel_data.google_maps_url,
        "google_rating": hotel_data.google_rating,
        "google_review_count": hotel_data.google_review_count,
        "latitude": hotel_data.latitude,
        "longitude": hotel_data.longitude,
        "status": "pending",
        "data_source": "manual",
    })
    
    await create_default_room_type(session, hotel_id)
    
    return hotel.to_dict()

@api_router.get("/hotels")
async def get_hotels(
    city: Optional[str] = None,
    search: Optional[str] = None,
    session: AsyncSession = Depends(get_db_session)
):
    query_params = {}
    if city:
        query_params["city"] = city
    if search:
        query_params["search"] = search
    
    # Only show verified/imported hotels to public
    hotels = await crud.get_hotels(session, **query_params, status=None)
    hotels = [h for h in hotels if h.status in ["verified", "imported"]]
    
    result = []
    for hotel in hotels:
        rooms = await crud.get_room_types_by_hotel(session, hotel.id)
        total_rooms = sum(r.total_rooms for r in rooms)
        available_rooms = sum(r.available_rooms for r in rooms)
        min_price = min((r.price_per_night for r in rooms), default=None)
        
        reviews = await crud.get_reviews_by_hotel(session, hotel.id)
        avg_rating = sum(r.rating for r in reviews) / len(reviews) if reviews else 0
        
        start_of_month = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        bookings_count = await crud.count_bookings(session, hotel_id=hotel.id, created_after=start_of_month)
        revenue = await crud.get_revenue_sum(session, hotel_id=hotel.id, created_after=start_of_month)
        
        hotel_dict = hotel.to_dict()
        hotel_dict.update({
            "total_rooms": total_rooms,
            "available_rooms": available_rooms,
            "min_price": min_price,
            "average_rating": round(avg_rating, 1),
            "review_count": len(reviews),
            "bookings_this_month": bookings_count,
            "revenue_this_month": revenue,
        })
        result.append(hotel_dict)
    
    return result

@api_router.get("/hotels/cities")
async def get_cities(session: AsyncSession = Depends(get_db_session)):
    cities = await crud.get_distinct_cities(session)
    return [{"city": c, "count": 0} for c in cities]

@api_router.get("/hotels/{hotel_id}")
async def get_hotel(hotel_id: str, session: AsyncSession = Depends(get_db_session)):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    rooms = await crud.get_room_types_by_hotel(session, hotel.id)
    reviews = await crud.get_reviews_by_hotel(session, hotel.id)
    avg_rating = sum(r.rating for r in reviews) / len(reviews) if reviews else 0
    
    hotel_dict = hotel.to_dict()
    hotel_dict.update({
        "room_types": [r.to_dict() for r in rooms],
        "reviews": [r.to_dict() for r in reviews],
        "average_rating": round(avg_rating, 1),
        "review_count": len(reviews),
    })
    
    return hotel_dict

@api_router.get("/hotels/{hotel_id}/full")
async def get_hotel_full(
    hotel_id: str,
    session: AsyncSession = Depends(get_db_session)
):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    rooms = await crud.get_room_types_by_hotel(session, hotel.id)
    reviews = await crud.get_reviews_by_hotel(session, hotel.id)
    owner = await crud.get_user_by_id(session, hotel.owner_id) if hotel.owner_id else None
    
    start_of_month = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    bookings_count = await crud.count_bookings(session, hotel_id=hotel.id, created_after=start_of_month)
    avg_rating = sum(r.rating for r in reviews) / len(reviews) if reviews else 0
    
    hotel_dict = hotel.to_dict()
    hotel_dict.update({
        "room_types": [r.to_dict() for r in rooms],
        "reviews": [r.to_dict() for r in reviews],
        "owner": owner.to_dict() if owner else None,
        "total_rooms": sum(r.total_rooms for r in rooms),
        "available_rooms": sum(r.available_rooms for r in rooms),
        "room_type_count": len(rooms),
        "has_default_rooms": any(r.is_default for r in rooms),
        "bookings_this_month": bookings_count,
        "average_rating": round(avg_rating, 1),
        "review_count": len(reviews),
    })
    
    return hotel_dict

@api_router.get("/owner/hotels")
async def get_owner_hotels(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    if current_user["role"] == "owner":
        hotels = await crud.get_hotels_by_owner(session, current_user["id"])
    else:
        hotels = await crud.get_hotels(session)
    
    result = []
    for hotel in hotels:
        rooms = await crud.get_room_types_by_hotel(session, hotel.id)
        total_rooms = sum(r.total_rooms for r in rooms)
        available_rooms = sum(r.available_rooms for r in rooms)
        
        start_of_month = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        bookings = await crud.count_bookings(session, hotel_id=hotel.id, created_after=start_of_month)
        revenue = await crud.get_revenue_sum(session, hotel_id=hotel.id, created_after=start_of_month)
        
        reviews = await crud.get_reviews_by_hotel(session, hotel.id)
        avg_rating = sum(r.rating for r in reviews) / len(reviews) if reviews else 0
        
        hotel_dict = hotel.to_dict()
        hotel_dict.update({
            "total_rooms": total_rooms,
            "available_rooms": available_rooms,
            "bookings_this_month": bookings,
            "revenue_this_month": revenue,
            "average_rating": round(avg_rating, 1),
            "review_count": len(reviews),
        })
        result.append(hotel_dict)
    
    return result

@api_router.put("/hotels/{hotel_id}")
async def update_hotel(
    hotel_id: str,
    hotel_data: HotelBase,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] == "owner" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    update_data = hotel_data.model_dump()
    updated_hotel = await crud.update_hotel(session, hotel_id, update_data)
    
    rooms = await crud.get_room_types_by_hotel(session, hotel_id)
    reviews = await crud.get_reviews_by_hotel(session, hotel_id)
    
    result = updated_hotel.to_dict()
    result["room_types"] = [r.to_dict() for r in rooms]
    result["reviews"] = [r.to_dict() for r in reviews]
    
    return result

@api_router.patch("/hotels/{hotel_id}")
async def patch_hotel(
    hotel_id: str,
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] == "owner" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    body = await request.json()
    
    allowed_fields = ["name", "description", "address", "city", "phone_number", "whatsapp_number",
                      "amenities", "photos", "cover_photo", "google_maps_url", "status"]
    update_data = {k: v for k, v in body.items() if k in allowed_fields}
    
    if not update_data:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    # Update cover_photo from photos if changed
    if "photos" in update_data:
        update_data["cover_photo"] = get_cover_url(update_data["photos"], "cloudinary_web")
    
    updated = await crud.update_hotel(session, hotel_id, update_data)
    return updated.to_dict()

# ===================== ROOM TYPE ENDPOINTS =====================

@api_router.post("/rooms")
@api_router.post("/room-types")
async def create_room_type(
    room_data: RoomTypeCreate,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    hotel = await crud.get_hotel_by_id(session, room_data.hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] == "owner" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    room = await crud.create_room_type(session, {
        "id": generate_uuid(),
        "hotel_id": room_data.hotel_id,
        "name": room_data.name,
        "description": room_data.description,
        "price_per_night": room_data.price_per_night,
        "capacity": room_data.capacity,
        "total_rooms": room_data.total_rooms,
        "available_rooms": room_data.available_rooms,
        "amenities": room_data.amenities,
        "photos": room_data.photos,
        "is_default": False,
    })
    
    return room.to_dict()

@api_router.post("/hotels/{hotel_id}/room-types")
async def create_room_type_for_hotel(
    hotel_id: str,
    room_data: RoomTypeBase,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    payload = RoomTypeCreate(**room_data.model_dump(), hotel_id=hotel_id)
    return await create_room_type(payload, session, current_user)


@api_router.get("/rooms/{room_id}")
@api_router.get("/room-types/{room_id}")
async def get_room_type(room_id: str, session: AsyncSession = Depends(get_db_session)):
    room = await crud.get_room_type_by_id(session, room_id)
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    return room.to_dict()


@api_router.get("/room-types")
async def list_room_types(
    hotel_id: Optional[str] = None,
    session: AsyncSession = Depends(get_db_session)
):
    if not hotel_id:
        raise HTTPException(status_code=400, detail="hotel_id inahitajika")
    rooms = await crud.get_room_types_by_hotel(session, hotel_id)
    return [r.to_dict() for r in rooms]

@api_router.patch("/room-types/{room_id}")
async def patch_room_type(
    room_id: str,
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    room = await crud.get_room_type_by_id(session, room_id)
    if not room:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    
    hotel = await crud.get_hotel_by_id(session, room.hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] != "admin" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    body = await request.json()
    updatable = [
        "name", "description", "price_per_night", "capacity",
        "total_rooms", "available_rooms", "amenities"
    ]
    update = {}
    for field in updatable:
        if field in body:
            update[field] = body[field]
    
    if "amenities" in update:
        update["amenities"] = [a for a in update["amenities"] if a in ROOM_AMENITIES]
    
    if not update:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    total = int(update.get("total_rooms", room.total_rooms))
    if "total_rooms" in update:
        update["total_rooms"] = total
    
    if "available_rooms" in update or "total_rooms" in update:
        available = int(update.get("available_rooms", room.available_rooms))
        available = max(0, min(total, available))
        update["available_rooms"] = available
    
    if room.is_default:
        update["is_default"] = False
    
    await crud.update_room_type(session, room_id, update)
    updated = await crud.get_room_type_by_id(session, room_id)
    return updated.to_dict()
@api_router.patch("/rooms/{room_id}")
async def patch_room_type(
    room_id: str,
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    room = await crud.get_room_type_by_id(session, room_id)
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    
    hotel = await crud.get_hotel_by_id(session, room.hotel_id)
    if current_user["role"] == "owner" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    body = await request.json()
    allowed_fields = ["name", "description", "price_per_night", "capacity", "total_rooms",
                      "available_rooms", "amenities", "photos", "is_default"]
    update_data = {k: v for k, v in body.items() if k in allowed_fields}
    
    if not update_data:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    updated = await crud.update_room_type(session, room_id, update_data)
    return updated.to_dict()

@api_router.get("/hotels/{hotel_id}/rooms")
async def get_hotel_rooms(hotel_id: str, session: AsyncSession = Depends(get_db_session)):
    rooms = await crud.get_room_types_by_hotel(session, hotel_id)
    return [r.to_dict() for r in rooms]

@api_router.put("/rooms/{room_id}")
@api_router.put("/room-types/{room_id}")
async def update_room_type(
    room_id: str,
    room_data: RoomTypeBase,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    room = await crud.get_room_type_by_id(session, room_id)
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    
    hotel = await crud.get_hotel_by_id(session, room.hotel_id)
    if current_user["role"] == "owner" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    updated = await crud.update_room_type(session, room_id, room_data.model_dump())
    return updated.to_dict()

@api_router.delete("/rooms/{room_id}")
@api_router.delete("/room-types/{room_id}")
async def delete_room_type(
    room_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    room = await crud.get_room_type_by_id(session, room_id)
    if not room:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    
    hotel = await crud.get_hotel_by_id(session, room.hotel_id)
    if current_user["role"] == "owner" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    # Check for active bookings
    active = await crud.count_bookings(
        session,
        hotel_id=hotel.id,
        status=["confirmed", "pending"],
        checkin_status="checked_in"
    )
    if active > 0:
        raise HTTPException(status_code=400, detail="Room type ina bukini hai, haiwezi kufutwa")
    
    await crud.delete_room_type(session, room_id)
    return {"message": "Room type imefutwa"}

# ===================== BOOKING ENDPOINTS =====================

@api_router.post("/bookings", response_model=BookingResponse)
async def create_booking(
    booking_data: BookingCreate,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user_optional)
):
    hotel = await crud.get_hotel_by_id(session, booking_data.hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    room_type = await crud.get_room_type_by_id(session, booking_data.room_type_id)
    if not room_type:
        raise HTTPException(status_code=404, detail="Room type haipatikani")
    
    if room_type.available_rooms < 1:
        raise HTTPException(status_code=400, detail="Hakuna vyumba vilivyopo")
    
    # Calculate nights and total
    checkin_dt = datetime.fromisoformat(booking_data.checkin_date)
    checkout_dt = datetime.fromisoformat(booking_data.checkout_date)
    nights = (checkout_dt - checkin_dt).days
    
    if nights < 1:
        raise HTTPException(status_code=400, detail="Tarehe si sahihi")
    
    total_amount = nights * room_type.price_per_night
    
    booking_id = generate_uuid()
    booking_ref = generate_booking_ref()
    now = datetime.now(timezone.utc)
    expires_at = now + timedelta(seconds=BOOKING_EXPIRY_SECONDS)
    
    booking = await crud.create_booking(session, {
        "id": booking_id,
        "booking_ref": booking_ref,
        "hotel_id": booking_data.hotel_id,
        "room_type_id": booking_data.room_type_id,
        "guest_name": booking_data.guest_name,
        "guest_phone": format_phone(booking_data.guest_phone),
        "guest_email": booking_data.guest_email,
        "checkin_date": booking_data.checkin_date,
        "checkout_date": booking_data.checkout_date,
        "nights": nights,
        "num_guests": booking_data.num_guests,
        "total_amount_tzs": total_amount,
        "booking_type": booking_data.booking_type,
        "payment_method": booking_data.payment_method,
        "payment_status": "unpaid",
        "status": "pending",
        "checkin_status": "not_checked_in",
        "created_by": current_user["id"] if current_user else None,
        "notes": booking_data.notes,
        "expires_at": expires_at,
    })
    
    # Decrement availability
    await crud.increment_room_availability(session, room_type.id, -1)
    
    # Send SMS
    if hotel.owner_id:
        owner = await crud.get_user_by_id(session, hotel.owner_id)
        if owner:
            sms_msg = f"HABARI STAYS: Booking mpya! {booking.guest_name}, {room_type.name}, TZS {total_amount:,}. Ref: {booking_ref}"
            background_tasks.add_task(send_sms, owner.phone, sms_msg)
    
    return BookingResponse(
        id=booking.id,
        booking_ref=booking.booking_ref,
        hotel_id=booking.hotel_id,
        hotel_name=hotel.name,
        room_type_id=booking.room_type_id,
        room_type_name=room_type.name,
        guest_name=booking.guest_name,
        guest_phone=booking.guest_phone,
        guest_email=booking.guest_email,
        checkin_date=booking.checkin_date,
        checkout_date=booking.checkout_date,
        nights=booking.nights,
        num_guests=booking.num_guests,
        total_amount_tzs=booking.total_amount_tzs,
        booking_type=booking.booking_type,
        payment_method=booking.payment_method,
        payment_status=booking.payment_status,
        payment_reference=booking.payment_reference,
        status=booking.status,
        checkin_status=booking.checkin_status,
        actual_checkin_time=None,
        actual_checkout_time=None,
        created_by=booking.created_by,
        created_by_name=None,
        notes=booking.notes,
        created_at=booking.created_at.isoformat(),
        expires_at=booking.expires_at.isoformat() if booking.expires_at else None
    )

@api_router.post("/bookings/{booking_id}/confirm-payment")
async def confirm_payment(
    booking_id: str,
    payment: PaymentConfirmation,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session)
):
    booking = await crud.get_booking_by_id(session, booking_id)
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    if booking.payment_status == "paid":
        raise HTTPException(status_code=400, detail="Booking tayari imelipwa")
    
    await crud.update_booking(session, booking_id, {
        "payment_status": "paid",
        "payment_reference": payment.payment_reference,
        "status": "confirmed"
    })
    
    hotel = await crud.get_hotel_by_id(session, booking.hotel_id)
    sms_msg = f"HABARI STAYS: Asante! Booking yako imethibitishwa. Hotel: {hotel.name}, Ref: {booking.booking_ref}. Karibu sana!"
    background_tasks.add_task(send_sms, booking.guest_phone, sms_msg)
    
    return {"message": "Malipo yamethibitishwa", "status": "confirmed"}

@api_router.get("/bookings")
async def get_bookings(
    hotel_id: Optional[str] = None,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] == "traveler":
        # Travelers see their own bookings
        bookings = await crud.get_bookings(session, guest_phone=current_user.get("phone"))
    elif current_user["role"] == "owner":
        owner_hotels = await crud.get_hotels_by_owner(session, current_user["id"])
        hotel_ids = [h.id for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            bookings = await crud.get_bookings(session, hotel_id=hotel_id)
        else:
            bookings = await crud.get_bookings(session, hotel_ids=hotel_ids)
    elif current_user["role"] == "cashier":
        bookings = await crud.get_bookings(session, hotel_id=current_user.get("assigned_hotel_id"))
    else:
        bookings = await crud.get_bookings(session, hotel_id=hotel_id)
    
    result = []
    for b in bookings:
        hotel = await crud.get_hotel_by_id(session, b.hotel_id)
        room_type = await crud.get_room_type_by_id(session, b.room_type_id) if b.room_type_id else None
        creator = await crud.get_user_by_id(session, b.created_by) if b.created_by else None
        
        result.append(BookingResponse(
            id=b.id,
            booking_ref=b.booking_ref,
            hotel_id=b.hotel_id,
            hotel_name=hotel.name if hotel else "Unknown",
            room_type_id=b.room_type_id or "",
            room_type_name=room_type.name if room_type else "Unknown",
            guest_name=b.guest_name,
            guest_phone=b.guest_phone,
            guest_email=b.guest_email,
            checkin_date=b.checkin_date,
            checkout_date=b.checkout_date,
            nights=b.nights,
            num_guests=b.num_guests,
            total_amount_tzs=b.total_amount_tzs,
            booking_type=b.booking_type,
            payment_method=b.payment_method,
            payment_status=b.payment_status,
            payment_reference=b.payment_reference,
            status=b.status,
            checkin_status=b.checkin_status,
            actual_checkin_time=b.actual_checkin_time.isoformat() if b.actual_checkin_time else None,
            actual_checkout_time=b.actual_checkout_time.isoformat() if b.actual_checkout_time else None,
            created_by=b.created_by,
            created_by_name=creator.full_name if creator else None,
            notes=b.notes,
            created_at=b.created_at.isoformat(),
            expires_at=b.expires_at.isoformat() if b.expires_at else None
        ))
    
    return result

@api_router.get("/bookings/recent")
async def get_recent_bookings(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] == "owner":
        owner_hotels = await crud.get_hotels_by_owner(session, current_user["id"])
        hotel_ids = [h.id for h in owner_hotels]
        bookings = await crud.get_bookings(session, hotel_ids=hotel_ids, limit=10)
    else:
        bookings = await crud.get_bookings(session, limit=10)
    
    return [b.to_dict() for b in bookings]

@api_router.get("/bookings/{booking_id}")
async def get_booking(booking_id: str, session: AsyncSession = Depends(get_db_session)):
    booking = await crud.get_booking_by_id(session, booking_id)
    if not booking:
        booking = await crud.get_booking_by_ref(session, booking_id)
    
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    hotel = await crud.get_hotel_by_id(session, booking.hotel_id)
    room_type = await crud.get_room_type_by_id(session, booking.room_type_id) if booking.room_type_id else None
    creator = await crud.get_user_by_id(session, booking.created_by) if booking.created_by else None
    
    return BookingResponse(
        id=booking.id,
        booking_ref=booking.booking_ref,
        hotel_id=booking.hotel_id,
        hotel_name=hotel.name if hotel else "Unknown",
        room_type_id=booking.room_type_id or "",
        room_type_name=room_type.name if room_type else "Unknown",
        guest_name=booking.guest_name,
        guest_phone=booking.guest_phone,
        guest_email=booking.guest_email,
        checkin_date=booking.checkin_date,
        checkout_date=booking.checkout_date,
        nights=booking.nights,
        num_guests=booking.num_guests,
        total_amount_tzs=booking.total_amount_tzs,
        booking_type=booking.booking_type,
        payment_method=booking.payment_method,
        payment_status=booking.payment_status,
        payment_reference=booking.payment_reference,
        status=booking.status,
        checkin_status=booking.checkin_status,
        actual_checkin_time=booking.actual_checkin_time.isoformat() if booking.actual_checkin_time else None,
        actual_checkout_time=booking.actual_checkout_time.isoformat() if booking.actual_checkout_time else None,
        created_by=booking.created_by,
        created_by_name=creator.full_name if creator else None,
        notes=booking.notes,
        created_at=booking.created_at.isoformat(),
        expires_at=booking.expires_at.isoformat() if booking.expires_at else None
    )

@api_router.post("/bookings/{booking_id}/checkin")
async def checkin(
    booking_id: str,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    booking = await crud.get_booking_by_id(session, booking_id)
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    if booking.payment_status != "paid":
        raise HTTPException(status_code=400, detail="Booking haijalipiwa")
    
    now = datetime.now(timezone.utc)
    await crud.update_booking(session, booking_id, {
        "checkin_status": "checked_in",
        "actual_checkin_time": now,
        "status": "confirmed"
    })
    
    hotel = await crud.get_hotel_by_id(session, booking.hotel_id)
    owner = await crud.get_user_by_id(session, hotel.owner_id) if hotel and hotel.owner_id else None
    
    if owner:
        sms_msg = f"HABARI STAYS: {booking.guest_name} amewasili. Ref: {booking.booking_ref}."
        background_tasks.add_task(send_sms, owner.phone, sms_msg)
    
    return {"message": f"{booking.guest_name} amethibitishwa"}

@api_router.post("/bookings/{booking_id}/checkout")
async def checkout(
    booking_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    booking = await crud.get_booking_by_id(session, booking_id)
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    if booking.checkin_status != "checked_in":
        raise HTTPException(status_code=400, detail="Mgeni hajaingizwa")
    
    now = datetime.now(timezone.utc)
    await crud.update_booking(session, booking_id, {
        "checkin_status": "checked_out",
        "actual_checkout_time": now
    })
    
    # Restore room availability
    await crud.increment_room_availability(session, booking.room_type_id, 1)
    
    return {"message": f"{booking.guest_name} ameondoka"}

# ===================== REVIEW ENDPOINTS =====================

@api_router.post("/reviews", response_model=ReviewResponse)
async def create_review(
    review_data: ReviewCreate,
    session: AsyncSession = Depends(get_db_session)
):
    booking = await crud.get_booking_by_id(session, review_data.booking_id)
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    
    existing_review = await crud.get_review_by_booking(session, review_data.booking_id)
    if existing_review:
        raise HTTPException(status_code=400, detail="Review tayari ipo")
    
    review = await crud.create_review(session, {
        "id": generate_uuid(),
        "booking_id": review_data.booking_id,
        "hotel_id": review_data.hotel_id,
        "guest_name": booking.guest_name,
        "guest_phone": booking.guest_phone,
        "rating": review_data.rating,
        "comment": review_data.comment,
    })
    
    return ReviewResponse(**review.to_dict())

@api_router.get("/hotels/{hotel_id}/reviews")
async def get_hotel_reviews(hotel_id: str, session: AsyncSession = Depends(get_db_session)):
    reviews = await crud.get_reviews_by_hotel(session, hotel_id)
    return [r.to_dict() for r in reviews]


@api_router.get("/reviews")
async def list_reviews(
    hotel_id: Optional[str] = None,
    session: AsyncSession = Depends(get_db_session)
):
    if not hotel_id:
        raise HTTPException(status_code=400, detail="hotel_id inahitajika")
    reviews = await crud.get_reviews_by_hotel(session, hotel_id)
    return [r.to_dict() for r in reviews]

# ===================== CASHIER ENDPOINTS =====================

@api_router.post("/cashiers", response_model=CashierResponse)
async def create_cashier(
    cashier_data: CashierCreate,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    hotel = await crud.get_hotel_by_id(session, cashier_data.hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] == "owner" and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hii si hotel yako")
    
    existing = await crud.get_user_by_email(session, cashier_data.email)
    if existing:
        raise HTTPException(status_code=400, detail="Email tayari imetumika")
    
    temp_password = generate_temp_password()
    cashier_id = generate_uuid()
    
    cashier = await crud.create_user(session, {
        "id": cashier_id,
        "email": cashier_data.email,
        "full_name": cashier_data.full_name,
        "phone": format_phone(cashier_data.phone),
        "role": "cashier",
        "assigned_hotel_id": cashier_data.hotel_id,
        "password_hash": get_password_hash(temp_password),
        "is_active": True,
        "is_verified": True,
    })
    
    # Create assignment record
    await crud.create_cashier_assignment(session, {
        "id": generate_uuid(),
        "cashier_id": cashier_id,
        "hotel_id": cashier_data.hotel_id,
        "assigned_by": current_user["id"],
        "is_active": True,
    })
    
    # Send SMS
    sms_msg = f"Habari {cashier_data.full_name}! Umewekwa kama Mweka Hazina wa {hotel.name} kwenye Habari Stays. Ingia: habaristays.com/login Email: {cashier_data.email} Neno la siri: {temp_password}"
    background_tasks.add_task(send_sms, cashier_data.phone, sms_msg)
    
    return CashierResponse(
        id=cashier.id,
        full_name=cashier.full_name,
        phone=cashier.phone,
        email=cashier.email,
        role=cashier.role,
        assigned_hotel_id=cashier.assigned_hotel_id,
        assigned_hotel_name=hotel.name,
        is_active=cashier.is_active,
        created_at=cashier.created_at.isoformat(),
        last_login=None
    )

@api_router.get("/cashiers")
async def get_cashiers(
    hotel_id: Optional[str] = None,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    if current_user["role"] == "owner":
        owner_hotels = await crud.get_hotels_by_owner(session, current_user["id"])
        hotel_ids = [h.id for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            cashiers = await crud.get_users_by_role(session, "cashier", [hotel_id])
        else:
            cashiers = await crud.get_users_by_role(session, "cashier", hotel_ids)
    elif hotel_id:
        cashiers = await crud.get_users_by_role(session, "cashier", [hotel_id])
    else:
        cashiers = await crud.get_users_by_role(session, "cashier")
    
    result = []
    for c in cashiers:
        hotel = await crud.get_hotel_by_id(session, c.assigned_hotel_id) if c.assigned_hotel_id else None
        result.append(CashierResponse(
            id=c.id,
            full_name=c.full_name,
            phone=c.phone,
            email=c.email,
            role=c.role,
            assigned_hotel_id=c.assigned_hotel_id,
            assigned_hotel_name=hotel.name if hotel else "Unknown",
            is_active=c.is_active,
            created_at=c.created_at.isoformat(),
            last_login=c.last_login.isoformat() if c.last_login else None
        ))
    
    return result

@api_router.put("/cashiers/{cashier_id}/toggle-status")
async def toggle_cashier_status(
    cashier_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    cashier = await crud.get_user_by_id(session, cashier_id)
    if not cashier or cashier.role != "cashier":
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    
    if current_user["role"] == "owner":
        hotel = await crud.get_hotel_by_id(session, cashier.assigned_hotel_id)
        if not hotel or hotel.owner_id != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    new_status = not cashier.is_active
    await crud.update_user(session, cashier_id, {"is_active": new_status})
    await crud.update_cashier_assignments(session, cashier_id, is_active=new_status)
    
    return {"message": "Hali imebadilishwa", "is_active": new_status}

@api_router.post("/cashiers/{cashier_id}/reset-password")
async def reset_cashier_password(
    cashier_id: str,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    cashier = await crud.get_user_by_id(session, cashier_id)
    if not cashier or cashier.role != "cashier":
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    
    if current_user["role"] == "owner":
        hotel = await crud.get_hotel_by_id(session, cashier.assigned_hotel_id)
        if not hotel or hotel.owner_id != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    new_password = generate_temp_password()
    await crud.update_user(session, cashier_id, {"password_hash": get_password_hash(new_password)})
    
    sms_msg = f"Habari {cashier.full_name}! Neno lako la siri la Habari Stays limewekwa upya. Neno jipya: {new_password} Ingia: habaristays.com/login"
    background_tasks.add_task(send_sms, cashier.phone, sms_msg)
    
    return {"message": "Neno la siri limewekwa upya", "new_password": new_password}

@api_router.patch("/cashiers/{cashier_id}")
async def edit_cashier(
    cashier_id: str,
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    cashier = await crud.get_user_by_id(session, cashier_id)
    if not cashier or cashier.role != "cashier":
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    
    if current_user["role"] == "owner":
        hotel = await crud.get_hotel_by_id(session, cashier.assigned_hotel_id)
        if not hotel or hotel.owner_id != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    body = await request.json()
    update = {}
    
    if "full_name" in body:
        update["full_name"] = body["full_name"]
    if "phone" in body:
        update["phone"] = format_phone(body["phone"])
    if "assigned_hotel_id" in body:
        new_hotel = await crud.get_hotel_by_id(session, body["assigned_hotel_id"])
        if not new_hotel:
            raise HTTPException(status_code=404, detail="Hotel haipatikani")
        if current_user["role"] == "owner" and new_hotel.owner_id != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hii si hotel yako")
        update["assigned_hotel_id"] = body["assigned_hotel_id"]
        
        # Update assignments
        await crud.update_cashier_assignments(session, cashier_id, deactivate_all=True)
        await crud.create_cashier_assignment(session, {
            "id": generate_uuid(),
            "cashier_id": cashier_id,
            "hotel_id": body["assigned_hotel_id"],
            "assigned_by": current_user["id"],
            "is_active": True,
        })
    
    if not update:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    await crud.update_user(session, cashier_id, update)
    return {"message": "Cashier amesasishwa"}

# ===================== CASHIER SHIFT ENDPOINTS =====================

@api_router.get("/cashier/todays-checkins")
async def cashier_todays_checkins(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    hotel_id = current_user["assigned_hotel_id"]
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    bookings = await crud.get_todays_checkins(session, hotel_id, today)
    
    result = []
    for b in bookings:
        room = await crud.get_room_type_by_id(session, b.room_type_id) if b.room_type_id else None
        result.append({
            "id": b.id,
            "booking_ref": b.booking_ref,
            "guest_name": b.guest_name,
            "guest_phone": b.guest_phone,
            "room_type_name": room.name if room else "Unknown",
            "nights": b.nights,
            "total_amount_tzs": b.total_amount_tzs,
            "payment_status": b.payment_status,
            "checkin_status": b.checkin_status,
            "checkin_date": b.checkin_date,
            "checkout_date": b.checkout_date,
            "actual_checkin_time": b.actual_checkin_time.isoformat() if b.actual_checkin_time else None,
            "actual_checkout_time": b.actual_checkout_time.isoformat() if b.actual_checkout_time else None,
            "status": b.status,
        })
    
    return result

@api_router.get("/cashier/active-guests")
async def cashier_active_guests(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    hotel_id = current_user["assigned_hotel_id"]
    bookings = await crud.get_active_guests(session, hotel_id)
    
    result = []
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    
    for b in bookings:
        room = await crud.get_room_type_by_id(session, b.room_type_id) if b.room_type_id else None
        try:
            nights_left = max(0, (datetime.fromisoformat(b.checkout_date) - datetime.fromisoformat(today)).days)
        except:
            nights_left = 0
        
        result.append({
            "id": b.id,
            "booking_ref": b.booking_ref,
            "guest_name": b.guest_name,
            "guest_phone": b.guest_phone,
            "room_type_name": room.name if room else "Unknown",
            "checkin_date": b.checkin_date,
            "checkout_date": b.checkout_date,
            "actual_checkin_time": b.actual_checkin_time.isoformat() if b.actual_checkin_time else None,
            "nights_left": nights_left,
            "booking_type": b.booking_type,
            "total_amount_tzs": b.total_amount_tzs,
            "payment_method": b.payment_method,
        })
    
    return result

@api_router.post("/cashier/confirm-checkin/{booking_id}")
async def cashier_confirm_checkin(
    booking_id: str,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    booking = await crud.get_booking_by_id(session, booking_id)
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    if booking.hotel_id != current_user["assigned_hotel_id"]:
        raise HTTPException(status_code=403, detail="Booking hii ni kwa hoteli nyingine")
    if booking.payment_status != "paid":
        raise HTTPException(status_code=400, detail="Booking hii haijalipiwa")
    if booking.checkin_status != "not_checked_in":
        raise HTTPException(status_code=400, detail="Mgeni tayari amefika au ameondoka")
    
    now = datetime.now(timezone.utc)
    await crud.update_booking(session, booking_id, {
        "checkin_status": "checked_in",
        "actual_checkin_time": now,
        "status": "confirmed"
    })
    
    room = await crud.get_room_type_by_id(session, booking.room_type_id) if booking.room_type_id else None
    await log_cashier_activity(
        session, current_user["id"], booking.hotel_id, booking_id, "checkin_confirmed",
        guest_name=booking.guest_name, room_type_name=room.name if room else ""
    )
    
    hotel = await crud.get_hotel_by_id(session, booking.hotel_id)
    owner = await crud.get_user_by_id(session, hotel.owner_id) if hotel and hotel.owner_id else None
    
    if owner:
        sms_msg = f"HABARI STAYS: {booking.guest_name} amewasili. Ref: {booking.booking_ref}. Mweka Hazina: {current_user['full_name']}."
        background_tasks.add_task(send_sms, owner.phone, sms_msg)
    
    return {"message": f"{booking.guest_name} amethibitishwa", "booking_id": booking_id}

@api_router.post("/cashier/confirm-checkout/{booking_id}")
async def cashier_confirm_checkout(
    booking_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    booking = await crud.get_booking_by_id(session, booking_id)
    if not booking:
        raise HTTPException(status_code=404, detail="Booking haipatikani")
    if booking.hotel_id != current_user["assigned_hotel_id"]:
        raise HTTPException(status_code=403, detail="Booking hii ni kwa hoteli nyingine")
    if booking.checkin_status != "checked_in":
        raise HTTPException(status_code=400, detail="Mgeni hajaingizwa au tayari ameondoka")
    
    now = datetime.now(timezone.utc)
    await crud.update_booking(session, booking_id, {
        "checkin_status": "checked_out",
        "actual_checkout_time": now
    })
    
    await crud.increment_room_availability(session, booking.room_type_id, 1)
    
    room = await crud.get_room_type_by_id(session, booking.room_type_id) if booking.room_type_id else None
    await log_cashier_activity(
        session, current_user["id"], booking.hotel_id, booking_id, "checkout_processed",
        guest_name=booking.guest_name, room_type_name=room.name if room else ""
    )
    
    return {"message": f"{booking.guest_name} ameondoka. Chumba kimerudishwa.", "booking_id": booking_id}

@api_router.post("/cashier/walkin")
async def cashier_walkin(
    request: Request,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    body = await request.json()
    hotel_id = current_user["assigned_hotel_id"]
    
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    room_type_id = body.get("room_type_id")
    room = await crud.get_room_type_by_id(session, room_type_id)
    if not room:
        raise HTTPException(status_code=404, detail="Aina ya chumba haipatikani")
    if room.available_rooms < 1:
        raise HTTPException(status_code=400, detail="Hakuna vyumba vilivyopo")
    
    guest_name = body.get("guest_name", "").strip()
    guest_phone = format_phone(body.get("guest_phone", ""))
    checkin_date = body.get("checkin_date", datetime.now(timezone.utc).strftime("%Y-%m-%d"))
    checkout_date = body.get("checkout_date")
    
    if not guest_name or not guest_phone or not checkout_date:
        raise HTTPException(status_code=400, detail="Taarifa za mgeni zinahitajika")
    
    checkin_dt = datetime.fromisoformat(checkin_date)
    checkout_dt = datetime.fromisoformat(checkout_date)
    nights = (checkout_dt - checkin_dt).days
    
    if nights < 1:
        raise HTTPException(status_code=400, detail="Tarehe si sahihi")
    
    total_amount = nights * room.price_per_night
    payment_method = body.get("payment_method", "cash")
    notes = body.get("notes", "")
    
    booking_id = generate_uuid()
    booking_ref = generate_booking_ref()
    now = datetime.now(timezone.utc)
    
    booking = await crud.create_booking(session, {
        "id": booking_id,
        "booking_ref": booking_ref,
        "hotel_id": hotel_id,
        "room_type_id": room_type_id,
        "guest_name": guest_name,
        "guest_phone": guest_phone,
        "guest_email": body.get("guest_email"),
        "checkin_date": checkin_date,
        "checkout_date": checkout_date,
        "nights": nights,
        "num_guests": body.get("num_guests", 1),
        "total_amount_tzs": total_amount,
        "booking_type": "walkin",
        "payment_method": payment_method,
        "payment_status": "paid",
        "status": "confirmed",
        "checkin_status": "checked_in",
        "actual_checkin_time": now,
        "created_by": current_user["id"],
        "notes": notes,
    })
    
    await crud.increment_room_availability(session, room_type_id, -1)
    
    await log_cashier_activity(
        session, current_user["id"], hotel_id, booking_id, "walkin_recorded",
        payment_method=payment_method, amount_tzs=total_amount,
        guest_name=guest_name, room_type_name=room.name
    )
    
    owner = await crud.get_user_by_id(session, hotel.owner_id) if hotel.owner_id else None
    if owner:
        sms_msg = f"HABARI STAYS: Walk-in Mpya! Mgeni: {guest_name}, Chumba: {room.name}, Usiku: {nights}, TZS {total_amount:,}. Mweka Hazina: {current_user['full_name']}. Ref: {booking_ref}"
        background_tasks.add_task(send_sms, owner.phone, sms_msg)
    
    return {
        "id": booking_id,
        "booking_ref": booking_ref,
        "hotel_name": hotel.name,
        "hotel_address": hotel.address,
        "guest_name": guest_name,
        "guest_phone": guest_phone,
        "room_type_name": room.name,
        "checkin_date": checkin_date,
        "checkout_date": checkout_date,
        "nights": nights,
        "total_amount_tzs": total_amount,
        "payment_method": payment_method,
        "cashier_name": current_user["full_name"],
        "created_at": now.isoformat()
    }

@api_router.get("/cashier/activity-log")
async def cashier_activity_log_today(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    start_of_day = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    logs = await crud.get_activity_logs(session, cashier_id=current_user["id"], after_timestamp=start_of_day)
    
    return [l.to_dict() for l in logs]

@api_router.get("/cashier/shift-summary", response_model=ShiftSummary)
async def get_shift_summary(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Hii ni kwa cashier tu")
    
    today = datetime.now().strftime("%Y-%m-%d")
    start_of_day = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0).isoformat()
    
    walkin_bookings = await crud.get_bookings(
        session,
        hotel_id=current_user["assigned_hotel_id"],
        booking_type="walkin",
        created_by=current_user["id"],
        created_after=start_of_day
    )
    
    walkins_count = len(walkin_bookings)
    walkins_total = sum(b.total_amount_tzs for b in walkin_bookings)
    cash_collected = sum(b.total_amount_tzs for b in walkin_bookings if b.payment_method == "cash")
    
    online_checkins = await crud.count_bookings(
        session,
        hotel_id=current_user["assigned_hotel_id"],
        booking_type="online",
        checkin_status=["checked_in", "checked_out"],
        created_after=start_of_day
    )
    
    checkouts = await crud.count_bookings(
        session,
        hotel_id=current_user["assigned_hotel_id"],
        checkin_status="checked_out",
        created_after=start_of_day
    )
    
    return ShiftSummary(
        cashier_id=current_user["id"],
        cashier_name=current_user["full_name"],
        date=today,
        walkins_count=walkins_count,
        walkins_total=walkins_total,
        online_checkins_count=online_checkins,
        checkouts_count=checkouts,
        cash_collected=cash_collected
    )

@api_router.get("/cashier/shift-summary-full")
async def cashier_shift_summary_full(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    hotel_id = current_user["assigned_hotel_id"]
    start_of_day = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    
    walkin_bookings = await crud.get_bookings(
        session,
        hotel_id=hotel_id,
        booking_type="walkin",
        created_by=current_user["id"],
        created_after=start_of_day.isoformat()
    )
    
    walkins_count = len(walkin_bookings)
    cash_total = sum(b.total_amount_tzs for b in walkin_bookings if b.payment_method == "cash")
    mpesa_total = sum(b.total_amount_tzs for b in walkin_bookings if b.payment_method == "mpesa")
    card_total = sum(b.total_amount_tzs for b in walkin_bookings if b.payment_method == "card")
    walkins_total = cash_total + mpesa_total + card_total
    
    online_checkins = await crud.count_bookings(
        session,
        hotel_id=hotel_id,
        booking_type="online",
        checkin_status=["checked_in", "checked_out"],
        created_after=start_of_day.isoformat()
    )
    
    checkouts = await crud.count_bookings(
        session,
        hotel_id=hotel_id,
        checkin_status="checked_out",
        created_after=start_of_day.isoformat()
    )
    
    activity_logs = await crud.get_activity_logs(
        session,
        cashier_id=current_user["id"],
        after_timestamp=start_of_day
    )
    
    return {
        "cashier_name": current_user["full_name"],
        "hotel_name": hotel.name if hotel else "",
        "hotel_address": hotel.address if hotel else "",
        "date": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        "walkins_count": walkins_count,
        "walkins_total": walkins_total,
        "online_checkins_count": online_checkins,
        "checkouts_count": checkouts,
        "cash_total": cash_total,
        "mpesa_total": mpesa_total,
        "card_total": card_total,
        "total_revenue": walkins_total,
        "activity_log": [l.to_dict() for l in activity_logs]
    }

@api_router.post("/cashier/verify-booking")
async def cashier_verify_booking(
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "cashier":
        raise HTTPException(status_code=403, detail="Cashier tu")
    
    body = await request.json()
    search_term = body.get("search", "").strip()
    
    if not search_term:
        raise HTTPException(status_code=400, detail="Tafadhali ingiza namba ya booking, simu, au jina")
    
    hotel_id = current_user["assigned_hotel_id"]
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    
    booking = await crud.search_bookings(session, search_term, hotel_id)
    
    if not booking:
        return {"found": False, "error": "not_found", "message": "Booking haikupatikana. Angalia namba na ujaribu tena."}
    
    room = await crud.get_room_type_by_id(session, booking.room_type_id) if booking.room_type_id else None
    hotel = await crud.get_hotel_by_id(session, booking.hotel_id)
    
    result = {
        "found": True,
        "id": booking.id,
        "booking_ref": booking.booking_ref,
        "guest_name": booking.guest_name,
        "guest_phone": booking.guest_phone,
        "room_type_name": room.name if room else "Unknown",
        "checkin_date": booking.checkin_date,
        "checkout_date": booking.checkout_date,
        "nights": booking.nights,
        "total_amount_tzs": booking.total_amount_tzs,
        "payment_status": booking.payment_status,
        "payment_method": booking.payment_method,
        "booking_type": booking.booking_type,
        "status": booking.status,
        "checkin_status": booking.checkin_status,
        "actual_checkin_time": booking.actual_checkin_time.isoformat() if booking.actual_checkin_time else None,
        "hotel_name": hotel.name if hotel else "Unknown",
        "hotel_id": booking.hotel_id,
        "can_checkin": False,
        "errors": []
    }
    
    if booking.hotel_id != hotel_id:
        result["errors"].append({"type": "wrong_hotel", "message": f"Booking hii ni kwa hoteli nyingine: {hotel.name if hotel else 'Unknown'}."})
    elif booking.status == "expired":
        result["errors"].append({"type": "expired", "message": "Booking hii imeisha muda na haijalipiwa."})
    elif booking.payment_status != "paid":
        result["errors"].append({"type": "unpaid", "message": "Booking hii haijalipiwa. Mwambie mgeni alipe kwanza."})
    elif booking.checkin_status == "checked_in":
        checkin_time = booking.actual_checkin_time.isoformat()[:16] if booking.actual_checkin_time else ""
        result["errors"].append({"type": "already_checked_in", "message": f"Mgeni huyu tayari amethibitishwa saa {checkin_time}."})
    elif booking.checkin_status == "checked_out":
        result["errors"].append({"type": "checked_out", "message": "Mgeni huyu tayari ameondoka."})
    elif booking.checkin_date != today:
        result["errors"].append({"type": "wrong_date", "message": f"Check-in ya booking hii ni {booking.checkin_date}, si leo."})
    else:
        result["can_checkin"] = True
    
    return result

@api_router.put("/cashiers/{cashier_id}/toggle-status-sms")
async def toggle_cashier_status_sms(
    cashier_id: str,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    cashier = await crud.get_user_by_id(session, cashier_id)
    if not cashier or cashier.role != "cashier":
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    
    if current_user["role"] == "owner":
        hotel = await crud.get_hotel_by_id(session, cashier.assigned_hotel_id)
        if not hotel or hotel.owner_id != current_user["id"]:
            raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    new_status = not cashier.is_active
    await crud.update_user(session, cashier_id, {"is_active": new_status})
    await crud.update_cashier_assignments(session, cashier_id, is_active=new_status)
    
    if new_status:
        sms_msg = f"Habari {cashier.full_name}! Akaunti yako ya Habari Stays imewashwa tena. Unaweza kuingia sasa."
    else:
        sms_msg = f"Habari {cashier.full_name}. Akaunti yako ya Habari Stays imezimwa. Wasiliana na msimamizi wako."
    
    background_tasks.add_task(send_sms, cashier.phone, sms_msg)
    
    return {"message": "Hali imebadilishwa", "is_active": new_status}

@api_router.get("/cashiers/{cashier_id}/performance")
async def cashier_performance_detail(
    cashier_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    cashier = await crud.get_user_by_id(session, cashier_id)
    if not cashier or cashier.role != "cashier":
        raise HTTPException(status_code=404, detail="Cashier hapatikani")
    
    hotel = await crud.get_hotel_by_id(session, cashier.assigned_hotel_id) if cashier.assigned_hotel_id else None
    
    start_of_month = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    logs = await crud.get_activity_logs(session, cashier_id=cashier_id, after_timestamp=start_of_month, limit=500)
    
    walkins = [l for l in logs if l.action_type == "walkin_recorded"]
    checkins = [l for l in logs if l.action_type == "checkin_confirmed"]
    checkouts = [l for l in logs if l.action_type == "checkout_processed"]
    
    cashier_dict = cashier.to_dict()
    cashier_dict["assigned_hotel_name"] = hotel.name if hotel else "Unknown"
    
    return {
        "cashier": cashier_dict,
        "this_month": {
            "walkins_recorded": len(walkins),
            "checkins_confirmed": len(checkins),
            "checkouts_processed": len(checkouts),
            "total_revenue": sum(l.amount_tzs for l in walkins),
        },
        "activity_log": [l.to_dict() for l in logs[:100]]
    }

# ===================== ANALYTICS ENDPOINTS =====================

@api_router.get("/analytics/revenue")
async def get_revenue_analytics(
    hotel_id: Optional[str] = None,
    period: str = "month",
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    now = datetime.now()
    if period == "week":
        start_date = (now - timedelta(days=7)).isoformat()
    elif period == "month":
        start_date = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    else:
        start_date = now.replace(month=1, day=1).isoformat()
    
    hotel_ids = None
    if current_user["role"] == "owner":
        owner_hotels = await crud.get_hotels_by_owner(session, current_user["id"])
        hotel_ids = [h.id for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            hotel_ids = [hotel_id]
    elif hotel_id:
        hotel_ids = [hotel_id]
    
    bookings = await crud.get_bookings(
        session,
        hotel_ids=hotel_ids,
        created_after=start_date
    )
    bookings = [b for b in bookings if b.payment_status == "paid"]
    
    total_revenue = sum(b.total_amount_tzs for b in bookings)
    online_revenue = sum(b.total_amount_tzs for b in bookings if b.booking_type == "online")
    walkin_revenue = sum(b.total_amount_tzs for b in bookings if b.booking_type == "walkin")
    
    daily_revenue = {}
    for b in bookings:
        date = b.created_at.strftime("%Y-%m-%d") if b.created_at else "Unknown"
        if date not in daily_revenue:
            daily_revenue[date] = {"date": date, "online": 0, "walkin": 0, "total": 0}
        amount = b.total_amount_tzs
        daily_revenue[date]["total"] += amount
        if b.booking_type == "online":
            daily_revenue[date]["online"] += amount
        else:
            daily_revenue[date]["walkin"] += amount
    
    return {
        "total_revenue": total_revenue,
        "online_revenue": online_revenue,
        "walkin_revenue": walkin_revenue,
        "bookings_count": len(bookings),
        "daily_breakdown": sorted(daily_revenue.values(), key=lambda x: x["date"])
    }

@api_router.get("/analytics/occupancy")
async def get_occupancy(
    hotel_id: Optional[str] = None,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] not in ["owner", "admin"]:
        raise HTTPException(status_code=403, detail="Hauruhusiwi")
    
    hotel_ids = None
    if current_user["role"] == "owner":
        owner_hotels = await crud.get_hotels_by_owner(session, current_user["id"])
        hotel_ids = [h.id for h in owner_hotels]
        if hotel_id and hotel_id in hotel_ids:
            hotel_ids = [hotel_id]
    elif hotel_id:
        hotel_ids = [hotel_id]
    
    rooms = await crud.get_room_types_with_availability(session, hotel_ids)
    
    total = sum(r.total_rooms for r in rooms)
    available = sum(r.available_rooms for r in rooms)
    occupied = total - available
    
    return {
        "total_rooms": total,
        "occupied_rooms": occupied,
        "available_rooms": available,
        "occupancy_rate": round((occupied / total * 100) if total > 0 else 0, 1)
    }

# ===================== ADMIN ENDPOINTS =====================

@api_router.get("/admin/stats")
async def get_admin_stats(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    total_hotels = await crud.count_hotels(session)
    verified_hotels = await crud.count_hotels(session, status="verified")
    pending_hotels = await crud.count_hotels(session, status="pending")
    
    total_bookings = await crud.count_bookings(session)
    online_bookings = await crud.count_bookings(session, booking_type="online")
    walkin_bookings = await crud.count_bookings(session, booking_type="walkin")
    
    total_cashiers = await crud.count_users_by_role(session, "cashier")
    total_owners = await crud.count_users_by_role(session, "owner")
    pending_owners = await crud.count_users_by_role(session, "owner", {"is_verified": False})
    
    total_revenue = await crud.get_revenue_sum(session)
    
    revenue_by_city = await crud.get_revenue_by_city(session)
    
    return {
        "total_hotels": total_hotels,
        "verified_hotels": verified_hotels,
        "pending_hotels": pending_hotels,
        "total_bookings": total_bookings,
        "online_bookings": online_bookings,
        "walkin_bookings": walkin_bookings,
        "total_cashiers": total_cashiers,
        "total_owners": total_owners,
        "pending_owners": pending_owners,
        "total_revenue": total_revenue,
        "commission_10_percent": int(total_revenue * 0.1),
        "revenue_by_city": revenue_by_city
    }

@api_router.put("/admin/hotels/{hotel_id}/verify")
async def verify_hotel(
    hotel_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    await crud.update_hotel(session, hotel_id, {"status": "verified"})
    return {"message": "Hotel imethibitishwa"}

@api_router.put("/admin/hotels/{hotel_id}/suspend")
async def suspend_hotel(
    hotel_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    await crud.update_hotel(session, hotel_id, {"status": "suspended"})
    return {"message": "Hotel imesimamishwa"}

@api_router.put("/admin/users/{user_id}/verify")
async def verify_owner(
    user_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    user = await crud.get_user_by_id(session, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User hapatikani")
    
    await crud.update_user(session, user_id, {"is_verified": True})
    return {"message": "Mmiliki amethibitishwa"}

@api_router.get("/admin/pending-owners")
async def get_pending_owners(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    owners = await crud.get_pending_owners(session)
    return [o.to_dict() for o in owners]

@api_router.get("/admin/all-hotels")
async def get_all_hotels_admin(
    status_filter: Optional[str] = None,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotels = await crud.get_hotels(session, status=status_filter if status_filter else None, limit=500)
    
    result = []
    for hotel in hotels:
        rooms = await crud.get_room_types_by_hotel(session, hotel.id)
        total_rooms = sum(r.total_rooms for r in rooms)
        
        owner = None
        if hotel.owner_id:
            owner = await crud.get_user_by_id(session, hotel.owner_id)
        
        result.append({
            "id": hotel.id,
            "hotel_code": hotel.hotel_code,
            "name": hotel.name,
            "city": hotel.city,
            "address": hotel.address,
            "phone_number": hotel.phone_number,
            "status": hotel.status,
            "data_source": hotel.data_source,
            "owner_id": hotel.owner_id,
            "owner_name": owner.full_name if owner else None,
            "owner_email": owner.email if owner else None,
            "total_rooms": total_rooms,
            "has_rooms": total_rooms > 0,
            "room_type_count": len(rooms),
            "has_default_rooms": any(r.is_default for r in rooms),
            "google_rating": hotel.google_rating,
            "photo_count": len(hotel.photos) if hotel.photos else 0,
            "cover_photo": get_cover_url(hotel.photos, "cloudinary_thumb") or hotel.cover_photo,
            "created_at": hotel.created_at.isoformat()
        })
    
    return result

@api_router.post("/admin/hotels/{hotel_id}/attach-owner")
async def attach_owner_to_hotel(
    hotel_id: str,
    request: Request,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    body = await request.json()
    owner_name = body.get("owner_name", "").strip()
    owner_phone = body.get("owner_phone", "").strip()
    owner_email = body.get("owner_email", "").strip().lower()
    
    if not owner_name or not owner_phone or not owner_email:
        raise HTTPException(status_code=400, detail="Owner name, phone na email vinahitajika")
    
    owner = await crud.get_user_by_email(session, owner_email)
    temp_password = None
    owner_id = None
    phone_formatted = format_phone(owner_phone)
    
    if owner:
        if owner.role != "owner":
            raise HTTPException(status_code=400, detail="User yupo lakini si mmiliki")
        owner_id = owner.id
        owner_name = owner.full_name or owner_name
        phone_formatted = owner.phone or phone_formatted
    else:
        temp_password = generate_temp_password()
        owner = await crud.create_user(session, {
            "email": owner_email,
            "full_name": owner_name,
            "phone": phone_formatted,
            "role": "owner",
            "assigned_hotel_id": None,
            "password_hash": get_password_hash(temp_password),
            "is_active": True,
            "is_verified": True,
        })
        owner_id = owner.id
    
    await crud.update_hotel(session, hotel_id, {"owner_id": owner_id})
    
    if temp_password:
        sms_msg = (
            f"Habari {owner_name}! Umealikwa kusimamia {hotel.name} kwenye Habari Stays. "
            f"Ingia habaristays.com/login Email: {owner_email} Neno la siri: {temp_password}"
        )
        background_tasks.add_task(send_sms, phone_formatted, sms_msg)
    
    return {
        "message": f"Mmiliki amewekwa kwa {hotel.name}",
        "owner_id": owner_id,
        "new_owner": temp_password is not None
    }

@api_router.post("/admin/hotels")
async def admin_create_hotel(
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    body = await request.json()
    name = body.get("name", "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Jina la hotel linahitajika")
    
    existing = await crud.get_hotel_by_name(session, name)
    if existing:
        raise HTTPException(status_code=400, detail="Hotel yenye jina hili ipo tayari")
    
    hotel_id = generate_uuid()
    hotel_code = await generate_unique_hotel_code(session)
    hotel = await crud.create_hotel(session, {
        "id": hotel_id,
        "hotel_code": hotel_code,
        "owner_id": body.get("owner_id"),
        "name": name,
        "description": body.get("description", ""),
        "address": body.get("address", ""),
        "city": body.get("city", ""),
        "phone_number": body.get("phone_number", ""),
        "whatsapp_number": body.get("whatsapp_number", body.get("phone_number", "")),
        "amenities": body.get("amenities", []),
        "google_maps_url": body.get("google_maps_url", ""),
        "google_rating": body.get("google_rating"),
        "google_review_count": body.get("google_review_count"),
        "latitude": body.get("latitude"),
        "longitude": body.get("longitude"),
        "website": body.get("website", ""),
        "status": body.get("status", "verified"),
        "data_source": "admin_created",
    })
    
    return hotel.to_dict()

@api_router.put("/admin/hotels/{hotel_id}")
async def admin_update_hotel(
    hotel_id: str,
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    body = await request.json()
    
    updatable = ["name", "description", "address", "city", "phone_number", "whatsapp_number",
                 "amenities", "google_maps_url", "google_rating", "google_review_count",
                 "latitude", "longitude", "website", "status"]
    
    update_data = {k: v for k, v in body.items() if k in updatable}
    
    if not update_data:
        raise HTTPException(status_code=400, detail="Hakuna mabadiliko")
    
    if "name" in update_data and update_data["name"].lower() != hotel.name.lower():
        existing = await crud.get_hotel_by_name(session, update_data["name"], exclude_id=hotel_id)
        if existing:
            raise HTTPException(status_code=400, detail="Hotel yenye jina hili ipo tayari")
    
    updated = await crud.update_hotel(session, hotel_id, update_data)
    return updated.to_dict()

@api_router.delete("/admin/hotels/{hotel_id}")
async def admin_delete_hotel(
    hotel_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin tu")
    
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    active_bookings = await crud.count_bookings(
        session,
        hotel_id=hotel_id,
        status=["confirmed", "pending"]
    )
    
    if active_bookings > 0:
        raise HTTPException(status_code=400, detail=f"Hotel ina bukini {active_bookings} hai. Haziwezi kufutwa.")
    
    # Delete room types
    await crud.delete_room_types_by_hotel(session, hotel_id)
    
    # Unassign cashiers
    cashiers = await crud.get_users_by_role(session, "cashier", [hotel_id])
    for c in cashiers:
        await crud.update_user(session, c.id, {"assigned_hotel_id": None})
    
    # Delete hotel
    await crud.delete_hotel(session, hotel_id)
    
    return {"message": f"Hotel '{hotel.name}' imefutwa"}

# ===================== PHOTO MANAGEMENT =====================

@api_router.get("/hotels/{hotel_id}/photos")
async def get_hotel_photos(hotel_id: str, session: AsyncSession = Depends(get_db_session)):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    photos = hotel.photos or []
    if photos and isinstance(photos[0], dict):
        photos.sort(key=lambda p: p.get("sort_order", 0))
    
    return photos

@api_router.post("/hotels/{hotel_id}/photos")
async def add_hotel_photo(
    hotel_id: str,
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] not in ["admin"] and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    body = await request.json()
    photos = hotel.photos or []
    
    current_count = len([p for p in photos if isinstance(p, dict)])
    if current_count >= 10:
        raise HTTPException(status_code=400, detail="Kiwango cha juu cha picha 10 kimefikia")
    
    is_first = current_count == 0
    photo_obj = {
        "id": generate_uuid(),
        "is_primary": is_first or body.get("is_primary", False),
        "sort_order": current_count,
        "local_path": body.get("local_path", ""),
        "cloudinary_original": body.get("cloudinary_original", ""),
        "cloudinary_thumb": body.get("cloudinary_thumb", ""),
        "cloudinary_mobile": body.get("cloudinary_mobile", ""),
        "cloudinary_web": body.get("cloudinary_web", ""),
        "cloudinary_hd": body.get("cloudinary_hd", ""),
        "source": body.get("source", "uploaded"),
        "uploaded_at": datetime.now(timezone.utc).isoformat()
    }
    
    if photo_obj["is_primary"]:
        for p in photos:
            if isinstance(p, dict):
                p["is_primary"] = False
    
    photos.append(photo_obj)
    cover = get_cover_url(photos, "cloudinary_web")
    
    await crud.update_hotel(session, hotel_id, {"photos": photos, "cover_photo": cover})
    
    return photo_obj

@api_router.patch("/hotels/{hotel_id}/photos/{photo_id}/set-primary")
async def set_primary_photo(
    hotel_id: str,
    photo_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] not in ["admin"] and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    photos = hotel.photos or []
    found = False
    for p in photos:
        if isinstance(p, dict):
            if p["id"] == photo_id:
                p["is_primary"] = True
                found = True
            else:
                p["is_primary"] = False
    
    if not found:
        raise HTTPException(status_code=404, detail="Picha haipatikani")
    
    cover = get_cover_url(photos, "cloudinary_web")
    await crud.update_hotel(session, hotel_id, {"photos": photos, "cover_photo": cover})
    
    return {"message": "Picha ya kwanza imebadilishwa"}

@api_router.delete("/hotels/{hotel_id}/photos/{photo_id}")
async def delete_hotel_photo(
    hotel_id: str,
    photo_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    hotel = await crud.get_hotel_by_id(session, hotel_id)
    if not hotel:
        raise HTTPException(status_code=404, detail="Hotel haipatikani")
    
    if current_user["role"] not in ["admin"] and hotel.owner_id != current_user["id"]:
        raise HTTPException(status_code=403, detail="Hauna ruhusa")
    
    photos = hotel.photos or []
    photo_dicts = [p for p in photos if isinstance(p, dict)]
    
    if len(photo_dicts) <= 1:
        raise HTTPException(status_code=400, detail="Lazima kuwe na picha moja angalau")
    
    target = next((p for p in photo_dicts if p["id"] == photo_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="Picha haipatikani")
    
    was_primary = target.get("is_primary", False)
    new_photos = [p for p in photos if not (isinstance(p, dict) and p["id"] == photo_id)]
    
    if was_primary and new_photos:
        for p in new_photos:
            if isinstance(p, dict):
                p["is_primary"] = True
                break
    
    for idx, p in enumerate(new_photos):
        if isinstance(p, dict):
            p["sort_order"] = idx
    
    cover = get_cover_url(new_photos, "cloudinary_web")
    await crud.update_hotel(session, hotel_id, {"photos": new_photos, "cover_photo": cover})
    
    return {"message": "Picha imefutwa"}

# ===================== IMPORT ENDPOINTS =====================

def parse_cloudinary_photos_from_row(row, headers):
    """Parse P1/P2/P3 photo columns from an Excel row into structured photo objects."""
    photos = []
    header_lower = [str(h or "").strip().lower() for h in headers]
    
    for pn in [1, 2, 3]:
        prefix = f"p{pn}"
        local_path = original = thumb = mobile = web = hd = None
        
        for i, h in enumerate(header_lower):
            val = str(row[i] or "").strip() if i < len(row) and row[i] else ""
            if not val:
                continue
            if h.startswith(f"p{pn}") and "local" in h:
                local_path = val
            elif h.startswith(f"p{pn}") and "original" in h:
                original = val
            elif h.startswith(f"p{pn}") and "thumb" in h:
                thumb = val
            elif h.startswith(f"p{pn}") and "mobile" in h:
                mobile = val
            elif h.startswith(f"p{pn}") and ("web" in h or "1280" in h):
                web = val
            elif h.startswith(f"p{pn}") and ("hd" in h or "1920" in h):
                hd = val
        
        if original or thumb or mobile or web or hd:
            photos.append({
                "id": str(uuid.uuid4()),
                "is_primary": pn == 1,
                "sort_order": pn - 1,
                "local_path": local_path or "",
                "cloudinary_original": original or "",
                "cloudinary_thumb": thumb or "",
                "cloudinary_mobile": mobile or "",
                "cloudinary_web": web or "",
                "cloudinary_hd": hd or "",
                "source": "imported",
                "uploaded_at": datetime.now(timezone.utc).isoformat()
            })
    
    return photos

def count_valid_photos(photos):
    """Count photos that have at least one valid Cloudinary URL."""
    count = 0
    for p in photos:
        if isinstance(p, dict):
            has_url = any(
                p.get(k, "").startswith("https://res.cloudinary.com/")
                for k in ["cloudinary_original", "cloudinary_thumb", "cloudinary_mobile", "cloudinary_web", "cloudinary_hd"]
            )
            if has_url:
                count += 1
    return count

def extract_city_from_address(address):
    """Extract city from address - take second-to-last segment before Tanzania."""
    if not address:
        return ""
    parts = [p.strip() for p in address.split(",")]
    tz_cities = ["Dodoma", "Dar es Salaam", "Arusha", "Zanzibar", "Mwanza", "Mbeya", 
                 "Morogoro", "Tanga", "Iringa", "Moshi", "Musoma", "Bukoba", "Kigoma",
                 "Songea", "Lindi", "Mtwara", "Tabora", "Singida", "Shinyanga", "Kagera"]
    for c in tz_cities:
        if c.lower() in address.lower():
            return c
    # Try second-to-last segment
    if len(parts) >= 2:
        candidate = parts[-2].strip() if parts[-1].strip().lower() == "tanzania" else parts[-1].strip()
        if candidate and candidate.lower() != "tanzania":
            return candidate
    return ""

# Store latest preview for confirm step - using database now for Cloud Run compatibility
# _import_previews = {}  # Deprecated: moved to database storage

@api_router.post("/admin/import/preview")
async def preview_import(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    """Preview hotels from Excel file before importing."""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin pekee anayeweza kuingiza data")
    
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    
    # Auto-detect header row: find the row containing "Name" column
    header_row = 1
    data_start_row = 2
    for r in range(1, min(6, ws.max_row + 1)):
        row_vals = [str(ws.cell(row=r, column=c+1).value or "").strip().lower() for c in range(min(20, ws.max_column))]
        if "name" in row_vals:
            header_row = r
            data_start_row = r + 1
            break
    
    headers_raw = [ws.cell(row=header_row, column=c+1).value for c in range(ws.max_column)]
    headers = [str(h or "").strip().lower() for h in headers_raw]
    header_lower = headers
    
    # Map columns with more flexible matching
    col_map = {}
    for i, h in enumerate(headers):
        if h in ['name', 'hotel name', 'jina', 'hotel']:
            col_map['name'] = i
        elif h in ['address', 'anwani', 'location']:
            col_map['address'] = i
        elif h in ['phone', 'simu', 'phone number', 'telephone']:
            col_map['phone'] = i
        elif h in ['rating', 'google rating', 'stars', 'star rating']:
            col_map['rating'] = i
        elif h in ['description', 'maelezo', 'summary']:
            col_map['description'] = i
        elif h in ['price', 'bei', 'price level', 'base price']:
            col_map['price'] = i
        elif h in ['city', 'mji']:
            col_map['city'] = i
    
    previews = []
    errors = []
    duplicates = []
    
    # Check if photo columns exist
    has_photo_columns = any("p1" in h or "p2" in h or "p3" in h for h in header_lower)
    
    total_data_rows = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=data_start_row, values_only=True), start=data_start_row):
        if not row or not any(row):
            continue
        total_data_rows += 1
        row_num = row_idx
        try:
            # Get name from detected column or first column as fallback
            name_col = col_map.get("name", 0)
            name = str(row[name_col] or "").strip() if name_col < len(row) else ""
            if not name:
                continue
            
            # Check for duplicates
            existing = await crud.get_hotel_by_name(session, name)
            if existing:
                duplicates.append({"row": row_num, "name": name, "error": "Hoteli ipo tayari", "status": "duplicate", "issues": ["Duplicate hotel"]})
                continue
            
            address_col = col_map.get("address")
            address = str(row[address_col] or "").strip() if address_col is not None and address_col < len(row) else ""
            
            phone_col = col_map.get("phone")
            phone = str(row[phone_col] or "").strip() if phone_col is not None and phone_col < len(row) else ""
            
            rating = 0
            rating_col = col_map.get("rating")
            if rating_col is not None and rating_col < len(row):
                try:
                    rating = int(float(str(row[rating_col] or 0)))
                except:
                    rating = 0
            
            desc_col = col_map.get("description")
            description = str(row[desc_col] or "").strip() if desc_col is not None and desc_col < len(row) else ""
            
            price = 0
            price_col = col_map.get("price")
            if price_col is not None and price_col < len(row):
                try:
                    price = int(float(str(row[price_col] or 0).replace(",", "")))
                except:
                    price = 0
            
            photos = parse_cloudinary_photos_from_row(row, headers)
            valid_photo_count = count_valid_photos(photos)
            
            # Get city from column or extract from address
            city_col = col_map.get("city")
            city = str(row[city_col] or "").strip() if city_col is not None and city_col < len(row) else ""
            if not city:
                city = extract_city_from_address(address)
            
            cover = get_cover_url(photos, "cloudinary_web")
            
            # Determine status and issues
            issues = []
            if not city:
                issues.append("Missing city")
            if not phone:
                issues.append("Missing phone")
            if valid_photo_count == 0:
                issues.append("No photos")
            
            if issues:
                status = "warning"
            else:
                status = "valid"
            
            previews.append({
                "row": row_num,
                "name": name,
                "address": address,
                "city": city,
                "phone": phone,
                "rating": rating,
                "description": description,
                "base_price": price,
                "photos": photos,
                "photo_count": valid_photo_count,
                "cover_photo": cover,
                "status": status,
                "issues": issues
            })
        except Exception as e:
            errors.append({"row": row_num, "name": "", "error": str(e), "status": "error", "issues": [str(e)]})
    
    preview_id = str(uuid.uuid4())
    
    # Store preview in database for persistence across Cloud Run instances
    import_preview = ImportPreview(
        id=preview_id,
        preview_data=previews,
        created_by=current_user["id"],
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1)  # Expires in 1 hour
    )
    session.add(import_preview)
    await session.commit()
    
    # Build combined preview list for frontend (valid + warnings + duplicates + errors)
    all_preview_rows = previews + duplicates + errors
    all_preview_rows.sort(key=lambda x: x.get("row", 0))
    
    valid_count = len([p for p in previews if p["status"] == "valid"])
    warning_count = len([p for p in previews if p["status"] == "warning"])
    
    return {
        "preview_id": preview_id,
        "total_rows": total_data_rows,
        "valid": valid_count,
        "warnings": warning_count,
        "duplicates": len(duplicates),
        "error_count": len(errors),
        "has_photo_columns": has_photo_columns,
        "preview": all_preview_rows[:50]
    }

@api_router.post("/admin/import/execute")
async def execute_import(
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    """Execute import of previewed hotels."""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin pekee anayeweza kuingiza data")
    
    body = await request.json()
    preview_id = body.get("preview_id")
    selected_rows = body.get("selected_rows")  # List of row numbers to import
    
    # Fetch preview from database
    result = await session.execute(
        select(ImportPreview).where(ImportPreview.id == preview_id)
    )
    import_preview = result.scalar_one_or_none()
    
    if not import_preview:
        raise HTTPException(status_code=400, detail="Preview haipatikani au imeisha muda")
    
    # Check if preview has expired
    if import_preview.expires_at and import_preview.expires_at < datetime.now(timezone.utc):
        await session.delete(import_preview)
        await session.commit()
        raise HTTPException(status_code=400, detail="Preview imeisha muda, tafadhali pakia faili tena")
    
    previews = import_preview.preview_data
    if selected_rows:
        previews = [p for p in previews if p["row"] in selected_rows]
    
    imported = []
    failed = []
    batch_id = str(uuid.uuid4())
    batch_name = body.get("batch_name") or f"Import {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    city_override = body.get("city_override") or None
    batch_record = await crud.create_import_batch(session, {
        "id": batch_id,
        "name": batch_name,
        "city_override": city_override,
        "total": len(previews),
        "imported": 0,
        "skipped": 0,
        "errors": [],
        "created_by": current_user["id"]
    })
    
    skipped_duplicates = []
    for preview in previews:
        try:
            # Check for duplicates by name+city or phone
            existing_by_name = await crud.get_hotel_by_name_and_city(session, preview["name"], preview["city"])
            if existing_by_name:
                skipped_duplicates.append({
                    "row": preview["row"],
                    "name": preview["name"],
                    "reason": f"Hotel with same name already exists in {preview['city']}",
                    "existing_id": existing_by_name.id
                })
                continue
            
            existing_by_phone = await crud.get_hotel_by_phone(session, preview.get("phone", ""))
            if existing_by_phone:
                skipped_duplicates.append({
                    "row": preview["row"],
                    "name": preview["name"],
                    "reason": f"Hotel with phone {preview['phone']} already exists",
                    "existing_id": existing_by_phone.id
                })
                continue
            
            # Generate unique hotel code (checks database for duplicates)
            code = await generate_unique_hotel_code(session)
            
            # Create hotel with correct field names matching Hotel model
            hotel_data = {
                "name": preview["name"],
                "hotel_code": code,
                "address": preview["address"],
                "city": preview["city"],
                "phone_number": preview["phone"],
                "google_rating": float(preview["rating"]) if preview["rating"] else None,
                "description": preview["description"],
                "photos": preview["photos"],
                "cover_photo": preview["cover_photo"],
                "amenities": [],
                "status": "imported",
                "data_source": "bulk_import",
                "import_batch_id": batch_id
            }
            
            hotel = await crud.create_hotel(session, hotel_data)
            
            # Create default room type for this hotel
            await create_default_room_type(session, hotel.id, preview["base_price"])
            
            imported.append({
                "row": preview["row"],
                "name": preview["name"],
                "hotel_id": hotel.id,
                "code": code
            })
        except Exception as e:
            # Rollback to clear the pending transaction state
            await session.rollback()
            failed.append({
                "row": preview["row"],
                "name": preview["name"],
                "error": str(e)
            })
    
    # Update batch record with final counts
    batch_record.imported = len(imported)
    batch_record.skipped = len(failed) + len(skipped_duplicates)
    batch_record.errors = failed + skipped_duplicates
    
    # Clean up preview from database
    await session.delete(import_preview)
    await session.commit()
    
    return {
        "batch_id": batch_id,
        "imported": len(imported),
        "skipped": len(failed) + len(skipped_duplicates),
        "duplicates": len(skipped_duplicates),
        "failed": len(failed),
        "errors": failed[:20],
        "duplicate_details": skipped_duplicates[:20],
        "imported_hotels": imported[:20],
        "failed_imports": failed[:20]
    }

@api_router.get("/admin/import/batches")
async def get_import_batches(
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    """List import batch history."""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin pekee")
    
    batches = await crud.get_import_batches(session, limit=50)
    return [
        {
            "id": b.id,
            "name": b.name,
            "created_at": b.created_at.isoformat() if b.created_at else None,
            "total": b.total,
            "imported": b.imported,
            "skipped": b.skipped,
            "errors": b.errors
        }
        for b in batches
    ]

@api_router.get("/admin/import/batches/{batch_id}")
async def get_import_batch(
    batch_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    """Get import batch details."""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin pekee")
    
    batch = await crud.get_import_batch_by_id(session, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Batch haipatikani")
    
    return {
        "id": batch.id,
        "name": batch.name,
        "city_override": batch.city_override,
        "created_at": batch.created_at.isoformat() if batch.created_at else None,
        "created_by": batch.created_by,
        "total": batch.total,
        "imported": batch.imported,
        "skipped": batch.skipped,
        "errors": batch.errors
    }

@api_router.get("/admin/import/batches/{batch_id}/errors")
async def get_import_batch_errors(
    batch_id: str,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    """Get import batch errors."""
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Admin pekee")
    
    batch = await crud.get_import_batch_by_id(session, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Batch haipatikani")
    
    return batch.errors or []

@api_router.post("/admin/import/confirm")
async def confirm_import(
    request: Request,
    session: AsyncSession = Depends(get_db_session),
    current_user: dict = Depends(get_current_user)
):
    """Confirm and finalize import (backward compatibility)."""
    body = await request.json()
    preview_id = body.get("preview_id")
    
    # Check if preview exists in database
    result = await session.execute(
        select(ImportPreview).where(ImportPreview.id == preview_id)
    )
    import_preview = result.scalar_one_or_none()
    
    if not import_preview:
        return {"message": "Preview expired, please re-upload", "status": "expired"}
    
    # Same as execute
    return await execute_import(request, session, current_user)

# ===================== CONFIG ENDPOINTS =====================

@api_router.get("/config/cloudinary")
async def get_cloudinary_config():
    return {
        "cloud_name": CLOUDINARY_CLOUD_NAME,
        "upload_preset": CLOUDINARY_UPLOAD_PRESET,
        "folder": "hotel_research/owner_uploads"
    }

@api_router.get("/config/payment")
async def get_payment_config():
    return {
        "selcom_till": SELCOM_TILL_NUMBER,
        "expiry_seconds": BOOKING_EXPIRY_SECONDS,
        "cloudinary_cloud_name": CLOUDINARY_CLOUD_NAME,
        "cloudinary_upload_preset": CLOUDINARY_UPLOAD_PRESET
    }

# ===================== HEALTH CHECK =====================

@api_router.get("/health")
async def health_check():
    return {"status": "healthy", "service": "Habari Stays API", "version": "2.1.0", "database": "PostgreSQL"}

# Include router
app.include_router(api_router)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("startup")
async def startup_event():
    # Initialize database
    await init_db()
    logger.info("Database initialized")
    
    # Start booking expiry background task
    asyncio.create_task(expire_pending_bookings())
    logger.info("Started booking expiry background task")

@app.on_event("shutdown")
async def shutdown_event():
    await close_db()
    logger.info("Database connections closed")
